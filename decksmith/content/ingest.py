"""Content pack ingestion: files -> ContentCorpus (text chunks, tables, numeric facts).

Also provides deterministic BM25 retrieval of the chunks most relevant to the
brief, so large content packs fit the model context without embeddings.
"""
from __future__ import annotations

import csv
import hashlib
import math
import re
from collections import Counter
from pathlib import Path

from decksmith.core.models import ContentChunk, ContentCorpus, DataTable

NUM_RE = re.compile(r"(?<![\w.])[-−+]?\d[\d\s ]*(?:[.,]\d+)?\s?(?:%|₽|\$|€|млн|млрд|тыс\.?|k|m|b|x|×)?", re.I)
WORD_RE = re.compile(r"[a-zA-Zа-яА-ЯёЁ0-9]{2,}")


def _chunks(text: str, source: str, page: int | None, size: int = 1400) -> list[ContentChunk]:
    paras = [p.strip() for p in re.split(r"\n\s*\n|\n(?=[•●\-–]\s)", text) if p.strip()]
    out, buf = [], ""
    for p in paras:
        if len(buf) + len(p) > size and buf:
            out.append(buf)
            buf = ""
        buf += ("\n" if buf else "") + p
    if buf:
        out.append(buf)
    return [
        ContentChunk(id=f"{Path(source).stem[:20]}_{page or 0}_{i}", source=Path(source).name, text=t, page=page)
        for i, t in enumerate(out)
    ]


def read_pdf(path: Path) -> tuple[list[ContentChunk], list[DataTable]]:
    import pymupdf

    chunks = []
    with pymupdf.open(path) as doc:
        for i, page in enumerate(doc):
            text = page.get_text("text")
            text = re.sub(r"[ \t]+\n", "\n", text)
            chunks += _chunks(text, str(path), i + 1)
    return chunks, []


def read_docx(path: Path) -> tuple[list[ContentChunk], list[DataTable]]:
    import docx

    d = docx.Document(str(path))
    text = "\n\n".join(p.text for p in d.paragraphs if p.text.strip())
    tables = []
    for ti, t in enumerate(d.tables):
        rows = [[c.text.strip() for c in r.cells] for r in t.rows]
        if len(rows) >= 2:
            tables.append(DataTable(id=f"{path.stem}_t{ti}", source=path.name, columns=rows[0], rows=rows[1:]))
    return _chunks(text, str(path), None), tables


def read_pptx(path: Path) -> tuple[list[ContentChunk], list[DataTable]]:
    from pptx import Presentation

    prs = Presentation(str(path))
    chunks, tables = [], []
    for i, s in enumerate(prs.slides):
        texts = []
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                texts.append(sh.text_frame.text.strip())
            if getattr(sh, "has_table", False) and sh.has_table:
                rows = [[c.text for c in r.cells] for r in sh.table.rows]
                if len(rows) >= 2:
                    tables.append(DataTable(id=f"{path.stem}_s{i}", source=path.name, columns=rows[0], rows=rows[1:]))
        if texts:
            chunks += _chunks("\n".join(texts), str(path), i + 1)
    return chunks, tables


def read_table(path: Path) -> tuple[list[ContentChunk], list[DataTable]]:
    rows: list[list] = []
    if path.suffix.lower() == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as f:
            sample = f.read(4096)
            f.seek(0)
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t") if sample else csv.excel
            rows = [r for r in csv.reader(f, dialect)]
    else:
        import openpyxl

        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        tables = []
        for ws in wb.worksheets:
            rs = [[c for c in r] for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]
            if len(rs) >= 2:
                tables.append(DataTable(id=f"{path.stem}_{ws.title}", source=path.name, title=ws.title,
                                        columns=[str(c or "") for c in rs[0]], rows=[list(r) for r in rs[1:]]))
        text = "\n".join(f"Таблица {t.title}: " + ", ".join(t.columns) for t in tables)
        return _chunks(text, str(path), None), tables
    if len(rows) < 2:
        return [], []
    t = DataTable(id=path.stem, source=path.name, title=path.stem, columns=rows[0], rows=rows[1:])
    return _chunks(f"Таблица {path.stem}: " + ", ".join(rows[0]), str(path), None), [t]


def read_text(path: Path) -> tuple[list[ContentChunk], list[DataTable]]:
    return _chunks(path.read_text(encoding="utf-8", errors="ignore"), str(path), None), []


READERS = {".pdf": read_pdf, ".docx": read_docx, ".pptx": read_pptx, ".csv": read_table, ".xlsx": read_table,
           ".md": read_text, ".txt": read_text}


def normalize_number(s: str) -> str:
    t = s.strip().replace(" ", "").replace(" ", "").replace("−", "-").replace(",", ".")
    return t.lower()


def ingest(paths: list[str | Path], extra_text: str = "") -> ContentCorpus:
    chunks: list[ContentChunk] = []
    tables: list[DataTable] = []
    files = []
    h = hashlib.sha256()
    for p in paths:
        p = Path(p)
        reader = READERS.get(p.suffix.lower())
        if not reader:
            continue
        c, t = reader(p)
        chunks += c
        tables += t
        files.append(str(p))
        h.update(p.read_bytes())
    if extra_text.strip():
        chunks += _chunks(extra_text, "brief.txt", None)
        h.update(extra_text.encode())
    numbers = sorted({normalize_number(m.group(0)) for c in chunks for m in NUM_RE.finditer(c.text) if any(ch.isdigit() for ch in m.group(0))})
    for t in tables:
        for r in t.rows:
            for v in r:
                if v is not None and re.search(r"\d", str(v)):
                    numbers.append(normalize_number(str(v)))
    cyr = sum(len(re.findall(r"[а-яё]", c.text, re.I)) for c in chunks)
    lat = sum(len(re.findall(r"[a-z]", c.text, re.I)) for c in chunks)
    return ContentCorpus(id=h.hexdigest()[:12], files=files, chunks=chunks, tables=tables, numbers=sorted(set(numbers)),
                         language="ru" if cyr >= lat else "en")


# ----------------------------------------------------------------------------
# Retrieval
# ----------------------------------------------------------------------------
def _tokens(text: str) -> list[str]:
    return [w.lower()[:7] for w in WORD_RE.findall(text)]  # crude stemming by prefix


def select_context(corpus: ContentCorpus, query: str, budget_chars: int = 18000) -> str:
    """BM25 top chunks (kept in document order) within a character budget."""
    if not corpus.chunks:
        return ""
    total = sum(len(c.text) for c in corpus.chunks)
    if total <= budget_chars:
        return "\n\n".join(f"[{c.source}{' p.' + str(c.page) if c.page else ''}]\n{c.text}" for c in corpus.chunks)
    docs = [_tokens(c.text) for c in corpus.chunks]
    df = Counter(t for d in docs for t in set(d))
    n = len(docs)
    avgdl = sum(len(d) for d in docs) / n
    q = _tokens(query)
    scores = []
    for i, d in enumerate(docs):
        tf = Counter(d)
        s = 0.0
        for t in q:
            if t not in tf:
                continue
            idf = math.log(1 + (n - df[t] + 0.5) / (df[t] + 0.5))
            s += idf * tf[t] * 2.2 / (tf[t] + 1.2 * (0.25 + 0.75 * len(d) / avgdl))
        scores.append((s + 0.01 * (n - i) / n, i))  # slight preference for early chunks
    chosen, used = set(), 0
    for s, i in sorted(scores, reverse=True):
        L = len(corpus.chunks[i].text)
        if used + L > budget_chars:
            continue
        chosen.add(i)
        used += L
    return "\n\n".join(f"[{corpus.chunks[i].source}{' p.' + str(corpus.chunks[i].page) if corpus.chunks[i].page else ''}]\n{corpus.chunks[i].text}"
                       for i in sorted(chosen))
