"""Модель входа аудита: на что может смотреть проверка (факты файла + рендер + правила шаблона)."""
from __future__ import annotations

from dataclasses import dataclass, field
from functools import cached_property
from pathlib import Path

from PIL import Image
from pptx import Presentation

from decksmith.core.models import Box, ContentCorpus, DeckPlan, TemplateProfile
from decksmith.layout.textfit import Measure, measure, measure_rich
from decksmith.parsing.elements import Element, extract_elements
from decksmith.parsing.ooxml import parse_theme, rgb_to_hex


def measure_element(e: Element, size_factor: float = 1.0) -> Measure:
    """Высота/число строк текста элемента так, как его раскладывают рендереры: каждый абзац со своим кеглем,
    насыщенностью, интервалами и межстрочным интервалом, внутри отступов рамки.
    """
    l, t, r, b = e.insets
    w = e.box.w - (l + r - 2 * 91440)
    base = e.style.size if e.style and e.style.size else 14
    font = (e.style.font if e.style else None) or "Arial"
    if e.para_styles and len(e.para_styles) == len(e.paragraphs):
        styles = [((s.size or base) * size_factor, s.bold, s.space_before, s.space_after, s.line, s.font or font)
                  for s in e.para_styles]
        return measure_rich(e.paragraphs, styles, w, inset_lr=91440, inset_tb=(t + b) // 2)
    return measure(e.paragraphs or [e.text], font, base * size_factor, w, bool(e.style and e.style.bold),
                   inset_lr=91440, inset_tb=(t + b) // 2)


@dataclass
class SlideFacts:
    index: int
    slide: object
    layout_name: str
    elements: list[Element]
    png: Path | None = None
    kind: str = ""  # вид паттерна / вид композиции, использованный сборщиком

    @property
    def texts(self) -> list[Element]:
        """Непустые текстовые элементы слайда."""
        return [e for e in self.elements if e.kind == "text" and e.text.strip()]

    @property
    def content(self) -> list[Element]:
        """Блоки, несущие содержание (не декор, не колонтитулы/номера страниц)."""
        return [e for e in self.elements if not e.brand and ((e.kind == "text" and e.text.strip()) or e.kind in ("table", "chart"))]

    def _measure(self, e: Element):
        """Замер текста элемента по реальным метрикам шрифта."""
        return measure_element(e)

    def effective_box(self, e: Element) -> Box:
        """Рамка, выросшая до измеренной высоты текста (рамки с автоподбором растут при рендере)."""
        if e.kind != "text" or not e.style:
            return e.box
        m = self._measure(e)
        if m.height_emu > e.box.h and e.autofit:
            return Box(x=e.box.x, y=e.box.y, w=e.box.w, h=m.height_emu)
        return e.box

    def glyph_box(self, e: Element) -> Box:
        """Где на самом деле стоят строки текста: измеренная высота, размещённая по вертикальной привязке
        рамки (текст, вытекающий из фиксированной рамки, всё равно рисуется, поэтому учитывается вся
        высота).
        """
        if e.kind != "text" or not e.style:
            return e.box
        h = self._measure(e).height_emu
        if e.autofit and h > e.box.h:
            return Box(x=e.box.x, y=e.box.y, w=e.box.w, h=h)
        y = e.box.y if e.anchor == "t" else (e.box.b - h if e.anchor == "b" else e.box.y + (e.box.h - h) // 2)
        return Box(x=e.box.x, y=y, w=e.box.w, h=max(h, 1))

    def text_extent(self, e: Element) -> Box:
        """Площадь, реально покрытая глифами (для доли заполнения)."""
        if e.kind != "text" or not e.style:
            return e.box
        m = self._measure(e)
        return Box(x=e.box.x, y=e.box.y, w=min(e.box.w, m.width_emu), h=min(max(e.box.h, 0), m.height_emu) if not e.autofit else m.height_emu)

    def lines(self, e: Element) -> tuple[int, int]:
        """(нужно строк, строк помещается в рамку)."""
        m = self._measure(e)
        _, t, _, b = e.insets
        line_h = (e.style.size or 14) * 1.2 * 12700
        fit = max(1, int(round((e.box.h - t - b) / line_h + 0.25)))
        return m.lines, fit

    def text_height_needed(self, e: Element) -> int:
        """Высота, нужная тексту элемента, в EMU."""
        return self._measure(e).height_emu


@dataclass
class AuditContext:
    pptx: Path
    profile: TemplateProfile
    plan: DeckPlan | None = None
    corpus: ContentCorpus | None = None
    pngs: list[Path] = field(default_factory=list)
    slide_kinds: list[str] = field(default_factory=list)
    # слайд шаблона, из которого клонирован каждый слайд результата
    slide_sources: list[int | None] = field(default_factory=list)
    brief_text: str = ""
    render_ok: bool = True

    @cached_property
    def template_prs(self):
        """Презентация шаблона (для сравнения с примерами)."""
        return Presentation(self.profile.file)

    def template_element(self, i: int, shape_id: int) -> Element | None:
        """Та же фигура на слайде-примере шаблона (для проверок «унаследовано от шаблона»)."""
        src = self.slide_sources[i] if i < len(self.slide_sources) else None
        if src is None:
            return None
        cache = self.__dict__.setdefault("_tpl_cache", {})
        if src not in cache:
            s = self.template_prs.slides[src]
            t = self.profile.tokens
            cache[src] = {e.shape_id: e for e in extract_elements(s, parse_theme(s.slide_layout.slide_master), t.slide_w, t.slide_h)}
        return cache[src].get(shape_id)

    @cached_property
    def prs(self):
        """Проверяемая презентация."""
        return Presentation(str(self.pptx))

    @cached_property
    def slides(self) -> list[SlideFacts]:
        """Факты по каждому слайду: элементы, макет, вид слайда, фон."""
        t = self.profile.tokens
        out = []
        for i, s in enumerate(self.prs.slides):
            theme = parse_theme(s.slide_layout.slide_master)
            els = extract_elements(s, theme, t.slide_w, t.slide_h, include_empty_placeholders=False)
            out.append(SlideFacts(index=i, slide=s, layout_name=s.slide_layout.name, elements=els,
                                  png=self.pngs[i] if i < len(self.pngs) else None,
                                  kind=self.slide_kinds[i] if i < len(self.slide_kinds) else ""))
        return out

    def image(self, i: int) -> Image.Image | None:
        """Рендер i-го слайда в RGB или None, если рендера нет."""
        if i < len(self.pngs) and self.pngs[i] and Path(self.pngs[i]).exists():
            return Image.open(self.pngs[i]).convert("RGB")
        return None

    def sample_background(self, i: int, box: Box) -> str | None:
        """Преобладающий цвет рендера внутри рамки (пиксели текста в меньшинстве)."""
        im = self.image(i)
        if im is None:
            return None
        t = self.profile.tokens
        sx, sy = im.width / t.slide_w, im.height / t.slide_h
        x0, y0 = max(0, int(box.x * sx)), max(0, int(box.y * sy))
        x1, y1 = min(im.width, int(box.r * sx)), min(im.height, int(box.b * sy))
        if x1 - x0 < 2 or y1 - y0 < 2:
            return None
        crop = im.crop((x0, y0, x1, y1)).resize((min(48, x1 - x0), min(24, y1 - y0)))
        q = crop.quantize(colors=4, method=Image.Quantize.MEDIANCUT)
        counts = sorted(q.getcolors() or [], reverse=True)
        if not counts:
            return None
        pal = q.getpalette()
        idx = counts[0][1]
        return rgb_to_hex((pal[idx * 3], pal[idx * 3 + 1], pal[idx * 3 + 2]))
