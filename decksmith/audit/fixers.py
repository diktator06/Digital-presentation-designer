"""Fixers for audit issues. Deterministic ones edit the PPTX directly;
contextual ones (fix_llm / shorten_llm) go through the versioned `fixer` skill.

apply_fixes() never mutates the input file: it writes deck_vN+1.pptx so the
UI can show before/after and roll back.
"""
from __future__ import annotations

import logging
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn

from decksmith.audit.rules import PLACEHOLDER_RE
from decksmith.core.models import TABLE_MAX_COLS, TABLE_MAX_ROWS, AuditIssue, TemplateProfile
from decksmith.generation.llm import LLMClient
from decksmith.generation.skills import load_skill
from decksmith.layout.compose import readable_on
from decksmith.layout.pptx_ops import delete_shape, delete_slides, scale_font_sizes, set_font_size, set_paragraphs, shape_by_id
from decksmith.audit.context import measure_element
from decksmith.layout.room import grow_frame, text_room
from decksmith.layout.textfit import max_chars_for, snap_down
from decksmith.parsing.elements import extract_elements
from decksmith.parsing.ooxml import color_distance, parse_theme

log = logging.getLogger(__name__)


class FixContext:
    def __init__(self, prs, profile: TemplateProfile, backgrounds: dict[tuple[int, int], str] | None = None):
        self.prs = prs
        self.profile = profile
        self.slides = list(prs.slides)
        self.backgrounds = backgrounds or {}
        self.to_drop: set[int] = set()

    def element(self, si: int, shape_id: int):
        s = self.slides[si]
        t = self.profile.tokens
        for e in extract_elements(s, parse_theme(s.slide_layout.slide_master), t.slide_w, t.slide_h, False):
            if e.shape_id == shape_id:
                return e
        return None


def _runs(shape):
    txb = shape._element.find(qn("p:txBody"))
    return list(txb.iter(qn("a:r"))) if txb is not None else []


# ----------------------------------------------------------------------------
def _fits_at(e, factor: float) -> bool:
    """The text, measured paragraph by paragraph with its spacing, fits the frame at size x factor."""
    m = measure_element(e, factor)
    l, _, r, _ = e.insets
    return m.height_emu <= e.box.h * 1.02 and m.longest_word_emu <= e.box.w - l - r


def fx_shrink_text(fc: FixContext, iss: AuditIssue) -> bool:
    """Smaller type until the text fits. A frame that mixes sizes (bold lead over body text)
    is scaled by one factor, so its hierarchy survives; a uniform one steps down the scale."""
    ok = False
    for sid in iss.shape_ids:
        sh = shape_by_id(fc.slides[iss.slide], sid)
        e = fc.element(iss.slide, sid)
        if sh is None or e is None or e.kind != "text" or not e.style:
            continue
        size = e.style.size or 14
        if len({round(p.size or size, 1) for p in e.para_styles}) > 1:
            f = next((k / 100 for k in range(95, 55, -5) if _fits_at(e, k / 100)), 0.6)
            scale_font_sizes(sh, f, default_size=size)
            ok = True
            continue
        scale = fc.profile.tokens.type_scale.sizes
        for s in sorted({x for x in scale if 0.6 * size <= x < size}, reverse=True):
            if _fits_at(e, s / size):
                set_font_size(sh, s)
                ok = True
                break
        else:
            set_font_size(sh, snap_down(size * 0.7, scale))
            ok = True
    return ok


def fx_grow_frame(fc: FixContext, iss: AuditIssue) -> bool:
    """Organisers' clarification: resizing a text frame inside its block is not a violation —
    so the frame grows to its text first; the type shrinks only when the block has no room."""
    ok = False
    for sid in iss.shape_ids:
        slide = fc.slides[iss.slide]
        sh = shape_by_id(slide, sid)
        e = fc.element(iss.slide, sid)
        if sh is None or e is None or e.kind != "text" or not e.style:
            continue
        li = next((l.index for l in fc.profile.layouts if l.name == slide.slide_layout.name), None)
        own = {f"layout:{li}", f"master:{fc.profile.layouts[li].master_index}"} if li is not None else set()
        zones = [b.box for b in fc.profile.brand_elements
                 if b.kind in ("logo", "footer", "page_number") and own & set(b.source.split(","))]
        room = text_room(slide, sh, fc.profile.tokens, keep_clear=zones)
        need = measure_element(e).height_emu
        if room is not None and need <= room[2] - room[1] and grow_frame(sh, room, need):
            ok = True
            continue
        ok = fx_shrink_text(fc, iss.model_copy(update={"shape_ids": [sid]})) or ok
    return ok


def fx_move_inside(fc: FixContext, iss: AuditIssue) -> bool:
    t = fc.profile.tokens
    m = t.margins
    ok = False
    for sid in iss.shape_ids:
        sh = shape_by_id(fc.slides[iss.slide], sid)
        if sh is None or sh.left is None:
            continue
        x0, y0 = m.left // 2, m.top // 2
        x1, y1 = t.slide_w - m.right // 2, t.slide_h - m.bottom // 2
        w, h = min(sh.width, x1 - x0), min(sh.height, y1 - y0)
        sh.width, sh.height = int(w), int(h)
        sh.left = int(min(max(sh.left, x0), x1 - w))
        sh.top = int(min(max(sh.top, y0), y1 - h))
        ok = True
    return ok


def fx_snap_align(fc: FixContext, iss: AuditIssue) -> bool:
    if len(iss.shape_ids) < 2:
        return False
    a = shape_by_id(fc.slides[iss.slide], iss.shape_ids[0])
    b = shape_by_id(fc.slides[iss.slide], iss.shape_ids[1])
    if a is None or b is None:
        return False
    b.left = a.left  # the first block (usually the title) is the guide
    return True


def fx_set_template_font(fc: FixContext, iss: AuditIssue) -> bool:
    t = fc.profile.tokens
    for sid in iss.shape_ids:
        sh = shape_by_id(fc.slides[iss.slide], sid)
        if sh is None:
            continue
        for r in _runs(sh):
            rpr = r.find(qn("a:rPr"))
            if rpr is None:
                continue
            lat = rpr.find(qn("a:latin"))
            if lat is not None:
                lat.set("typeface", t.body_font)
    return bool(iss.shape_ids)


def fx_snap_size(fc: FixContext, iss: AuditIssue) -> bool:
    """Off-scale runs move to the nearest size of the template scale; a larger one only if the text
    still fits its frame (otherwise the next smaller one). Other runs of the frame keep their size."""
    size, nearest = iss.data.get("size"), iss.data.get("nearest")
    if not size or not nearest:
        return False
    below = max((s for s in fc.profile.tokens.type_scale.sizes if s <= size), default=None)
    ok = False
    for sid in iss.shape_ids:
        sh = shape_by_id(fc.slides[iss.slide], sid)
        e = fc.element(iss.slide, sid)
        if sh is None:
            continue
        target = float(nearest)
        if target > size and below and e is not None and e.kind == "text" and e.style and not _fits_at(e, target / size):
            target = float(below)
        txb = sh._element.find(qn("p:txBody"))
        for el in (list(txb.iter(qn("a:rPr"))) + list(txb.iter(qn("a:endParaRPr")))) if txb is not None else []:
            cur = int(el.get("sz")) / 100 if el.get("sz") else None
            if cur is None or abs(cur - size) < 0.3:
                el.set("sz", str(int(round(target * 100))))
                ok = True
    return ok


def _nearest_palette(fc: FixContext, hex_: str) -> str:
    pal = [c.hex for c in fc.profile.tokens.palette]
    return min(pal, key=lambda p: color_distance(p, hex_)) if pal else hex_


def fx_snap_color(fc: FixContext, iss: AuditIssue) -> bool:
    c = iss.data.get("color")
    if not c:
        return False
    new = _nearest_palette(fc, c)
    for sid in iss.shape_ids:
        sh = shape_by_id(fc.slides[iss.slide], sid)
        if sh is None:
            continue
        for el in sh._element.iter(qn("a:srgbClr")):
            if el.get("val", "").upper() == c.upper():
                el.set("val", new)
    return True


def fx_fix_contrast(fc: FixContext, iss: AuditIssue) -> bool:
    bg = iss.data.get("bg")
    if not bg:
        return False
    pal = [c.hex for c in fc.profile.tokens.palette[:10]] + ["FFFFFF", "000000"]
    new = readable_on(bg, pal)
    for sid in iss.shape_ids:
        sh = shape_by_id(fc.slides[iss.slide], sid)
        if sh is None:
            continue
        for r in _runs(sh):
            rpr = r.find(qn("a:rPr"))
            if rpr is None:
                continue
            for f in rpr.findall(qn("a:solidFill")):
                rpr.remove(f)
            sf = rpr.makeelement(qn("a:solidFill"), {})
            clr = sf.makeelement(qn("a:srgbClr"), {"val": new})
            sf.append(clr)
            rpr.insert(0, sf)
    return True


def fx_trim_bullets(fc: FixContext, iss: AuditIssue) -> bool:
    for sid in iss.shape_ids:
        sh = shape_by_id(fc.slides[iss.slide], sid)
        if sh is None or not sh.has_text_frame:
            continue
        paras = [p.text for p in sh.text_frame.paragraphs if p.text.strip()][:6]
        set_paragraphs(sh, paras)
    return True


def fx_trim_table(fc: FixContext, iss: AuditIssue) -> bool:
    for sid in iss.shape_ids:
        sh = shape_by_id(fc.slides[iss.slide], sid)
        if sh is None or not getattr(sh, "has_table", False):
            continue
        tbl = sh._element.graphic.graphicData.tbl
        rows = tbl.findall(qn("a:tr"))
        for tr in rows[TABLE_MAX_ROWS:]:
            tbl.remove(tr)
        grid = tbl.find(qn("a:tblGrid"))
        cols = grid.findall(qn("a:gridCol"))
        if len(cols) > TABLE_MAX_COLS:
            for gc in cols[TABLE_MAX_COLS:]:
                grid.remove(gc)
            for tr in tbl.findall(qn("a:tr")):
                for tc in tr.findall(qn("a:tc"))[TABLE_MAX_COLS:]:
                    tr.remove(tc)
    return True


def fx_remove_placeholder(fc: FixContext, iss: AuditIssue) -> bool:
    for sid in iss.shape_ids:
        sh = shape_by_id(fc.slides[iss.slide], sid)
        if sh is None or not sh.has_text_frame:
            continue
        keep = [p.text for p in sh.text_frame.paragraphs if p.text.strip() and not PLACEHOLDER_RE.search(p.text)]
        if keep:
            set_paragraphs(sh, keep)
        else:
            delete_shape(sh)
    return True


def fx_drop_slide(fc: FixContext, iss: AuditIssue) -> bool:
    fc.to_drop.add(iss.slide)
    return True


def fx_fix_aspect(fc: FixContext, iss: AuditIssue) -> bool:
    for sid in iss.shape_ids:
        sh = shape_by_id(fc.slides[iss.slide], sid)
        if sh is None:
            continue
        iw, ih = sh.image.size
        box_r, img_r = sh.width / sh.height, iw / ih
        sh.crop_left = sh.crop_right = sh.crop_top = sh.crop_bottom = 0
        if img_r > box_r:
            c = (1 - box_r / img_r) / 2
            sh.crop_left = sh.crop_right = c
        else:
            c = (1 - img_r / box_r) / 2
            sh.crop_top = sh.crop_bottom = c
    return True


DETERMINISTIC = {
    "shrink_text": fx_shrink_text,
    "grow_frame": fx_grow_frame,
    "move_inside": fx_move_inside,
    "snap_align": fx_snap_align,
    "set_template_font": fx_set_template_font,
    "snap_size": fx_snap_size,
    "snap_color": fx_snap_color,
    "fix_contrast": fx_fix_contrast,
    "trim_bullets": fx_trim_bullets,
    "trim_table": fx_trim_table,
    "remove_placeholder": fx_remove_placeholder,
    "drop_slide": fx_drop_slide,
    "fix_aspect": fx_fix_aspect,
}
AUTO_SAFE = {"shrink_text", "grow_frame", "remove_placeholder", "trim_table", "trim_bullets", "fix_aspect", "move_inside"}


async def fx_contextual(fc: FixContext, issues: list[AuditIssue], llm: LLMClient, language: str) -> int:
    """Group LLM-fixable issues per slide and ask the `fixer` skill once per slide."""
    if not llm.enabled:
        return 0
    skill = load_skill("fixer")
    by_slide: dict[int, list[AuditIssue]] = {}
    for i in issues:
        by_slide.setdefault(i.slide, []).append(i)
    changed = 0
    for si, iss in by_slide.items():
        slide = fc.slides[si]
        texts, budgets, shapes = {}, {}, {}
        t = fc.profile.tokens
        for e in extract_elements(slide, parse_theme(slide.slide_layout.slide_master), t.slide_w, t.slide_h, False):
            if e.kind == "text" and e.text.strip():
                key = f"sh{e.shape_id}"
                texts[key] = e.paragraphs or [e.text]
                budgets[key] = max_chars_for(e.box.w, e.box.h, e.style.font or "Arial", e.style.size or 14) if e.style else 200
                shapes[key] = e.shape_id
        try:
            res = await llm.run_skill(skill, None, language=language, texts=texts, issues=[i.model_dump() for i in iss], budgets=budgets)
        except Exception as ex:
            log.warning("fixer failed on slide %d: %s", si + 1, ex)
            continue
        for key, paras in (res or {}).get("texts", {}).items():
            if key in shapes and isinstance(paras, list) and paras:
                sh = shape_by_id(slide, shapes[key])
                if sh is not None:
                    set_paragraphs(sh, [str(p) for p in paras])
                    changed += 1
    return changed


async def apply_fixes(pptx: str | Path, issues: list[AuditIssue], profile: TemplateProfile, out: str | Path,
                      llm: LLMClient | None = None, language: str = "ru") -> dict:
    prs = Presentation(str(pptx))
    fc = FixContext(prs, profile)
    applied, skipped, contextual = [], [], []
    for iss in issues:
        if not iss.fix:
            skipped.append(iss.id)
            continue
        fn = DETERMINISTIC.get(iss.fix)
        if fn is None:
            contextual.append(iss)
            continue
        try:
            (applied if fn(fc, iss) else skipped).append(iss.id)
        except Exception as e:
            log.warning("fixer %s failed: %s", iss.fix, e)
            skipped.append(iss.id)
    n_ctx = 0
    if contextual and llm is not None:
        n_ctx = await fx_contextual(fc, contextual, llm, language)
        applied += [i.id for i in contextual] if n_ctx else []
        skipped += [] if n_ctx else [i.id for i in contextual]
    elif contextual:
        skipped += [i.id for i in contextual]
    if fc.to_drop:
        delete_slides(prs, sorted(fc.to_drop))
    out = Path(out)
    prs.save(str(out))
    return {"output": str(out), "applied": applied, "skipped": skipped, "contextual_changed": n_ctx}
