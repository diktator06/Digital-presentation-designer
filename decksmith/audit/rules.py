"""Deterministic audit checks (same slide -> same result).

They rely only on what is in the file (coordinates, sizes, color codes, layout
references), on the template profile (rules) and on the rendered PNG (pixels).
Each check is registered with an id, category and the fixer that can repair it.
"""
from __future__ import annotations

import colorsys
import re
from dataclasses import dataclass
from typing import Callable

from pptx.oxml.ns import qn

from decksmith.audit.context import AuditContext, SlideFacts
from decksmith.content.ingest import NUM_RE, normalize_number
from decksmith.core.models import TABLE_MAX_COLS, TABLE_MAX_ROWS, AuditIssue, Box, Severity
from decksmith.parsing.elements import Element
from decksmith.parsing.ooxml import color_distance, contrast_ratio, hex_to_rgb


@dataclass
class CheckMeta:
    id: str
    category: str
    title: str
    fixer: str | None
    fn: Callable[[AuditContext], list[AuditIssue]]
    deterministic: bool = True


CHECKS: dict[str, CheckMeta] = {}
COVER_KINDS = {"title", "section", "thanks", "quote", "contacts"}


def check(cid: str, category: str, title: str, fixer: str | None = None):
    def deco(fn):
        CHECKS[cid] = CheckMeta(cid, category, title, fixer, fn)
        return fn
    return deco


_counter = [0]


def issue(cid: str, slide: int, message: str, *, boxes=(), shape_ids=(), severity=Severity.warning, data=None) -> AuditIssue:
    meta = CHECKS[cid]
    _counter[0] += 1
    return AuditIssue(
        id=f"{cid}#{slide}#{_counter[0]}", check=cid, category=meta.category, deterministic=True, severity=severity,
        slide=slide, message=message, boxes=list(boxes), shape_ids=list(shape_ids), fixable=meta.fixer is not None,
        fix=meta.fixer, data=data or {},
    )


def _tol(ctx: AuditContext, frac: float = 0.004) -> int:
    return int(ctx.profile.tokens.slide_w * frac)


def _same_box(a: Box, b: Box, tol: int = 12700 * 2) -> bool:
    return abs(a.x - b.x) <= tol and abs(a.y - b.y) <= tol and abs(a.w - b.w) <= tol and abs(a.h - b.h) <= tol


def inherited_geometry(ctx: AuditContext, sf: SlideFacts, e: Element) -> bool:
    """True when the element sits exactly where the template put it (layout placeholder
    geometry or the example slide it was cloned from): a property of the template itself."""
    tpl = ctx.template_element(sf.index, e.shape_id)
    if tpl is not None and _same_box(tpl.box, e.box):
        return True
    if e.placeholder:
        for ph in sf.slide.slide_layout.placeholders:
            try:
                same_type = str(ph.placeholder_format.type).split(".")[-1].split(" ")[0] == e.placeholder
                box = Box(x=int(ph.left), y=int(ph.top), w=int(ph.width), h=int(ph.height))
            except Exception:
                continue
            if same_type and _same_box(box, e.box):
                return True
    return False


# ============================================================================ layout
@check("layout.out_of_bounds", "layout", "Элемент вышел за границы слайда", fixer="move_inside")
def out_of_bounds(ctx):
    t = ctx.profile.tokens
    out, tol = [], _tol(ctx)
    for sf in ctx.slides:
        for e in sf.content:
            b = e.box
            if b.x < -tol or b.y < -tol or b.r > t.slide_w + tol or b.b > t.slide_h + tol:
                inh = inherited_geometry(ctx, sf, e)
                out.append(issue("layout.out_of_bounds", sf.index,
                                 f"«{e.text[:40] or e.kind}» выходит за край слайда" + (" (так задано в шаблоне)" if inh else ""),
                                 boxes=[b], shape_ids=[e.shape_id], severity=Severity.info if inh else Severity.error,
                                 data={"inherited": inh}))
    return out


@check("layout.overlap", "layout", "Два блока наложились друг на друга", fixer="shrink_text")
def overlap(ctx):
    """Colliding text lines / objects are an error; frames that only overlap (lines apart) are a
    layout-hygiene warning. Both are template design when the frames sit exactly where the
    template example put them and the text did not grow them."""
    out = []
    for sf in ctx.slides:
        # text is compared where its lines are (a card-sized frame may contain another by design)
        blocks = [(e, sf.glyph_box(e) if e.kind == "text" else sf.effective_box(e)) for e in sf.content]
        for i in range(len(blocks)):
            for j in range(i + 1, len(blocks)):
                (a, ga), (b, gb) = blocks[i], blocks[j]
                frames = a.box.intersection(b.box) > 0.08 * (min(a.box.area, b.box.area) or 1)
                lines = ga.intersection(gb) > 0.08 * (min(ga.area, gb.area) or 1)
                if not frames and not lines:
                    continue
                inherited = inherited_geometry(ctx, sf, a) and inherited_geometry(ctx, sf, b)
                grown = not (a.box.contains(ga, tol=12700) and b.box.contains(gb, tol=12700))
                if lines:
                    inh = frames and inherited and not grown
                    sev, what = (Severity.info if inh else Severity.error), " (так в шаблоне)" if inh else ""
                elif inherited or a.box.contains(b.box, tol=_tol(ctx)) or b.box.contains(a.box, tol=_tol(ctx)):
                    continue  # frames nest (template design, or a frame grown inside its card), the lines are apart
                else:
                    sev, what = Severity.warning, " (рамки; строки не пересекаются)"
                inter = (ga if lines else a.box).intersection(gb if lines else b.box)
                small = min((ga if lines else a.box).area, (gb if lines else b.box).area) or 1
                out.append(issue("layout.overlap", sf.index,
                                 f"«{a.text[:30] or a.kind}» перекрывает «{b.text[:30] or b.kind}»" + what,
                                 boxes=[ga, gb] if lines else [a.box, b.box], shape_ids=[a.shape_id, b.shape_id],
                                 severity=sev, data={"ratio": round(inter / small, 2), "inherited": sev == Severity.info,
                                                     "lines": lines}))
    return out


@check("layout.text_overflow", "layout", "Текст не поместился в свою рамку", fixer="grow_frame")
def text_overflow(ctx):
    """The fixer first resizes the frame inside its block (allowed by the organisers'
    clarification), and shrinks the type only when the block has no room left."""
    t = ctx.profile.tokens
    area, tol = t.slide_w * t.slide_h, _tol(ctx)
    out = []
    for sf in ctx.slides:
        plates = _plates(sf, area)
        for e in sf.texts:
            if not e.style or e.autofit:
                continue
            need_lines, fit_lines = sf.lines(e)
            if need_lines > fit_lines and not _leaves_block(sf, e, plates, tol):
                need = sf.text_height_needed(e)
                out.append(issue("layout.text_overflow", sf.index,
                                 f"Текст «{e.text[:40]}» занимает {need_lines} стр. при месте на {fit_lines}",
                                 boxes=[Box(x=e.box.x, y=e.box.y, w=e.box.w, h=max(need, e.box.h))], shape_ids=[e.shape_id],
                                 data={"need": need, "have": e.box.h, "lines": need_lines, "fit": fit_lines}))
    return out


def _plates(sf: SlideFacts, area: int) -> list[Element]:
    """Visible surfaces that can frame text: filled/outlined shapes and pictures (not icons)."""
    return [e for e in sf.elements if not e.brand and 0 < e.box.area < 0.85 * area and
            ((e.kind == "decor" and not e.graphic and (e.fill_hex or e.has_line)) or (e.kind == "picture" and not e.is_icon))]


def text_block(sf: SlideFacts, e: Element, plates: list[Element], tol: int) -> Element | None:
    """The block a text belongs to: the smallest plate spanning the frame's width and holding its top."""
    best = None
    for p in plates:
        b = p.box
        if b.area > e.box.area and b.x - tol <= e.box.x and e.box.r <= b.r + tol and b.y - tol <= e.box.y < b.b:
            best = p if best is None or b.area < best.box.area else best
    return best


def _leaves_block(sf: SlideFacts, e: Element, plates: list[Element], tol: int) -> bool:
    """The text spills out of its block: `layout.block_overflow` reports it (one finding per defect)."""
    block = text_block(sf, e, plates, tol)
    if block is None:
        return False
    g = sf.glyph_box(e)
    return max(g.b - block.box.b, block.box.y - g.y) > tol


@check("layout.block_overflow", "layout", "Текст вышел за пределы своего блока (плашки, карточки)", fixer="shrink_text")
def block_overflow(ctx):
    """Organisers' clarification of «текст не поместился в свою рамку»: the frame may grow inside
    its block, the text must stay within the block. Lines are placed by the frame's anchor
    (auto-fit frames grow downwards). The template's own example spilling the same way is design."""
    t = ctx.profile.tokens
    area, tol = t.slide_w * t.slide_h, _tol(ctx)
    out = []
    for sf in ctx.slides:
        plates = _plates(sf, area)
        for e in sf.texts:
            if e.brand or not e.style:
                continue
            block = text_block(sf, e, plates, tol)
            if block is None:
                continue
            g = sf.glyph_box(e)
            over = max(g.b - block.box.b, block.box.y - g.y)
            if over <= tol:
                continue
            tpl = ctx.template_element(sf.index, e.shape_id)
            inherited = tpl is not None and inherited_geometry(ctx, sf, e) and tpl.style is not None and (
                max(sf.glyph_box(tpl).b - block.box.b, block.box.y - sf.glyph_box(tpl).y) >= over - tol)
            out.append(issue("layout.block_overflow", sf.index,
                             f"Текст «{e.text[:40]}» выходит за свой блок на {over / 12700:.0f} pt" + (" (так в шаблоне)" if inherited else ""),
                             boxes=[g, block.box], shape_ids=[e.shape_id],
                             severity=Severity.info if inherited else Severity.error,
                             data={"over": over, "block": block.shape_id, "inherited": inherited}))
    return out


@check("layout.text_clipped", "layout", "Текст обрезан краем слайда", fixer="shrink_text")
def text_clipped(ctx):
    t = ctx.profile.tokens
    out = []
    for sf in ctx.slides:
        for e in sf.texts:
            eb = sf.effective_box(e)
            if eb.b > t.slide_h + _tol(ctx) or eb.r > t.slide_w + _tol(ctx):
                # the frame itself sits across the edge in the template and our text did not grow it
                inh = inherited_geometry(ctx, sf, e) and (e.box.b > t.slide_h or e.box.r > t.slide_w) and eb.area <= e.box.area * 1.05
                out.append(issue("layout.text_clipped", sf.index,
                                 f"Текст «{e.text[:40]}» уходит за край слайда" + (" (рамка так стоит в шаблоне)" if inh else ""),
                                 boxes=[eb], shape_ids=[e.shape_id], severity=Severity.info if inh else Severity.error,
                                 data={"inherited": inh}))
    return out


@check("layout.misaligned", "layout", "Блоки не выровнены по направляющим", fixer="snap_align")
def misaligned(ctx):
    """Left edges that are almost-but-not-quite equal (0.2%..1.5% of width) are misalignments."""
    t = ctx.profile.tokens
    lo, hi = 0.002 * t.slide_w, 0.015 * t.slide_w
    out = []
    for sf in ctx.slides:
        blocks = [e for e in sf.content if e.box.w > 0.08 * t.slide_w]
        for i in range(len(blocks)):
            for j in range(i + 1, len(blocks)):
                a, b = blocks[i], blocks[j]
                ax = a.box.x + (a.insets[0] if a.kind == "text" else 0)
                bx = b.box.x + (b.insets[0] if b.kind == "text" else 0)
                d = abs(ax - bx)
                if lo < d < hi and not (a.box.intersection(b.box)):
                    out.append(issue("layout.misaligned", sf.index, f"Левые края «{a.text[:20] or a.kind}» и «{b.text[:20] or b.kind}» расходятся на {d / 12700:.1f} pt",
                                     boxes=[a.box, b.box], shape_ids=[a.shape_id, b.shape_id], severity=Severity.info,
                                     data={"dx": d}))
    return out


@check("layout.margins", "layout", "Контент заходит в поля у краёв", fixer="move_inside")
def margins(ctx):
    t = ctx.profile.tokens
    m = t.margins
    out = []
    for sf in ctx.slides:
        for e in sf.content:
            b = e.box
            bad = b.x < m.left * 0.5 or b.r > t.slide_w - m.right * 0.5 or b.y < m.top * 0.4
            if bad and b.w < 0.95 * t.slide_w:
                out.append(issue("layout.margins", sf.index, f"«{e.text[:30] or e.kind}» заходит в поле слайда",
                                 boxes=[b], shape_ids=[e.shape_id], severity=Severity.info))
    return out


@check("layout.image_distorted", "layout", "Картинка растянута, пропорции нарушены", fixer="fix_aspect")
def image_distorted(ctx):
    out = []
    for sf in ctx.slides:
        for sh in sf.slide.shapes:
            if sh.shape_type != 13:
                continue
            try:
                iw, ih = sh.image.size
            except Exception:
                continue
            crop = sh.crop_left + sh.crop_right, sh.crop_top + sh.crop_bottom
            vis_w, vis_h = iw * (1 - crop[0]), ih * (1 - crop[1])
            if not (vis_w and vis_h and sh.width and sh.height):
                continue
            ratio = (sh.width / sh.height) / (vis_w / vis_h)
            if abs(ratio - 1) > 0.04:
                b = Box(x=int(sh.left), y=int(sh.top), w=int(sh.width), h=int(sh.height))
                tpl = ctx.template_element(sf.index, sh.shape_id)
                inherited = tpl is not None and tpl.kind == "picture" and abs(tpl.box.w - b.w) < 12700 and abs(tpl.box.h - b.h) < 12700
                out.append(issue("layout.image_distorted", sf.index,
                                 f"Изображение искажено в {ratio:.2f} раза" + (" (так в шаблоне)" if inherited else ""),
                                 boxes=[b], shape_ids=[sh.shape_id], severity=Severity.info if inherited else Severity.warning,
                                 data={"ratio": round(ratio, 3), "inherited": inherited}))
    return out


# ============================================================================ template
def _allowed_fonts(ctx) -> set[str]:
    fams = {f.family.lower() for f in ctx.profile.tokens.fonts}
    fams |= {ctx.profile.tokens.heading_font.lower(), ctx.profile.tokens.body_font.lower()}
    return fams


@check("template.font", "template", "Шрифт не из шаблона или гарнитур больше двух", fixer="set_template_font")
def fonts(ctx):
    allowed = _allowed_fonts(ctx)
    out = []
    for sf in ctx.slides:
        fams = set()
        for e in sf.texts:
            f = (e.style.font if e.style else None) or ""
            fams.add(f.lower())
            if f and f.lower() not in allowed:
                out.append(issue("template.font", sf.index, f"Шрифт «{f}» отсутствует в шаблоне", boxes=[e.box], shape_ids=[e.shape_id]))
        if len(fams - {""}) > 2:
            out.append(issue("template.font", sf.index, f"На слайде {len(fams)} гарнитуры: {', '.join(sorted(fams))}",
                             data={"families": sorted(fams)}))
    return out


@check("template.type_scale", "template", "Кегль не из типографической шкалы шаблона", fixer="snap_size")
def type_scale(ctx):
    scale = ctx.profile.tokens.type_scale.sizes
    out = []
    if not scale:
        return out
    for sf in ctx.slides:
        for e in sf.texts:
            if not e.style or not e.style.size:
                continue
            nearest = min(scale, key=lambda s: abs(s - e.style.size))
            if abs(nearest - e.style.size) > 0.6:
                out.append(issue("template.type_scale", sf.index, f"Кегль {e.style.size:g} pt вне шкалы (ближайший {nearest:g})",
                                 boxes=[e.box], shape_ids=[e.shape_id], severity=Severity.info, data={"size": e.style.size, "nearest": nearest}))
    return out


def _sat_light(h: str):
    r, g, b = (c / 255 for c in hex_to_rgb(h))
    hh, l, s = colorsys.rgb_to_hls(r, g, b)
    return hh, l, s


def color_in_palette(ctx, hex_: str) -> bool:
    pal = [c.hex for c in ctx.profile.tokens.palette] + ["FFFFFF", "000000"]
    if any(color_distance(hex_, p) < 40 for p in pal):
        return True
    h, l, s = _sat_light(hex_)
    if s < 0.12 or l > 0.93 or l < 0.07:
        return True  # neutrals / near-white / near-black
    for p in pal:  # tint or shade of a palette hue
        ph, pl, ps = _sat_light(p)
        if ps > 0.2 and min(abs(ph - h), 1 - abs(ph - h)) < 0.04:
            return True
    return False


@check("template.palette", "template", "Цвет не из палитры шаблона", fixer="snap_color")
def palette(ctx):
    out = []
    for sf in ctx.slides:
        for e in sf.elements:
            cols = []
            if e.kind == "text" and e.style and e.style.color and e.text.strip():
                cols.append(("текст", e.style.color))
            if e.fill_hex and e.kind in ("text", "decor"):
                cols.append(("заливка", e.fill_hex))
            for what, c in cols:
                if not color_in_palette(ctx, c):
                    out.append(issue("template.palette", sf.index, f"{what} #{c} не из палитры шаблона", boxes=[e.box],
                                     shape_ids=[e.shape_id], severity=Severity.info, data={"color": c}))
    return out


@check("template.layout", "template", "Слайд собран не на макете из шаблона")
def layout_from_template(ctx):
    names = {l.name for l in ctx.profile.layouts}
    return [issue("template.layout", sf.index, f"Макет «{sf.layout_name}» отсутствует в шаблоне", severity=Severity.error)
            for sf in ctx.slides if sf.layout_name not in names]


@check("template.brand_zone", "template", "Логотип или колонтитул перекрыт/сдвинут", fixer="move_inside")
def brand_zone(ctx):
    out = []
    brand = [b for b in ctx.profile.brand_elements if b.kind in ("logo", "footer", "page_number")]
    for sf in ctx.slides:
        layout_idx = next((l.index for l in ctx.profile.layouts if l.name == sf.layout_name), None)
        # footer / page-number placeholders of a layout are drawn only if the slide carries them
        on_slide = {e.placeholder for e in sf.elements if e.placeholder in ("FOOTER", "SLIDE_NUMBER", "DATE")}
        mine = [b for b in brand if (f"layout:{layout_idx}" in b.source or "master" in b.source)
                and (b.kind == "logo" or {"footer": "FOOTER", "page_number": "SLIDE_NUMBER"}[b.kind] in on_slide)]
        for e in sf.content:
            if e.brand:
                continue
            for b in mine:
                inter = sf.effective_box(e).intersection(b.box)
                if inter > 0.15 * max(b.box.area, 1):
                    inh = inherited_geometry(ctx, sf, e) and sf.effective_box(e).area <= e.box.area * 1.05
                    out.append(issue("template.brand_zone", sf.index,
                                     f"«{e.text[:30] or e.kind}» перекрывает {b.kind} шаблона" + (" (так в шаблоне)" if inh else ""),
                                     boxes=[b.box, e.box], shape_ids=[e.shape_id], severity=Severity.info if inh else Severity.warning,
                                     data={"inherited": inh}))
    return out


@check("template.contrast", "template", "Контраст текста к фону ниже 4.5:1", fixer="fix_contrast")
def contrast(ctx):
    out = []
    if not ctx.pngs:
        return out
    for sf in ctx.slides:
        for e in sf.texts:
            if not e.style or not e.style.color:
                continue
            bg = ctx.sample_background(sf.index, e.box)
            if not bg:
                continue
            cr = contrast_ratio(e.style.color, bg)
            large = (e.style.size or 0) >= 18 or ((e.style.size or 0) >= 14 and e.style.bold)
            need = 3.0 if large else 4.5
            if cr < need and color_distance(e.style.color, bg) > 1:
                # the same colour on the same shape in the template example, or a placeholder whose
                # colour comes from the layout/master (we never set it) = a rule of the template itself
                tpl = ctx.template_element(sf.index, e.shape_id)
                inherited = bool(tpl and tpl.style and tpl.style.color == e.style.color) or e.style.color in _template_text_colors(ctx) \
                    or (e.placeholder is not None and not _explicit_run_color(sf.slide, e.shape_id))
                out.append(issue("template.contrast", sf.index,
                                 f"Контраст {cr:.1f}:1 (#{e.style.color} на #{bg}), нужно ≥ {need}:1" + (" — цвет задан шаблоном" if inherited else ""),
                                 boxes=[e.box], shape_ids=[e.shape_id], severity=Severity.info if inherited else Severity.warning,
                                 data={"ratio": round(cr, 2), "bg": bg, "fg": e.style.color, "inherited": inherited}))
    return out


def _explicit_run_color(slide, shape_id: int) -> bool:
    from decksmith.layout.pptx_ops import shape_by_id

    sh = shape_by_id(slide, shape_id)
    if sh is None:
        return False
    return any(r.find(qn("a:rPr")) is not None and r.find(qn("a:rPr")).find(qn("a:solidFill")) is not None
               for r in sh._element.iter(qn("a:r")))


def _template_text_colors(ctx) -> set[str]:
    return {c.hex for c in ctx.profile.tokens.palette if "text" in c.roles}


# ============================================================================ density
def _is_list(e: Element) -> bool:
    return len(e.paragraphs) >= 2


@check("density.bullets", "density", "Больше 6 буллетов на слайде", fixer="trim_bullets")
def bullets(ctx):
    out = []
    for sf in ctx.slides:
        for e in sf.texts:
            if len(e.paragraphs) > 6:
                out.append(issue("density.bullets", sf.index, f"{len(e.paragraphs)} пунктов в одном блоке (> 6)", boxes=[e.box],
                                 shape_ids=[e.shape_id], data={"n": len(e.paragraphs)}))
    return out


@check("density.bullet_words", "density", "Буллет длиннее 15 слов", fixer="shorten_llm")
def bullet_words(ctx):
    out = []
    for sf in ctx.slides:
        for e in sf.texts:
            if not _is_list(e):
                continue
            for p in e.paragraphs:
                n = len(p.split())
                if n > 15:
                    out.append(issue("density.bullet_words", sf.index, f"Пункт из {n} слов: «{p[:50]}…»", boxes=[e.box],
                                     shape_ids=[e.shape_id], severity=Severity.info, data={"words": n}))
    return out


@check("density.table", "density", "Таблица больше 10 строк или 5 колонок", fixer="trim_table")
def table_size(ctx):
    """TZ Appendix 1 names 7 rows; per the organisers' clarification tables up to 10 rows pass."""
    out = []
    for sf in ctx.slides:
        for e in sf.elements:
            if e.kind == "table" and e.table_shape:
                r, c = e.table_shape
                if r > TABLE_MAX_ROWS or c > TABLE_MAX_COLS:
                    out.append(issue("density.table", sf.index, f"Таблица {r}×{c} (макс. {TABLE_MAX_ROWS} строк × {TABLE_MAX_COLS} колонок)",
                                     boxes=[e.box], shape_ids=[e.shape_id]))
    return out


@check("density.chart_series", "density", "Больше 5 серий на диаграмме")
def chart_series(ctx):
    out = []
    for sf in ctx.slides:
        for sh in sf.slide.shapes:
            if getattr(sh, "has_chart", False) and sh.has_chart:
                n = sum(len(p.series) for p in sh.chart.plots)
                if n > 5:
                    out.append(issue("density.chart_series", sf.index, f"{n} серий на диаграмме", shape_ids=[sh.shape_id]))
    return out


@check("density.fill", "density", "Слайд заполнен меньше чем на ¼ или больше чем на ¾")
def fill(ctx):
    t = ctx.profile.tokens
    m = t.margins
    area = (t.slide_w - m.left - m.right) * (t.slide_h - m.top - m.bottom)
    out = []
    for sf in ctx.slides:
        if sf.kind in COVER_KINDS:
            continue
        boxes = [sf.text_extent(e) for e in sf.content]
        boxes += [e.box for e in sf.elements if e.kind == "picture" and not e.is_icon and e.name in ("Illustration", "Picture")]
        # union area on a coarse grid
        g = 40
        cells = set()
        for b in boxes:
            for gx in range(max(0, int(b.x / t.slide_w * g)), min(g, int(b.r / t.slide_w * g) + 1)):
                for gy in range(max(0, int(b.y / t.slide_h * g)), min(g, int(b.b / t.slide_h * g) + 1)):
                    cells.add((gx, gy))
        ratio = len(cells) * (t.slide_w / g) * (t.slide_h / g) / max(area, 1)
        if ratio < 0.25:
            out.append(issue("density.fill", sf.index, f"Заполнено ~{ratio:.0%} полезной площади (< 25%)", severity=Severity.info,
                             data={"ratio": round(ratio, 2)}))
        elif ratio > 0.75:
            out.append(issue("density.fill", sf.index, f"Заполнено ~{min(ratio, 1):.0%} полезной площади (> 75%)", severity=Severity.info,
                             data={"ratio": round(ratio, 2)}))
    return out


# ============================================================================ integrity
@check("integrity.open", "integrity", "Файл не открывается")
def opens(ctx):
    try:
        n = len(ctx.prs.slides)
    except Exception as e:
        return [issue("integrity.open", 0, f"python-pptx не открыл файл: {e}", severity=Severity.error)]
    if ctx.pngs and len(ctx.pngs) != n:
        return [issue("integrity.open", 0, f"Рендер дал {len(ctx.pngs)} страниц из {n}", severity=Severity.error)]
    if not ctx.render_ok:
        return [issue("integrity.open", 0, "LibreOffice не смог отрисовать файл", severity=Severity.error)]
    return []


PLACEHOLDER_RE = re.compile(
    r"lorem ipsum|\bXXX+\b|\bTODO\b|\bTBD\b|вставьте текст|введите текст|click to add|нажмите, чтобы|"
    r"^\s*(заголовок|подзаголовок|текст|пункт|описание|имя фамилия|должность|показатель)\s*$|\[.*?(вставить|insert).*?\]",
    re.I | re.M,
)


@check("integrity.placeholder_text", "integrity", "Остался текст-заглушка", fixer="remove_placeholder")
def placeholder_text(ctx):
    out = []
    for sf in ctx.slides:
        for e in sf.texts:
            for p in e.paragraphs or [e.text]:
                if PLACEHOLDER_RE.search(p):
                    out.append(issue("integrity.placeholder_text", sf.index, f"Заглушка: «{p[:50]}»", boxes=[e.box],
                                     shape_ids=[e.shape_id], severity=Severity.error, data={"text": p}))
                    break
    return out


@check("integrity.empty", "integrity", "Пустой слайд или слайд с одним заголовком")
def empty(ctx):
    out = []
    for sf in ctx.slides:
        if sf.kind in COVER_KINDS:
            continue
        non_title = [e for e in sf.content if not e.is_title_ph]
        pics = [e for e in sf.elements if e.kind == "picture" and not e.is_icon]
        if len(non_title) == 0 and not pics:
            out.append(issue("integrity.empty", sf.index, "На слайде только заголовок", severity=Severity.error))
    return out


@check("integrity.raster", "integrity", "Слайд оказался картинкой, а не редактируемыми объектами")
def raster(ctx):
    t = ctx.profile.tokens
    out = []
    for sf in ctx.slides:
        texts = sf.texts
        big = [e for e in sf.elements if e.kind == "picture" and e.box.area > 0.85 * t.slide_w * t.slide_h]
        if big and not texts:
            out.append(issue("integrity.raster", sf.index, "Слайд состоит из одного растрового изображения", severity=Severity.error))
    return out


@check("integrity.chart_labels", "integrity", "У диаграммы нет подписей осей, единиц или легенды")
def chart_labels(ctx):
    out = []
    for sf in ctx.slides:
        for sh in sf.slide.shapes:
            if not (getattr(sh, "has_chart", False) and sh.has_chart):
                continue
            ch = sh.chart
            xml = ch._chartSpace.xml
            n_series = sum(len(p.series) for p in ch.plots)
            pie = "pieChart" in xml or "doughnutChart" in xml
            has_axis_title = "<c:title>" in xml.split("<c:valAx>")[-1] if "<c:valAx>" in xml else pie
            has_labels = "<c:dLbls>" in xml and 'c:showVal val="1"' in xml
            problems = []
            if not pie and not (has_axis_title or has_labels):
                problems.append("нет подписи оси/единиц")
            if (n_series > 1 or pie) and not ch.has_legend:
                problems.append("нет легенды")
            if problems:
                b = Box(x=int(sh.left), y=int(sh.top), w=int(sh.width), h=int(sh.height))
                out.append(issue("integrity.chart_labels", sf.index, "Диаграмма: " + ", ".join(problems), boxes=[b], shape_ids=[sh.shape_id]))
    return out


def _words(sf: SlideFacts) -> set[str]:
    return {w.lower() for e in sf.texts for w in re.findall(r"[a-zA-Zа-яА-ЯёЁ]{4,}", e.text)}


@check("integrity.duplicate", "integrity", "Два слайда дублируют друг друга", fixer="drop_slide")
def duplicate(ctx):
    out = []
    ws = [_words(sf) for sf in ctx.slides]
    for i in range(len(ws)):
        for j in range(i + 1, len(ws)):
            a, b = ws[i], ws[j]
            if len(a) >= 6 and len(b) >= 6:
                jac = len(a & b) / len(a | b)
                if jac > 0.7:
                    out.append(issue("integrity.duplicate", j, f"Слайд {j + 1} повторяет слайд {i + 1} ({jac:.0%} общих слов)",
                                     data={"other": i, "jaccard": round(jac, 2)}))
    return out


def _norm_words(s: str) -> list[str]:
    return re.findall(r"[a-zA-Zа-яА-ЯёЁ0-9]+", s.lower())


def _present(fragment: str, haystack_words: set[str], min_hit: float = 0.6) -> bool:
    words = [w for w in _norm_words(fragment) if len(w) > 2 or any(ch.isdigit() for ch in w)][:6]
    if not words:
        return True
    return sum(w in haystack_words for w in words) / len(words) >= min_hit


def _slide_words(sf) -> set[str]:
    words = set()
    for e in sf.texts:
        words |= set(_norm_words(e.text))
    for sh in sf.slide.shapes:
        if getattr(sh, "has_table", False) and sh.has_table:
            for r in sh.table.rows:
                for c in r.cells:
                    words |= set(_norm_words(c.text))
        if getattr(sh, "has_chart", False) and sh.has_chart:
            try:
                for p in sh.chart.plots:
                    words |= set(_norm_words(" ".join(str(c) for c in p.categories)))
            except Exception:
                pass
    return words


@check("integrity.title_missing", "integrity", "На слайде нет заголовка из плана")
def title_missing(ctx):
    if not ctx.plan:
        return []
    out = []
    for sf in ctx.slides:
        if sf.index >= len(ctx.plan.slides):
            break
        spec = ctx.plan.slides[sf.index]
        title = ctx.plan.title if spec.intent.value == "title" else (spec.quote if spec.intent.value == "quote" and spec.quote else spec.title)
        if title and not _present(title, _slide_words(sf), 0.5):
            out.append(issue("integrity.title_missing", sf.index, f"Заголовок «{title[:50]}» не выведен на слайд", severity=Severity.error))
    return out


@check("integrity.content_lost", "integrity", "Часть контента плана не попала на слайд")
def content_lost(ctx):
    """Every bullet / item / table row of the plan must be visible (silent drops are layout bugs)."""
    if not ctx.plan:
        return []
    out = []
    for sf in ctx.slides:
        if sf.index >= len(ctx.plan.slides):
            break
        spec = ctx.plan.slides[sf.index]
        frags = list(spec.bullets[:6])
        frags += [it.title or it.text or it.value for it in spec.items[:8]]
        frags += [it.value for it in spec.items[:8] if it.value and it.title]  # KPI values are the point
        if spec.table:
            frags += [str(r[0]) for r in spec.table.rows[:TABLE_MAX_ROWS - 1] if r]
        if spec.chart:
            frags += spec.chart.categories[:8]
        words = _slide_words(sf)
        missing = [f for f in frags if f and not _present(f, words)]
        if missing:
            out.append(issue("integrity.content_lost", sf.index, f"Не видно {len(missing)} из {len(frags)} элементов: «{missing[0][:40]}»…",
                             severity=Severity.error if len(missing) > len(frags) / 2 else Severity.warning,
                             data={"missing": missing[:10]}))
    return out


# ============================================================================ content (deterministic part)
@check("content.language", "content", "Вся колода на одном языке", fixer="fix_llm")
def language(ctx):
    lang = ctx.plan.language if ctx.plan else "ru"
    out = []
    for sf in ctx.slides:
        text = " ".join(e.text for e in sf.texts)
        words = re.findall(r"[a-zA-Zа-яА-ЯёЁ]{3,}", text)
        if len(words) < 5:
            continue
        lat = [w for w in words if re.match(r"[a-zA-Z]", w) and not w.isupper()]
        share = len(lat) / len(words)
        if (lang == "ru" and share > 0.35) or (lang == "en" and share < 0.65):
            out.append(issue("content.language", sf.index, f"Смешение языков: {share:.0%} латиницы", severity=Severity.info,
                             data={"latin_share": round(share, 2)}))
    return out


@check("content.numbers_sourced", "content", "Цифры на слайде есть в исходных материалах", fixer="fix_llm")
def numbers_sourced(ctx):
    if not ctx.corpus:
        return []
    known = set(ctx.corpus.numbers)
    known |= {normalize_number(m.group(0)) for m in NUM_RE.finditer(ctx.brief_text or "")}
    known_digits = {re.sub(r"[^\d.]", "", k) for k in known}
    out = []
    for sf in ctx.slides:
        for e in sf.texts:
            if e.brand:
                continue  # page numbers, dates, footers are not content
            for m in NUM_RE.finditer(e.text):
                raw = m.group(0).strip()
                digits = re.sub(r"[^\d.]", "", normalize_number(raw))
                if not digits or len(digits.replace(".", "")) < 2:
                    continue  # step numbers 1..9
                if digits.lstrip("0") in {d.lstrip("0") for d in known_digits}:
                    continue
                if re.fullmatch(r"(19|20)\d\d", digits):
                    continue  # years are validated contextually
                out.append(issue("content.numbers_sourced", sf.index, f"Число «{raw}» не найдено в материалах", boxes=[e.box],
                                 shape_ids=[e.shape_id], data={"number": raw}))
    return out


def run_rules(ctx: AuditContext, only: list[str] | None = None) -> tuple[list[AuditIssue], list[str]]:
    issues: list[AuditIssue] = []
    ran = []
    for cid, meta in CHECKS.items():
        if only and cid not in only:
            continue
        try:
            issues += meta.fn(ctx)
            ran.append(cid)
        except Exception as e:  # a broken check must not break the pipeline
            issues.append(AuditIssue(id=f"{cid}#err", check=cid, category=meta.category, slide=0,
                                     message=f"проверка упала: {e}", severity=Severity.info))
    return issues, ran
