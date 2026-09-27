"""Оркестратор разбора шаблона: .pptx -> TemplateProfile (кэш по sha256).

Этапы (все детерминированные):
  1. структура  : мастера, макеты, плейсхолдеры, темы
  2. рендер     : слайды шаблона + пробная колода «пустых макетов» через LibreOffice
  3. CV         : яркость / цвет фона каждого слайда, свободная зона контента каждого
                  макета по сетке занятости отрендеренного пустого макета
  4. паттерны   : по каждому слайду-примеру (+ макеты с большим числом плейсхолдеров)
  5. токены     : палитра, шкала кеглей, поля, сетка
  6. бренд      : логотипы / колонтитулы / номера страниц, закреплённые мастерами и макетами
  7. ресурсы    : наборы иконок со слайдов-инструкций

Необязательный контекстный этап (разметка паттернов с помощью VLM) находится в
generation/labeler.py и лишь уточняет `Pattern.kind` и теги.
"""
from __future__ import annotations

import hashlib
import io
import json
import logging
import re
import shutil
import subprocess
import tempfile
import time
import zipfile
from collections import Counter
from pathlib import Path

from PIL import Image, ImageFilter, ImageStat
from pptx import Presentation
from pptx.oxml.ns import qn

from decksmith.core.config import settings
from decksmith.core.fonts import DOWNLOAD_BUDGET_S, ensure_font
from decksmith.core.models import (
    Box,
    BrandElement,
    IconAsset,
    LayoutInfo,
    Pattern,
    PatternKind,
    TemplateProfile,
)
from decksmith.parsing.elements import Element, extract_elements
from decksmith.parsing.ooxml import Theme, color_distance, paints, parse_theme, rgb_to_hex
from decksmith.parsing.patterns import build_pattern
from decksmith.parsing.tokens import build_tokens
from decksmith.render.soffice import RenderError, pdf_to_pngs, pptx_to_pdf, soffice_path

log = logging.getLogger(__name__)


def _parser_version() -> str:
    """Ключ кэша: меняется при любом изменении модулей разбора."""
    h = hashlib.sha256()
    for p in sorted(Path(__file__).parent.glob("*.py")):
        h.update(p.read_bytes())
    return "0.4-" + h.hexdigest()[:10]


PARSER_VERSION = _parser_version()


def file_sha256(path: str | Path) -> str:
    """sha256 файла."""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def all_layouts(prs) -> list[tuple[int, int, object]]:
    """Все макеты: (глобальный индекс, индекс мастера, макет)."""
    out = []
    for mi, m in enumerate(prs.slide_masters):
        for layout in m.slide_layouts:
            out.append((len(out), mi, layout))
    return out


def layout_global_index(prs, layout) -> int:
    """Глобальный индекс макета среди всех мастеров."""
    for gi, _, lay in all_layouts(prs):
        if lay.part is layout.part:
            return gi
    return -1


# ----------------------------------------------------------------------------
# Вспомогательные функции CV
# ----------------------------------------------------------------------------
def image_stats(png: Path) -> tuple[bool, str, float]:
    """(тёмный ли, hex фона по краю, средняя яркость) отрендеренного слайда."""
    im = Image.open(png).convert("RGB")
    small = im.resize((160, 90))
    lum = ImageStat.Stat(small.convert("L")).mean[0] / 255
    w, h = small.size
    ring = []
    px = small.load()
    for x in range(w):
        ring += [px[x, 1], px[x, h - 2]]
    for y in range(h):
        ring += [px[1, y], px[w - 2, y]]
    ring.sort(key=lambda c: sum(c))
    med = ring[len(ring) // 2]
    return lum < 0.42, rgb_to_hex(med), round(lum, 3)


def occupancy_grid(png: Path, bg_hex: str, gw: int = 64, gh: int = 36, edges_only: bool = False) -> list[list[bool]]:
    """Клетки отрендеренного слайда, на которых что-то нарисовано. Клетка занята, если в ней есть края (по
    каждому каналу RGB, поэтому границы, отличающиеся только оттенком, тоже считаются) или цвет, далёкий от
    фона. Области далёкого цвета почти без краёв по границе — это плавные свечения и градиенты самого фона,
    а не объекты: они остаются свободными. Плоские панели имеют резкие границы и остаются занятыми.
    """
    im = Image.open(png).convert("RGB")
    small = im.resize((gw * 4 + 2, gh * 4 + 2))  # поле в 1 px: фильтр краёв даёт кольцо по границе изображения
    edges = small.filter(ImageFilter.FIND_EDGES)
    px, ex = small.load(), edges.load()
    edge = [[False] * gw for _ in range(gh)]
    far = [[False] * gw for _ in range(gh)]
    for gy in range(gh):
        for gx in range(gw):
            ne = nf = 0
            for dy in range(4):
                for dx in range(4):
                    x, y = gx * 4 + dx + 1, gy * 4 + dy + 1
                    if max(ex[x, y]) > 40:
                        ne += 1
                    elif not edges_only and color_distance(rgb_to_hex(px[x, y]), bg_hex) > 120:
                        nf += 1
            edge[gy][gx] = ne >= 3
            far[gy][gx] = not edge[gy][gx] and ne + nf >= 3
    if edges_only:
        return edge
    occ = [[edge[gy][gx] or far[gy][gx] for gx in range(gw)] for gy in range(gh)]
    seen = [[False] * gw for _ in range(gh)]
    for sy in range(gh):
        for sx in range(gw):
            if not far[sy][sx] or seen[sy][sx]:
                continue
            comp, stack = [], [(sx, sy)]
            seen[sy][sx] = True
            while stack:
                x, y = stack.pop()
                comp.append((x, y))
                for nx, ny in ((x + 1, y), (x - 1, y), (x, y + 1), (x, y - 1)):
                    if 0 <= nx < gw and 0 <= ny < gh and far[ny][nx] and not seen[ny][nx]:
                        seen[ny][nx] = True
                        stack.append((nx, ny))
            members = set(comp)
            border = hard = 0
            for x, y in comp:
                outside = [(nx, ny) for nx, ny in ((x + 1, y), (x - 1, y), (x, y + 1), (x, y - 1))
                           if 0 <= nx < gw and 0 <= ny < gh and (nx, ny) not in members]
                if outside:
                    border += 1
                    hard += any(edge[ny][nx] for nx, ny in outside)
            if border and hard < 0.25 * border:
                for x, y in comp:
                    occ[y][x] = False
    return occ


def title_clear_box(png: Path, box: Box, slide_w: int, slide_h: int, bg_hex: str) -> Box | None:
    """Часть рамки заголовка, свободная от фоновой графики (полосы логотипов в фоновой картинке, угловая
    графика): заголовок не должен заходить под неё.
    """
    gw, gh = 64, 36
    occ = occupancy_grid(png, bg_hex, gw, gh, edges_only=True)  # у градиентов нет краёв, у логотипов есть
    # полоса расширяется на строку в каждую сторону: графика, лишь задевающая рамку, всё равно
    # сталкивается с текстом, прижатым к этому краю; две занятые клетки блокируют колонку (это не случайный
    # край)
    y0, y1 = max(0, int(box.y / slide_h * gh) - 1), min(gh, int(box.b / slide_h * gh) + 2)
    x0, x1 = max(0, int(box.x / slide_w * gw)), min(gw, int(box.r / slide_w * gw))
    # длинные горизонтальные отрезки — линейки и подчёркивания, а не графика, которую надо обходить
    run_min = max(8, int(0.15 * gw))
    for gy in range(y0, y1):
        gx = 0
        while gx < gw:
            if not occ[gy][gx]:
                gx += 1
                continue
            end = gx
            while end < gw and occ[gy][end]:
                end += 1
            if end - gx >= run_min:
                for k in range(gx, end):
                    occ[gy][k] = False
            gx = end
    start = x0 + max(2, (x1 - x0) // 3)  # в первой трети может стоять собственная акцентная графика заголовка
    for gx in range(start, x1):
        if sum(occ[gy][gx] for gy in range(y0, y1)) >= 2:
            new_r = int(gx / gw * slide_w) - int(0.01 * slide_w)
            if new_r - box.x >= 0.45 * box.w:
                return Box(x=box.x, y=box.y, w=new_r - box.x, h=box.h)
            return None
    return None


def free_content_box(png: Path, slide_w: int, slide_h: int, below_y: int, margins, bg_hex: str) -> tuple[Box | None, bool]:
    """Наибольший пустой прямоугольник ниже `below_y` на отрендеренном пустом макете -> (рамка, текстура).
    Занятость (`occupancy_grid`) считается на сетке 64×36; наибольший полностью свободный прямоугольник
    находится методом гистограмм. Фон, занятый почти везде (фото, текстура, сетка чертежа), помечается как
    `textured`: тогда годится вся область под заголовком, а компоновщик ставит контент на подложку для
    читаемости.
    """
    gw, gh = 64, 36
    occ = occupancy_grid(png, bg_hex, gw, gh)
    y0 = max(0, int(below_y / slide_h * gh) + 1)
    x_min = int(margins.left / slide_w * gw)
    x_max = gw - int(margins.right / slide_w * gw)
    y_max = gh - int(margins.bottom / slide_h * gh)
    cells = [(gx, gy) for gy in range(y0, y_max) for gx in range(x_min, x_max)]
    busy_share = sum(occ[gy][gx] for gx, gy in cells) / max(len(cells), 1)
    default = Box(x=int(x_min / gw * slide_w), y=int(y0 / gh * slide_h), w=int((x_max - x_min) / gw * slide_w),
                  h=int(max(y_max - y0, 1) / gh * slide_h))
    if busy_share > 0.45:
        return default, True
    best = (0, 0, 0, 0, 0)  # площадь, x, y, w, h
    heights = [0] * gw
    for gy in range(y0, y_max):
        for gx in range(gw):
            heights[gx] = heights[gx] + 1 if (x_min <= gx < x_max and not occ[gy][gx]) else 0
        stack: list[int] = []
        for gx in range(gw + 1):
            hcur = heights[gx] if gx < gw else 0
            while stack and heights[stack[-1]] >= hcur:
                hh = heights[stack.pop()]
                left = stack[-1] + 1 if stack else 0
                width = gx - left
                if hh * width > best[0]:
                    best = (hh * width, left, gy - hh + 1, width, hh)
            stack.append(gx)
    if best[0] == 0:
        return default, True
    _, bx, by, bw, bh = best
    box = Box(x=int(bx / gw * slide_w), y=int(by / gh * slide_h), w=int(bw / gw * slide_w), h=int(bh / gh * slide_h))
    if box.area < 0.18 * default.area:  # свободны лишь узкие полосы: считаем фон текстурой
        return default, True
    return box, False


def region_color(png: Path, box: Box, slide_w: int, slide_h: int) -> str | None:
    """Преобладающий цвет рендера внутри области (то, на чём окажется поставленный туда текст)."""
    im = Image.open(png).convert("RGB")
    sx, sy = im.width / slide_w, im.height / slide_h
    crop = im.crop((int(box.x * sx), int(box.y * sy), max(int(box.r * sx), int(box.x * sx) + 2), max(int(box.b * sy), int(box.y * sy) + 2)))
    q = crop.resize((48, 27)).quantize(colors=4)
    counts = sorted(q.getcolors() or [], reverse=True)
    if not counts:
        return None
    pal = q.getpalette()
    i = counts[0][1]
    return rgb_to_hex((pal[i * 3], pal[i * 3 + 1], pal[i * 3 + 2]))


def layout_photo_share(layout, slide_w: int, slide_h: int) -> float:
    """Доля слайда под непрозрачными растровыми картинками, которые рисует сам макет (не плейсхолдеры):
    фотоколлаж, вшитый в макет, виден на каждом слайде на его основе, о чём бы слайд ни был. Прозрачная
    PNG-графика (логотипы, 3D-фигуры) не считается.
    """
    from io import BytesIO

    total = 0
    stack = list(layout.shapes)
    while stack:
        shp = stack.pop()
        if shp.shape_type == 6:  # группа
            stack.extend(shp.shapes)
            continue
        if getattr(shp, "is_placeholder", False) or shp.width is None or shp.height is None:
            continue
        blip = shp._element.find(".//" + qn("a:blip"))
        rid = blip.get(qn("r:embed")) if blip is not None else None
        if not rid:
            continue
        try:
            im = Image.open(BytesIO(layout.part.related_part(rid).blob))
            if im.mode in ("RGBA", "LA") or (im.mode == "P" and "transparency" in im.info):
                if im.convert("RGBA").getchannel("A").resize((32, 32)).getextrema()[0] < 200:
                    continue  # вырезанная графика, а не фото
        except Exception:
            continue
        x0, y0 = max(int(shp.left), 0), max(int(shp.top), 0)
        x1, y1 = min(int(shp.left + shp.width), slide_w), min(int(shp.top + shp.height), slide_h)
        area = max(x1 - x0, 0) * max(y1 - y0, 0)
        if area > 0.85 * slide_w * slide_h:
            continue  # фоновая картинка во весь слайд, а не изображение на слайде
        total += area
    return round(min(total / (slide_w * slide_h), 1.0), 3)


def region_colors(png: Path, box: Box, slide_w: int, slide_h: int, min_share: float = 0.1) -> list[str]:
    """Цвета рендера, занимающие не меньше `min_share` области (градиент даёт несколько): текст, поставленный
    туда, должен читаться на каждом из них.
    """
    im = Image.open(png).convert("RGB")
    sx, sy = im.width / slide_w, im.height / slide_h
    crop = im.crop((int(box.x * sx), int(box.y * sy), max(int(box.r * sx), int(box.x * sx) + 2), max(int(box.b * sy), int(box.y * sy) + 2)))
    q = crop.resize((64, 36)).quantize(colors=6)
    pal = q.getpalette()
    total = 64 * 36
    return [rgb_to_hex((pal[i * 3], pal[i * 3 + 1], pal[i * 3 + 2])) for n, i in sorted(q.getcolors() or [], reverse=True)
            if n >= min_share * total]


def rendered_palette(pngs: list[Path]) -> Counter:
    """Цвета, которые шаблон действительно показывает (фоны, графика, полосы), с весом по площади."""
    acc: Counter = Counter()
    for p in pngs:
        try:
            q = Image.open(p).convert("RGB").resize((96, 54)).quantize(colors=8)
        except Exception:
            continue
        pal = q.getpalette()
        total = 96 * 54
        for n, i in q.getcolors() or []:
            acc[rgb_to_hex((pal[i * 3], pal[i * 3 + 1], pal[i * 3 + 2]))] += n / total
    return acc


def mark_repeated_brand_texts(slide_elements: list[list[Element]], sw: int, sh: int) -> None:
    """Текст, повторяющийся в одном и том же месте у края на большинстве слайдов-примеров («Конфиденциально»,
    «© Компания», название продукта в углу), — бренд-элемент, а не слот контента.
    """
    from decksmith.parsing.elements import is_filler_text

    n = len(slide_elements)
    if n < 3:
        return

    def key(e: Element):
        """Ключ повторяющегося текста: нормализованный текст и позиция на сетке."""
        return (re.sub(r"\s+", " ", e.text.strip().lower()), round(e.box.x / (0.02 * sw)), round(e.box.y / (0.02 * sh)))

    counts: Counter = Counter()
    for els in slide_elements:
        counts.update({key(e) for e in els if e.kind == "text" and e.text.strip()})
    need = max(3, -(-n // 2))
    for els in slide_elements:
        for e in els:
            if e.kind != "text" or not e.text.strip() or e.is_title_ph or is_filler_text(e.text):
                continue
            near_edge = e.box.y > 0.85 * sh or e.box.b < 0.15 * sh
            if near_edge and len(e.text) <= 80 and counts[key(e)] >= need:
                e.brand = True


def normalize_input(path: Path, workdir: Path) -> Path:
    """Принимает .pptx, .potx/.pptm/.potm (замена content-type) и всё, что открывает LibreOffice (.ppt, .odp,
    .otp, .key...), с конвертацией в .pptx.
    """
    suffix = path.suffix.lower()
    out = workdir / "template.pptx"
    if suffix == ".pptx":
        shutil.copy(path, out)
        return out
    if suffix in (".potx", ".pptm", ".potm"):
        with zipfile.ZipFile(path) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    data = re.sub(rb"application/vnd\.ms-powerpoint\.(template|presentation)\.macroEnabled\.main\+xml|"
                                  rb"application/vnd\.openxmlformats-officedocument\.presentationml\.template\.main\+xml",
                                  b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml", data)
                if item.filename.lower().endswith("vbaproject.bin"):
                    continue  # макросы никогда не выполняются и не сохраняются
                zout.writestr(item, data)
        return out
    tmp = Path(tempfile.mkdtemp(prefix="decksmith_convert_"))
    src = tmp / f"in{suffix}"
    shutil.copy(path, src)
    subprocess.run([soffice_path(), f"-env:UserInstallation=file://{tmp}/profile", "--headless", "--convert-to",
                    "pptx:Impress MS PowerPoint 2007 XML", "--outdir", str(tmp), str(src)], capture_output=True, timeout=300)
    conv = tmp / "in.pptx"
    if not conv.exists():
        raise RenderError(f"cannot convert {path.name} to .pptx")
    shutil.move(str(conv), out)
    shutil.rmtree(tmp, ignore_errors=True)
    return out


# ----------------------------------------------------------------------------
# Этап рендера
# ----------------------------------------------------------------------------
def _render_template(path: Path, out: Path) -> list[Path]:
    """Рендер слайдов шаблона в PNG (миниатюры паттернов)."""
    pdf = pptx_to_pdf(path, out)
    pngs = pdf_to_pngs(pdf, out / "slides", dpi=48, prefix="slide")
    pdf.unlink(missing_ok=True)
    return pngs


def _render_layout_probe(path: Path, out: Path) -> list[Path]:
    """По одному пустому слайду на макет (плейсхолдеры удалены) -> PNG на каждый макет."""
    prs = Presentation(str(path))
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        prs.part.drop_rel(sldId.rId)
        sldIdLst.remove(sldId)
    from pptx.oxml.ns import qn

    for container in list(prs.slide_masters) + [lay for _, _, lay in all_layouts(prs)]:
        for ph in container.placeholders:  # текст-подсказка выглядел бы занятой областью
            if "FOOTER" in str(ph.placeholder_format.type) or "SLIDE_NUMBER" in str(ph.placeholder_format.type):
                continue
            for t in ph._element.iter(qn("a:t")):
                t.text = ""
    for _, _, layout in all_layouts(prs):
        s = prs.slides.add_slide(layout)
        for ph in list(s.placeholders):
            ph._element.getparent().remove(ph._element)
    tmp = Path(tempfile.mkdtemp(prefix="decksmith_probe_"))
    probe = tmp / "layouts.pptx"
    prs.save(str(probe))
    pdf = pptx_to_pdf(probe, tmp)
    pngs = pdf_to_pngs(pdf, out / "layouts", dpi=48, prefix="layout")
    shutil.rmtree(tmp, ignore_errors=True)
    return pngs


# ----------------------------------------------------------------------------
# Бренд-элементы и иконки
# ----------------------------------------------------------------------------
def _brand_elements(container, theme: Theme, sw: int, sh: int, source: str) -> list[BrandElement]:
    out = []
    for e in extract_elements(container, theme, sw, sh, include_empty_placeholders=True):
        # логотипы и надписи-знаки: небольшие объекты у края (крупная графика у края — оформление, а не
        # логотип)
        near_edge = e.box.y > 0.85 * sh or e.box.b < 0.15 * sh or e.box.x > 0.85 * sw or e.box.r < 0.15 * sw
        small = e.box.area < 0.025 * sw * sh and 0.1 < (e.box.w / max(e.box.h, 1)) < 12
        if e.placeholder == "SLIDE_NUMBER":
            out.append(BrandElement(kind="page_number", box=e.box, source=source, name=e.name))
        elif e.placeholder == "FOOTER":
            out.append(BrandElement(kind="footer", box=e.box, source=source, name=e.name))
        elif e.placeholder == "DATE":
            out.append(BrandElement(kind="date", box=e.box, source=source, name=e.name))
        elif e.placeholder is None and e.kind in ("picture", "group", "text") and small and near_edge and e.box.w > 0:
            out.append(BrandElement(kind="logo", box=e.box, source=source, name=e.name))
    return out


def _harvest_icons(prs, patterns: list[Pattern], out: Path, sw: int, sh: int) -> list[IconAsset]:
    """Иконки со слайдов-библиотек шаблона (для карточек и списков)."""
    icons: list[IconAsset] = []
    seen: set[str] = set()
    out.mkdir(parents=True, exist_ok=True)
    guide_slides = {p.slide_index for p in patterns if p.kind == PatternKind.guide and "icons" in p.tags}
    for si, slide in enumerate(prs.slides):
        if si not in guide_slides:
            continue
        for sh_ in slide.shapes:
            try:
                img = sh_.image
            except Exception:
                continue
            if img.sha1 in seen:
                continue
            seen.add(img.sha1)
            try:
                pil = Image.open(io.BytesIO(img.blob))
                if max(pil.size) > 600:
                    continue
                p = out / f"icon_{len(icons):03d}.png"
                pil.convert("RGBA").save(p)
            except Exception:
                continue
            icons.append(IconAsset(id=f"icon_{len(icons):03d}", path=str(p), source_slide=si))
    return icons


# ----------------------------------------------------------------------------
# Главная точка входа
# ----------------------------------------------------------------------------
def analyze_template(path: str | Path, *, name: str | None = None, force: bool = False, render: bool = True) -> TemplateProfile:
    t0 = time.time()
    path = Path(path)
    sha = file_sha256(path)
    tid = sha[:12]
    workdir = settings().workspace / "templates" / tid
    prof_path = workdir / "profile.json"
    if prof_path.exists() and not force:
        prof = TemplateProfile.model_validate_json(prof_path.read_text(encoding="utf-8"))
        if prof.parser_version == PARSER_VERSION:
            return prof
    workdir.mkdir(parents=True, exist_ok=True)
    stored = workdir / "template.pptx"
    if not stored.exists() or force:
        stored = normalize_input(path, workdir)
    warnings: list[str] = []

    prs = Presentation(str(stored))
    sw, sh = int(prs.slide_width), int(prs.slide_height)
    themes = [parse_theme(m) for m in prs.slide_masters]
    layouts_raw = all_layouts(prs)

    # --- 2. рендер ---------------------------------------------------------------
    slide_pngs: list[Path] = []
    layout_pngs: list[Path] = []
    if render:
        try:
            slide_pngs = _render_template(stored, workdir)
            layout_pngs = _render_layout_probe(stored, workdir)
        except RenderError as e:
            warnings.append(f"render failed: {e}")
    slide_cv = [image_stats(p) for p in slide_pngs]
    layout_cv = [image_stats(p) for p in layout_pngs]

    # --- 1. макеты ----------------------------------------------------------------
    layouts: list[LayoutInfo] = []
    for gi, mi, layout in layouts_raw:
        phs = []
        for ph in layout.placeholders:
            pf = ph.placeholder_format
            phs.append({
                "idx": pf.idx,
                "type": str(pf.type).split(".")[-1].split(" ")[0],
                "name": ph.name,
                "box": Box(x=int(ph.left or 0), y=int(ph.top or 0), w=int(ph.width or 0), h=int(ph.height or 0)).model_dump(),
                "painted": paints(ph._element),
            })
        types = [p["type"] for p in phs]
        title_ph = next((p for p in phs if p["type"] in ("TITLE", "CENTER_TITLE")), None)
        li = LayoutInfo(
            index=gi,
            master_index=mi,
            name=layout.name,
            placeholders=phs,
            has_title=title_ph is not None,
            body_count=sum(t in ("BODY", "OBJECT") for t in types),
            picture_count=sum(t == "PICTURE" for t in types),
            title_box=Box(**title_ph["box"]) if title_ph else None,
            painted_idx=[p["idx"] for p in phs if p["painted"] and p["type"] not in
                         ("TITLE", "CENTER_TITLE", "FOOTER", "DATE", "SLIDE_NUMBER")],
            photo_share=layout_photo_share(layout, sw, sh),
        )
        if gi < len(layout_cv):
            li.dark, li.background_hex, _ = layout_cv[gi]
            li.thumbnail = str(layout_pngs[gi])
        layouts.append(li)

    # --- 4. паттерны ------------------------------------------------------------------
    slide_elements: list[list[Element]] = []
    patterns: list[Pattern] = []
    slide_dark, slide_bg = [], []
    slide_layout_idx = []
    for si, slide in enumerate(prs.slides):
        gi = layout_global_index(prs, slide.slide_layout)
        mi = layouts[gi].master_index if gi >= 0 else 0
        if gi >= 0:
            layouts[gi].used_by_slides.append(si)
        slide_layout_idx.append(gi)
        slide_elements.append(extract_elements(slide, themes[mi], sw, sh))
    mark_repeated_brand_texts(slide_elements, sw, sh)
    furniture: dict[tuple, dict] = {}
    for si, els in enumerate(slide_elements):
        for e in els:
            if e.brand and e.kind == "text" and e.placeholder is None and not e.group_path:
                key = (e.text.strip(), round(e.box.x / 50000), round(e.box.y / 50000))
                furniture.setdefault(key, {"slide": si, "shape_id": e.shape_id, "text": e.text.strip()[:60],
                                           "master": layouts[slide_layout_idx[si]].master_index if slide_layout_idx[si] >= 0 else 0})
    for si, slide in enumerate(prs.slides):
        gi = slide_layout_idx[si]
        els = slide_elements[si]
        dark, bg, _ = slide_cv[si] if si < len(slide_cv) else (False, themes[0].scheme("bg1") or "FFFFFF", 1.0)
        slide_dark.append(dark)
        slide_bg.append(bg)
        pat = build_pattern(
            f"s{si + 1:02d}", els, slide_index=si, layout_index=gi, layout_name=slide.slide_layout.name,
            slide_w=sw, slide_h=sh, bottom_margin=int(0.06 * sh), dark=dark,
        )
        if si < len(slide_pngs):
            pat.thumbnail = str(slide_pngs[si])
        patterns.append(pat)

    # макеты с большим числом плейсхолдеров тоже становятся паттернами (шаблоны с малым числом примеров)
    for li in layouts:
        if li.body_count + li.picture_count == 0:
            continue
        layout = layouts_raw[li.index][2]
        els = [e for e in extract_elements(layout, themes[li.master_index], sw, sh) if e.placeholder is not None]
        pat = build_pattern(
            f"l{li.index:02d}", els, slide_index=None, layout_index=li.index, layout_name=li.name,
            slide_w=sw, slide_h=sh, bottom_margin=int(0.06 * sh), dark=li.dark, source="layout",
        )
        pat.thumbnail = li.thumbnail
        pat.score_hint *= 0.85  # примеры богаче голых макетов
        patterns.append(pat)

    # --- шрифты -----------------------------------------------------------------------
    fams = Counter(e.style.font for els in slide_elements for e in els if e.kind == "text" and e.style and e.style.font)
    for li in layouts:  # шаблоны с малым числом примеров: шрифты плейсхолдеров макетов тоже учитываются
        for e in extract_elements(layouts_raw[li.index][2], themes[li.master_index], sw, sh):
            if e.kind == "text" and e.style and e.style.font:
                fams[e.style.font] += 1
    deadline = time.monotonic() + DOWNLOAD_BUDGET_S
    font_avail = {f: ensure_font(f, deadline=deadline) for f, _ in fams.most_common(6)}

    # --- 5. токены --------------------------------------------------------------------
    render_colors = rendered_palette(slide_pngs or layout_pngs)
    tokens = build_tokens(
        slide_w=sw, slide_h=sh, theme=themes[0], slide_elements=slide_elements, patterns=patterns,
        slide_dark=slide_dark, slide_bg=slide_bg or [c[1] for c in layout_cv], font_available=font_avail,
        render_colors=render_colors,
    )

    # --- 3. свободная зона контента по макетам + выбор холста ---------------------------------
    for li in layouts:
        if li.thumbnail:
            below = (li.title_box.b if li.title_box else (tokens.title_box.b if tokens.title_box else int(0.2 * sh)))
            li.content_box, li.textured = free_content_box(Path(li.thumbnail), sw, sh, below, tokens.margins,
                                                           li.background_hex or tokens.background_hex)
            if li.content_box:
                li.content_bg_hex = region_color(Path(li.thumbnail), li.content_box, sw, sh)
            tb_ = li.title_box or tokens.title_box
            if tb_ is not None:
                li.title_clear = title_clear_box(Path(li.thumbnail), tb_, sw, sh, li.background_hex or tokens.background_hex)
            # окрашенный плейсхолдер учитывается, только если его видно на пробном рендере (на пробном слайде
            # нет
            # плейсхолдеров, так что всё нарисованное в его рамке идёт от самого макета)
            bg_ = li.background_hex or tokens.background_hex
            boxes = {ph["idx"]: Box(**ph["box"]) for ph in li.placeholders}
            li.painted_idx = [i for i in li.painted_idx if i in boxes and boxes[i].area > 0
                              and color_distance(region_color(Path(li.thumbnail), boxes[i], sw, sh) or bg_, bg_) > 20]
    for pat in patterns:  # слоты на заметно окрашенных плейсхолдерах макета (см. Slot.painted)
        if pat.layout_index is not None and 0 <= pat.layout_index < len(layouts):
            painted = set(layouts[pat.layout_index].painted_idx)
            for sl in pat.slots:
                sl.painted = sl.placeholder_idx in painted
    content_kinds = {PatternKind.cards, PatternKind.text, PatternKind.steps, PatternKind.stats, PatternKind.table,
                     PatternKind.chart, PatternKind.image_text, PatternKind.two_column}

    def canvas_score(li: LayoutInfo) -> float:
        uses = sum(1 for p in patterns if p.layout_index == li.index and p.kind in content_kinds)
        area = li.content_box.area / (sw * sh) if li.content_box else 0
        simple = 1.0 if li.body_count + li.picture_count == 0 else 0.7
        titled = 1.0 if li.has_title else 0.55  # плейсхолдер заголовка сохраняет стиль заголовка шаблона
        calm = 0.6 if li.textured else 1.0  # предпочтение — макетам с действительно пустой зоной контента
        clean = 0.3 if li.painted_idx else 1.0  # окрашенные плейсхолдеры выглядели бы пустыми карточками
        return (uses + 1) * area * simple * titled * calm * clean

    canvas: dict[str, int] = {}
    for tone in ("light", "dark"):
        cands = [li for li in layouts if li.content_box and li.dark == (tone == "dark") and li.content_box.area > 0.2 * sw * sh]
        if cands:
            canvas[tone] = max(cands, key=canvas_score).index
    if not canvas:
        cands = [li for li in layouts if li.content_box] or layouts
        best = max(cands, key=canvas_score) if cands else None
        if best is not None:
            canvas["dark" if best.dark else "light"] = best.index
        warnings.append("no clean canvas layout found; composed slides use the largest free area available")

    # --- 6. бренд -----------------------------------------------------------------------
    brand: list[BrandElement] = []
    for mi, m in enumerate(prs.slide_masters):
        brand += _brand_elements(m, themes[mi], sw, sh, f"master:{mi}")
    for li in layouts:
        brand += _brand_elements(layouts_raw[li.index][2], themes[li.master_index], sw, sh, f"layout:{li.index}")

    # убираем дубли бренд-элементов, повторяющихся на многих макетах (тот же вид + та же рамка)
    uniq: dict[tuple, BrandElement] = {}
    for b in brand:
        key = (b.kind, round(b.box.x / 50000), round(b.box.y / 50000), round(b.box.w / 50000), round(b.box.h / 50000))
        if key in uniq:
            if b.source not in uniq[key].source:
                uniq[key].source += "," + b.source
        else:
            uniq[key] = b
    brand = list(uniq.values())

    # --- 7. ресурсы ----------------------------------------------------------------------
    icons = _harvest_icons(prs, patterns, workdir / "icons", sw, sh)

    for f in tokens.fonts:
        if not f.available:
            warnings.append(f"font '{f.family}' not available locally: rendering will substitute it")

    prof = TemplateProfile(
        id=tid,
        name=name or path.stem,
        file=str(stored),
        sha256=sha,
        tokens=tokens,
        layouts=layouts,
        patterns=patterns,
        brand_elements=brand,
        icons=icons,
        canvas_layouts=canvas,
        brand_furniture=list(furniture.values()),
        slide_thumbnails=[str(p) for p in slide_pngs],
        workdir=str(workdir),
        parser_version=PARSER_VERSION,
        warnings=warnings,
    )
    prof_path.write_text(prof.model_dump_json(indent=1), encoding="utf-8")
    log.info("template %s analysed in %.1fs: %d patterns", tid, time.time() - t0, len(patterns))
    return prof


def summarize(prof: TemplateProfile) -> dict:
    """Краткая сводка профиля для UI и CLI."""
    kinds = Counter(p.kind.value for p in prof.patterns)
    t = prof.tokens
    return {
        "id": prof.id,
        "name": prof.name,
        "slide_size_in": [round(t.slide_w / 914400, 2), round(t.slide_h / 914400, 2)],
        "patterns": dict(kinds),
        "fonts": {"heading": t.heading_font, "body": t.body_font, "all": [(f.family, f.available) for f in t.fonts]},
        "type_scale": t.type_scale.model_dump(),
        "accent": t.accent_hex,
        "chart_colors": t.chart_colors,
        "background": t.background_hex,
        "dark": t.dark_background,
        "text": t.text_hex,
        "palette": [c.hex for c in t.palette],
        "margins_in": {k: round(v / 914400, 2) for k, v in t.margins.model_dump().items()},
        "canvas_layouts": {k: prof.layouts[v].name for k, v in prof.canvas_layouts.items()},
        "brand_elements": len(prof.brand_elements),
        "icons": len(prof.icons),
        "warnings": prof.warnings,
    }


if __name__ == "__main__":  # pragma: no cover
    import sys

    from decksmith.core.logging import setup_logging

    setup_logging()
    p = analyze_template(sys.argv[1], force="--force" in sys.argv)
    print(json.dumps(summarize(p), ensure_ascii=False, indent=1))
