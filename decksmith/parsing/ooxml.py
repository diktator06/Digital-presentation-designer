"""Низкоуровневые помощники OOXML: разворачивание фигур, наследование стилей, цвета темы.

python-pptx отдаёт только то, что явно задано у элемента. Но дизайн-система
живёт в цепочке наследования:

    rPr фрагмента -> pPr/defRPr абзаца -> lstStyle фигуры -> плейсхолдер макета
            -> плейсхолдер мастера -> txStyles мастера -> шрифты/цвета темы

Модуль детерминированно разрешает *эффективные* значения, чтобы разборщик
(а затем аудит) видел то, что нарисует рендерер.
"""
from __future__ import annotations

import colorsys
from dataclasses import dataclass, field
from typing import Iterator

from lxml import etree
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn

from decksmith.core.models import Box

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
}


# ----------------------------------------------------------------------------
# Цвета
# ----------------------------------------------------------------------------
def hex_to_rgb(h: str) -> tuple[int, int, int]:
    """hex -> (r, g, b)."""
    h = h.lstrip("#")
    return int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)


def rgb_to_hex(rgb: tuple[float, float, float]) -> str:
    """(r, g, b) -> hex."""
    return "".join(f"{max(0, min(255, round(c))):02X}" for c in rgb)


def rel_luminance(h: str) -> float:
    """Относительная яркость цвета по WCAG."""
    def ch(c: int) -> float:
        """Линеаризация канала sRGB."""
        c = c / 255
        return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4

    r, g, b = hex_to_rgb(h)
    return 0.2126 * ch(r) + 0.7152 * ch(g) + 0.0722 * ch(b)


def contrast_ratio(a: str, b: str) -> float:
    """Контраст двух цветов по WCAG (1..21)."""
    la, lb = rel_luminance(a), rel_luminance(b)
    hi, lo = max(la, lb), min(la, lb)
    return (hi + 0.05) / (lo + 0.05)


def color_distance(a: str, b: str) -> float:
    """Почти перцептивное расстояние (приближение redmean), 0..~765."""
    r1, g1, b1 = hex_to_rgb(a)
    r2, g2, b2 = hex_to_rgb(b)
    rm = (r1 + r2) / 2
    dr, dg, db = r1 - r2, g1 - g2, b1 - b2
    return ((2 + rm / 256) * dr * dr + 4 * dg * dg + (2 + (255 - rm) / 256) * db * db) ** 0.5


def _apply_mods(hex_: str, el) -> str:
    """Применяет модификаторы цвета DrawingML (lumMod/lumOff/tint/shade; alpha игнорируется)."""
    r, g, b = (c / 255 for c in hex_to_rgb(hex_))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    for m in el:
        tag = etree.QName(m).localname
        val = m.get("val")
        if val is None:
            continue
        v = int(val) / 100000
        if tag == "lumMod":
            l = l * v
        elif tag == "lumOff":
            l = l + v
        elif tag == "tint":
            r, g, b = colorsys.hls_to_rgb(h, l, s)
            r, g, b = (c + (1 - c) * (1 - v) for c in (r, g, b))
            h, l, s = colorsys.rgb_to_hls(r, g, b)
        elif tag == "shade":
            r, g, b = colorsys.hls_to_rgb(h, l, s)
            r, g, b = (c * v for c in (r, g, b))
            h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = max(0.0, min(1.0, l))
    r, g, b = colorsys.hls_to_rgb(h, l, s)
    return rgb_to_hex((r * 255, g * 255, b * 255))


@dataclass
class Theme:
    colors: dict[str, str] = field(default_factory=dict)  # dk1, lt1, accent1..6, hlink
    clr_map: dict[str, str] = field(default_factory=dict)  # bg1->lt1, tx1->dk1 ...
    major_font: str = "Arial"
    minor_font: str = "Arial"

    def scheme(self, name: str) -> str | None:
        """Цвет схемы темы с учётом clrMap."""
        name = self.clr_map.get(name, name)
        return self.colors.get(name)


def parse_theme(master) -> Theme:
    """Тема, привязанная к мастеру слайдов (+ его clrMap)."""
    th = Theme()
    try:
        theme_part = master.part.part_related_by(
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme"
        )
        root = etree.fromstring(theme_part.blob)
    except Exception:  # pragma: no cover — повреждённый шаблон
        return th
    cs = root.find(".//a:clrScheme", NS)
    if cs is not None:
        for c in cs:
            name = etree.QName(c).localname
            if len(c):
                inner = c[0]
                val = inner.get("val") if etree.QName(inner).localname == "srgbClr" else inner.get("lastClr")
                if val and len(val) == 6:
                    th.colors[name] = val.upper()
    fs = root.find(".//a:fontScheme", NS)
    if fs is not None:
        maj = fs.find("a:majorFont/a:latin", NS)
        mnr = fs.find("a:minorFont/a:latin", NS)
        th.major_font = maj.get("typeface") if maj is not None else "Arial"
        th.minor_font = mnr.get("typeface") if mnr is not None else "Arial"
    cm = master._element.find(qn("p:clrMap"))
    if cm is not None:
        th.clr_map = dict(cm.attrib)
    else:
        th.clr_map = {"bg1": "lt1", "tx1": "dk1", "bg2": "lt2", "tx2": "dk2"}
    return th


def color_from_fill(el, theme: Theme) -> str | None:
    """Разрешает элемент, содержащий srgbClr/schemeClr/sysClr/prstClr."""
    if el is None:
        return None
    for child in el.iter():
        tag = etree.QName(child).localname
        if tag == "srgbClr":
            return _apply_mods(child.get("val", "000000").upper(), child)
        if tag == "schemeClr":
            base = theme.scheme(child.get("val", "tx1"))
            return _apply_mods(base, child) if base else None
        if tag == "sysClr":
            return _apply_mods((child.get("lastClr") or "000000").upper(), child)
        if tag == "prstClr":
            return {"black": "000000", "white": "FFFFFF"}.get(child.get("val"), None)
    return None


def resolve_font_name(name: str | None, theme: Theme) -> str | None:
    """Реальное имя шрифта: ссылки +mj/+mn заменяются шрифтами темы."""
    if not name:
        return None
    if name.startswith("+mj"):
        return theme.major_font
    if name.startswith("+mn"):
        return theme.minor_font
    return name


# ----------------------------------------------------------------------------
# Разворачивание фигур
# ----------------------------------------------------------------------------
@dataclass
class ShapeRec:
    shape: object
    path: list[int]
    box: Box
    depth: int
    group_path: list[int]

    @property
    def shape_id(self) -> int:
        """id фигуры."""
        return self.shape.shape_id


def _xfrm_box(shape) -> Box | None:
    """Рамка фигуры в её собственных координатах (None, если размер не задан)."""
    try:
        x, y, w, h = shape.left, shape.top, shape.width, shape.height
    except Exception:
        return None
    if x is None or y is None or w is None or h is None:
        return None
    return Box(x=int(x), y=int(y), w=int(w), h=int(h))


def _group_transform(group):
    """Преобразование координат группы: смещение и масштаб дочернего пространства."""
    xfrm = group._element.find(qn("p:grpSpPr")).find(qn("a:xfrm"))
    if xfrm is None:
        return (0, 0, 1.0, 1.0, 0, 0)
    off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
    choff, chext = xfrm.find(qn("a:chOff")), xfrm.find(qn("a:chExt"))
    ox, oy = int(off.get("x")), int(off.get("y"))
    ex, ey = int(ext.get("cx")), int(ext.get("cy"))
    cx0 = int(choff.get("x")) if choff is not None else ox
    cy0 = int(choff.get("y")) if choff is not None else oy
    cex = int(chext.get("cx")) if chext is not None else ex
    cey = int(chext.get("cy")) if chext is not None else ey
    sx = ex / cex if cex else 1.0
    sy = ey / cey if cey else 1.0
    return (ox, oy, sx, sy, cx0, cy0)


def flatten_shapes(shapes, _path=None, _tf=None, _depth=0, _gpath=None) -> Iterator[ShapeRec]:
    """Отдаёт каждую листовую фигуру (и группу) с её абсолютной рамкой на слайде."""
    _path = _path or []
    _gpath = _gpath or []
    for i, sh in enumerate(shapes):
        box = _xfrm_box(sh)
        if box is not None and _tf is not None:
            ox, oy, sx, sy, cx0, cy0 = _tf
            box = Box(
                x=int(ox + (box.x - cx0) * sx),
                y=int(oy + (box.y - cy0) * sy),
                w=int(box.w * sx),
                h=int(box.h * sy),
            )
        if box is None:
            box = Box(x=0, y=0, w=0, h=0)
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield ShapeRec(sh, _path + [i], box, _depth, _gpath)
            local = _group_transform(sh)
            if _tf is not None:
                # композиция: родитель отображает дочернее пространство этой группы
                ox, oy, sx, sy, cx0, cy0 = _tf
                gox, goy, gsx, gsy, gcx0, gcy0 = local
                local = (
                    ox + (gox - cx0) * sx,
                    oy + (goy - cy0) * sy,
                    gsx * sx,
                    gsy * sy,
                    gcx0,
                    gcy0,
                )
            yield from flatten_shapes(sh.shapes, _path + [i], local, _depth + 1, _gpath + [sh.shape_id])
        else:
            yield ShapeRec(sh, _path + [i], box, _depth, _gpath)


# ----------------------------------------------------------------------------
# Эффективный стиль текста
# ----------------------------------------------------------------------------
_TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}


def placeholder_type(shape):
    """Тип плейсхолдера фигуры или None."""
    try:
        return shape.placeholder_format.type if shape.is_placeholder else None
    except Exception:
        return None


def _find_inherited_placeholder(shape, container):
    """Находит плейсхолдер, от которого наследуется плейсхолдер слайда/макета."""
    try:
        pf = shape.placeholder_format
    except Exception:
        return None
    for ph in container.placeholders:
        try:
            if ph.placeholder_format.idx == pf.idx and pf.idx != 0:
                return ph
        except Exception:
            continue
    t = pf.type
    for ph in container.placeholders:
        pt = ph.placeholder_format.type
        if pt == t or (t in _TITLE_TYPES and pt in _TITLE_TYPES):
            return ph
        if t in (PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.BODY) and pt in (PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.BODY):
            return ph
    return None


def _lvl_prop(txbody_el, lvl: int, attr: str, child: str | None = None):
    """Читает атрибут (или дочерний элемент) lstStyle/lvlNpPr/defRPr из txBody."""
    if txbody_el is None:
        return None
    lst = txbody_el.find(qn("a:lstStyle"))
    if lst is None:
        return None
    lp = lst.find(qn(f"a:lvl{lvl + 1}pPr"))
    if lp is None:
        return None
    d = lp.find(qn("a:defRPr"))
    if d is None:
        return None
    if child:
        return d.find(qn(child))
    return d.get(attr)


def _spacing_pt(el, size: float, line: bool = False) -> float | None:
    """Значение spcBef / spcAft / lnSpc: пункты или доля строки (для lnSpc — как множитель)."""
    if el is None:
        return None
    pts, pct = el.find(qn("a:spcPts")), el.find(qn("a:spcPct"))
    if pts is not None:
        v = int(pts.get("val", "0")) / 100
        return v / (size * 1.2) if line else v
    if pct is not None:
        v = int(pct.get("val", "0")) / 100000
        return v if line else v * size * 1.2
    return None


def _txstyle_ppr(master, kind: str, lvl: int):
    """pPr уровня lvl из текстового стиля мастера."""
    tx = master._element.find(qn("p:txStyles"))
    st = tx.find(qn(f"p:{kind}")) if tx is not None else None
    return st.find(qn(f"a:lvl{lvl + 1}pPr")) if st is not None else None


def _txstyle(master, kind: str, lvl: int):
    """defRPr уровня lvl из текстового стиля мастера."""
    tx = master._element.find(qn("p:txStyles"))
    if tx is None:
        return None
    st = tx.find(qn(f"p:{kind}"))
    if st is None:
        return None
    lp = st.find(qn(f"a:lvl{lvl + 1}pPr"))
    if lp is None:
        return None
    return lp.find(qn("a:defRPr"))


@dataclass
class EffStyle:
    size: float | None = None
    font: str | None = None
    color: str | None = None
    bold: bool = False
    space_before: float = 0.0  # пт, интервалы абзаца (разрешаются так же, как стиль фрагмента)
    space_after: float = 0.0
    line: float = 1.0  # множитель межстрочного интервала (1.0 = одинарный)


class StyleResolver:
    """Разрешает эффективные стили фрагментов для фигур одного слайда (или макета)."""

    def __init__(self, slide, theme: Theme):
        self.slide = slide
        self.theme = theme
        if hasattr(slide, "slide_layout"):  # обычный слайд
            self.layout = slide.slide_layout
            self.master = self.layout.slide_master
        elif hasattr(slide, "slide_master"):  # макет
            self.layout = None
            self.master = slide.slide_master
        else:  # мастер
            self.layout = None
            self.master = slide

    def _chain(self, shape):
        """Цепочка наследования: фигура -> плейсхолдер макета -> плейсхолдер мастера."""
        chain = [shape]
        if not getattr(shape, "is_placeholder", False):
            return chain
        cur = shape
        for container in [self.layout, self.master]:
            if container is None or container is self.slide:
                continue
            inh = _find_inherited_placeholder(cur, container)
            if inh is not None:
                chain.append(inh)
                cur = inh
        return chain

    def _master_style_kind(self, shape) -> str:
        """Какой текстовый стиль мастера наследует фигура: заголовка, тела или прочий."""
        pt = placeholder_type(shape)
        if pt in _TITLE_TYPES:
            return "titleStyle"
        if pt is not None:
            return "bodyStyle"
        return "otherStyle"

    def resolve(self, shape, paragraph=None, run=None) -> EffStyle:
        st = EffStyle()
        lvl = paragraph.level if paragraph is not None else 0
        # 1. явно у фрагмента / абзаца
        if run is not None:
            rpr = run._r.find(qn("a:rPr"))
            if rpr is not None:
                if rpr.get("sz"):
                    st.size = int(rpr.get("sz")) / 100
                if rpr.get("b") in ("1", "true"):
                    st.bold = True
                lat = rpr.find(qn("a:latin"))
                if lat is not None:
                    st.font = resolve_font_name(lat.get("typeface"), self.theme)
                fill = rpr.find(qn("a:solidFill"))
                if fill is not None:
                    st.color = color_from_fill(fill, self.theme)
        if paragraph is not None:
            ppr = paragraph._p.find(qn("a:pPr"))
            if ppr is not None:
                d = ppr.find(qn("a:defRPr"))
                if d is not None:
                    if st.size is None and d.get("sz"):
                        st.size = int(d.get("sz")) / 100
                    if st.font is None and d.find(qn("a:latin")) is not None:
                        st.font = resolve_font_name(d.find(qn("a:latin")).get("typeface"), self.theme)
                    if st.color is None and d.find(qn("a:solidFill")) is not None:
                        st.color = color_from_fill(d.find(qn("a:solidFill")), self.theme)
            if st.size is None:
                epr = paragraph._p.find(qn("a:endParaRPr"))
                if epr is not None and epr.get("sz") and run is None:
                    st.size = int(epr.get("sz")) / 100
        # 2. lstStyle фигуры, затем унаследованные плейсхолдеры
        for sh in self._chain(shape):
            txb = sh._element.find(qn("p:txBody"))
            if st.size is None:
                v = _lvl_prop(txb, lvl, "sz")
                if v:
                    st.size = int(v) / 100
            if st.font is None:
                lat = _lvl_prop(txb, lvl, "", "a:latin")
                if lat is not None:
                    st.font = resolve_font_name(lat.get("typeface"), self.theme)
            if st.color is None:
                f = _lvl_prop(txb, lvl, "", "a:solidFill")
                if f is not None:
                    st.color = color_from_fill(f, self.theme)
            if sh is not shape and txb is not None and (st.size is None or st.font is None):
                # плейсхолдеры макетов часто несут стиль на своих образцовых фрагментах
                for p in txb.findall(qn("a:p")):
                    for r in p.findall(qn("a:r")):
                        rpr = r.find(qn("a:rPr"))
                        if rpr is None:
                            continue
                        if st.size is None and rpr.get("sz"):
                            st.size = int(rpr.get("sz")) / 100
                        lat = rpr.find(qn("a:latin"))
                        if st.font is None and lat is not None:
                            st.font = resolve_font_name(lat.get("typeface"), self.theme)
                        if st.color is None and rpr.find(qn("a:solidFill")) is not None:
                            st.color = color_from_fill(rpr.find(qn("a:solidFill")), self.theme)
                        break
                    break
        # 3. текстовые стили мастера
        d = _txstyle(self.master, self._master_style_kind(shape), lvl)
        if d is not None:
            if st.size is None and d.get("sz"):
                st.size = int(d.get("sz")) / 100
            if st.font is None and d.find(qn("a:latin")) is not None:
                st.font = resolve_font_name(d.find(qn("a:latin")).get("typeface"), self.theme)
            if st.color is None and d.find(qn("a:solidFill")) is not None:
                st.color = color_from_fill(d.find(qn("a:solidFill")), self.theme)
        # 4. значения по умолчанию
        if st.size is None:
            st.size = 18.0
        if st.font is None:
            st.font = self.theme.major_font if self._master_style_kind(shape) == "titleStyle" else self.theme.minor_font
        if st.color is None:
            st.color = self.theme.scheme("tx1") or "000000"
        # масштаб автоподбора
        txb = shape._element.find(qn("p:txBody"))
        if txb is not None:
            bp = txb.find(qn("a:bodyPr"))
            if bp is not None:
                na = bp.find(qn("a:normAutofit"))
                if na is not None and na.get("fontScale"):
                    st.size = st.size * int(na.get("fontScale")) / 100000
        st.size = round(st.size, 2)
        if paragraph is not None:
            self._spacing(shape, paragraph, lvl, st)
        return st

    def _spacing(self, shape, paragraph, lvl: int, st: EffStyle) -> None:
        """Интервалы до/после абзаца и межстрочный интервал: абзац, фигура и унаследованные списки стилей
        плейсхолдеров, затем текстовый стиль мастера (так, как их разрешают рендереры).
        """
        sources = [paragraph._p.find(qn("a:pPr"))]
        for sh in self._chain(shape):
            txb = sh._element.find(qn("p:txBody"))
            lst = txb.find(qn("a:lstStyle")) if txb is not None else None
            sources.append(lst.find(qn(f"a:lvl{lvl + 1}pPr")) if lst is not None else None)
        sources.append(_txstyle_ppr(self.master, self._master_style_kind(shape), lvl))
        found = {}
        for src in sources:
            for tag in ("spcBef", "spcAft", "lnSpc"):
                if tag not in found and src is not None and src.find(qn(f"a:{tag}")) is not None:
                    found[tag] = src.find(qn(f"a:{tag}"))
        size = st.size or 18.0
        st.space_before = _spacing_pt(found.get("spcBef"), size) or 0.0
        st.space_after = _spacing_pt(found.get("spcAft"), size) or 0.0
        st.line = _spacing_pt(found.get("lnSpc"), size, line=True) or 1.0
        txb = shape._element.find(qn("p:txBody"))
        na = txb.find(qn("a:bodyPr") + "/" + qn("a:normAutofit")) if txb is not None else None
        if na is not None and na.get("lnSpcReduction"):
            st.line *= 1 - int(na.get("lnSpcReduction")) / 100000

    def dominant(self, shape) -> EffStyle | None:
        """Стиль первого непустого фрагмента (представительный для фигуры)."""
        if not getattr(shape, "has_text_frame", False) or not shape.has_text_frame:
            return None
        for p in shape.text_frame.paragraphs:
            for r in p.runs:
                if r.text.strip():
                    return self.resolve(shape, p, r)
        if shape.text_frame.paragraphs:
            return self.resolve(shape, shape.text_frame.paragraphs[0], None)
        return None


# ----------------------------------------------------------------------------
# Заливки / фоны
# ----------------------------------------------------------------------------
def shape_fill_hex(shape, theme: Theme) -> str | None:
    """Сплошная заливка автофигуры (явная или через fillRef стиля)."""
    el = shape._element
    sppr = el.find(qn("p:spPr"))
    if sppr is not None:
        if sppr.find(qn("a:noFill")) is not None:
            return None
        sf = sppr.find(qn("a:solidFill"))
        if sf is not None:
            return color_from_fill(sf, theme)
        gf = sppr.find(qn("a:gradFill"))
        if gf is not None:
            return color_from_fill(gf, theme)
        if sppr.find(qn("a:blipFill")) is not None:
            return None
    style = el.find(qn("p:style"))
    if style is not None:
        fr = style.find(qn("a:fillRef"))
        if fr is not None and fr.get("idx") not in (None, "0"):
            return color_from_fill(fr, theme)
    return None


def background_hex(container, theme: Theme) -> str | None:
    """Цвет сплошного фона (или первой точки градиента) слайда/макета/мастера."""
    csld = container._element.find(qn("p:cSld"))
    if csld is None:
        return None
    bg = csld.find(qn("p:bg"))
    if bg is None:
        return None
    bgpr = bg.find(qn("p:bgPr"))
    if bgpr is not None:
        return color_from_fill(bgpr, theme)
    bgref = bg.find(qn("p:bgRef"))
    if bgref is not None:
        return color_from_fill(bgref, theme)
    return None


_FILLS = ("solidFill", "gradFill", "blipFill", "pattFill")


def paints(el) -> bool:
    """Рисует ли фигура что-то сама: заливку, контур или заливку из стиля темы?"""
    sp_pr = el.find(qn("p:spPr"))
    for c in sp_pr if sp_pr is not None else []:
        name = etree.QName(c).localname
        if name in _FILLS or (name == "ln" and any(etree.QName(x).localname in _FILLS for x in c)):
            return True
    ref = el.find(qn("p:style") + "/" + qn("a:fillRef"))
    return ref is not None and ref.get("idx", "0") != "0"
