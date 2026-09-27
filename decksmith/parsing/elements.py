"""Slide -> flat list of typed elements with absolute geometry and effective style."""
from __future__ import annotations

import re
from dataclasses import dataclass, field

from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from decksmith.core.models import Box
from decksmith.parsing.ooxml import (
    EffStyle,
    StyleResolver,
    Theme,
    flatten_shapes,
    placeholder_type,
    shape_fill_hex,
)

MONO_FONTS = {"consolas", "courier", "courier new", "menlo", "monaco", "jetbrains mono", "fira code", "source code pro"}


@dataclass
class Element:
    kind: str  # text | picture | table | chart | decor | line | group
    box: Box
    shape_id: int
    path: list[int]
    group_path: list[int]
    name: str = ""
    text: str = ""
    paragraphs: list[str] = field(default_factory=list)
    para_styles: list[EffStyle] = field(default_factory=list)
    style: EffStyle | None = None
    placeholder: str | None = None  # TITLE / BODY / PICTURE ...
    placeholder_idx: int | None = None
    fill_hex: str | None = None
    rounded: bool = False
    autofit: bool = False
    is_icon: bool = False
    image_hash: str | None = None
    image_ext: str | None = None
    table_shape: tuple[int, int] | None = None
    chart_type: str | None = None
    auto_shape: str | None = None
    has_line: bool = False
    insets: tuple[int, int, int, int] = (91440, 45720, 91440, 45720)  # l, t, r, b (EMU)
    anchor: str = "t"  # vertical text anchor: t | ctr | b
    brand: bool = False  # footer / date / slide number / text repeated on every slide: never a slot
    graphic: str | None = None  # SmartArt / OLE / media frames that cannot be refilled

    @property
    def is_title_ph(self) -> bool:
        return self.placeholder in ("TITLE", "CENTER_TITLE", "VERTICAL_TITLE")

    @property
    def font_size(self) -> float:
        return self.style.size if self.style and self.style.size else 0.0

    @property
    def is_mono(self) -> bool:
        return bool(self.style and self.style.font and self.style.font.lower() in MONO_FONTS)


def _ph_name(shape) -> str | None:
    pt = placeholder_type(shape)
    if pt is None:
        return None
    return str(pt).split(".")[-1].split(" ")[0]


def _autoshape_name(shape) -> str | None:
    geom = shape._element.find(".//" + qn("a:prstGeom"))
    return geom.get("prst") if geom is not None else None


def _has_autofit(shape) -> bool:
    txb = shape._element.find(qn("p:txBody"))
    if txb is None:
        return False
    bp = txb.find(qn("a:bodyPr"))
    return bp is not None and (bp.find(qn("a:spAutoFit")) is not None or bp.find(qn("a:normAutofit")) is not None)


def _table_box(shape, box: Box) -> Box:
    """Renderers draw a table at its grid size (column widths, row heights); some exporters
    (Google Slides) leave a dummy 3000000 EMU frame, so the frame size is not trusted."""
    tbl = shape._element.find(".//" + qn("a:tbl"))
    grid = tbl.find(qn("a:tblGrid")) if tbl is not None else None
    if grid is None:
        return box
    w = sum(int(gc.get("w", "0")) for gc in grid.findall(qn("a:gridCol")))
    h = sum(int(tr.get("h", "0")) for tr in tbl.findall(qn("a:tr")))
    try:
        sx, sy = box.w / int(shape.width), box.h / int(shape.height)  # group scaling
    except (TypeError, ValueError, ZeroDivisionError):
        sx = sy = 1.0
    return Box(x=box.x, y=box.y, w=int(w * sx) or box.w, h=int(h * sy) or box.h)


def text_insets(shape) -> tuple[int, int, int, int]:
    txb = shape._element.find(qn("p:txBody"))
    bp = txb.find(qn("a:bodyPr")) if txb is not None else None
    d = (91440, 45720, 91440, 45720)
    if bp is None:
        return d
    return tuple(int(bp.get(k)) if bp.get(k) is not None else dv for k, dv in zip(("lIns", "tIns", "rIns", "bIns"), d))  # type: ignore[return-value]


def text_anchor(shape) -> str:
    """Vertical text anchor (t / ctr / b), following placeholder inheritance to layout and master."""
    node = shape
    for _ in range(3):
        txb = node._element.find(qn("p:txBody"))
        bp = txb.find(qn("a:bodyPr")) if txb is not None else None
        if bp is not None and bp.get("anchor"):
            return bp.get("anchor")
        if not getattr(node, "is_placeholder", False):
            break
        try:
            node = node._base_placeholder
        except Exception:
            break
        if node is None:
            break
    return "t"


def _has_field(shape) -> bool:
    """Text boxes carrying slide-number / date fields behave like footers."""
    txb = shape._element.find(qn("p:txBody"))
    if txb is None:
        return False
    return any((f.get("type") or "").startswith(("slidenum", "datetime")) for f in txb.iter(qn("a:fld")))


def _line_visible(shape) -> bool:
    ln = shape._element.find(".//" + qn("a:ln"))
    if ln is None:
        return False
    return ln.find(qn("a:noFill")) is None


def extract_elements(slide, theme: Theme, slide_w: int, slide_h: int, include_empty_placeholders: bool = True) -> list[Element]:
    resolver = StyleResolver(slide, theme)
    out: list[Element] = []
    icon_limit = 0.11 * min(slide_w, slide_h) * 1.6
    for rec in flatten_shapes(slide.shapes):
        sh = rec.shape
        st = sh.shape_type
        base = dict(box=rec.box, shape_id=sh.shape_id, path=rec.path, group_path=rec.group_path, name=sh.name or "")
        if st == MSO_SHAPE_TYPE.GROUP:
            out.append(Element(kind="group", **base))
            continue
        ph = _ph_name(sh)
        ph_idx = None
        if ph is not None:
            try:
                ph_idx = sh.placeholder_format.idx
            except Exception:
                ph_idx = None
        if getattr(sh, "has_table", False) and sh.has_table:
            tbl = sh.table
            base["box"] = _table_box(sh, rec.box)
            out.append(Element(kind="table", table_shape=(len(tbl.rows), len(tbl.columns)), placeholder=ph, placeholder_idx=ph_idx, **base))
            continue
        if getattr(sh, "has_chart", False) and sh.has_chart:
            try:
                ct = str(sh.chart.chart_type).split(".")[-1].split(" ")[0]
            except Exception:
                ct = "unknown"
            out.append(Element(kind="chart", chart_type=ct, placeholder=ph, placeholder_idx=ph_idx, **base))
            continue
        if sh._element.tag == qn("p:graphicFrame"):
            gd = sh._element.find(".//" + qn("a:graphicData"))
            uri = gd.get("uri", "") if gd is not None else ""
            kind = "smartart" if "diagram" in uri else "ole" if "ole" in uri else "graphic"
            out.append(Element(kind="decor", graphic=kind, placeholder=ph, placeholder_idx=ph_idx, **base))
            continue
        if st == MSO_SHAPE_TYPE.PICTURE or (ph == "PICTURE" and not getattr(sh, "has_text_frame", False)):
            h = ext = None
            try:
                h = sh.image.sha1
                ext = sh.image.ext
            except Exception:
                pass
            is_icon = max(rec.box.w, rec.box.h) <= icon_limit
            out.append(Element(kind="picture", placeholder=ph, placeholder_idx=ph_idx, is_icon=is_icon, image_hash=h, image_ext=ext, **base))
            continue
        if st == MSO_SHAPE_TYPE.LINE or (st == MSO_SHAPE_TYPE.AUTO_SHAPE and _autoshape_name(sh) in ("line", "straightConnector1")):
            out.append(Element(kind="line", has_line=True, **base))
            continue
        text = ""
        paragraphs: list[str] = []
        if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
            paragraphs = [p.text for p in sh.text_frame.paragraphs]
            text = "\n".join(paragraphs).strip()
        fill = None
        if st in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.TEXT_BOX, MSO_SHAPE_TYPE.FREEFORM) or ph is not None:
            try:
                fill = shape_fill_hex(sh, theme)
            except Exception:
                fill = None
        prst = _autoshape_name(sh)
        if text or (ph is not None and include_empty_placeholders and ph not in ("SLIDE_NUMBER", "DATE", "FOOTER")):
            style = resolver.dominant(sh) if getattr(sh, "has_text_frame", False) and sh.has_text_frame else None
            para_styles = []
            if text:
                for p in sh.text_frame.paragraphs:
                    runs = [r for r in p.runs if r.text.strip()]
                    if runs:
                        para_styles.append(resolver.resolve(sh, p, runs[0]))
            if ph == "PICTURE":
                out.append(Element(kind="picture", placeholder=ph, placeholder_idx=ph_idx, **base))
                continue
            out.append(
                Element(
                    kind="text",
                    text=text,
                    paragraphs=[p for p in paragraphs if p.strip()],
                    para_styles=para_styles,
                    style=style,
                    placeholder=ph,
                    placeholder_idx=ph_idx,
                    fill_hex=fill,
                    rounded=prst in ("roundRect", "round2SameRect", "snipRoundRect"),
                    autofit=_has_autofit(sh),
                    auto_shape=prst,
                    has_line=_line_visible(sh),
                    insets=text_insets(sh),
                    anchor=text_anchor(sh),
                    brand=ph in ("FOOTER", "DATE", "SLIDE_NUMBER") or _has_field(sh),
                    **base,
                )
            )
        else:
            out.append(
                Element(
                    kind="decor",
                    fill_hex=fill,
                    rounded=prst in ("roundRect", "round2SameRect", "snipRoundRect"),
                    auto_shape=prst,
                    has_line=_line_visible(sh),
                    placeholder=ph,
                    placeholder_idx=ph_idx,
                    **base,
                )
            )
    return out


# ----------------------------------------------------------------------------
# Placeholder-text heuristics (language-agnostic-ish, RU + EN)
# ----------------------------------------------------------------------------
_FILLER_RE = re.compile(
    r"^(заголовок|подзаголовок|текст|пункт|описание|название|имя фамилия|должность|показатель|итого|"
    r"title|subtitle|text|heading|header|body|item|description|caption|name|position|lorem ipsum.*|"
    r"x+%?|х+%?|xx+.*|хх+.*|\d{1,2}|0\d|\d+%|>?\d+\*?|кейс|дата|date|заметка|примечание|"
    r"вставить\s+(фото|qr)|qr-?code|qr-код|ссылка|link|кнопка|button|перейти|графики|таймлайн|timeline)$",
    re.IGNORECASE,
)


def is_filler_text(text: str) -> bool:
    t = re.sub(r"\s+", " ", text.strip().replace("​", ""))
    if not t:
        return True
    if t.lower().startswith("lorem ipsum"):
        return True
    return bool(_FILLER_RE.match(t))


def is_numberish(text: str) -> bool:
    t = text.strip().replace(" ", "").replace(" ", "")
    return bool(re.fullmatch(r"[<>~≈+\-]?[\d.,]+\s*(%|₽|\$|€|млн|млрд|тыс|k|m|x|×|\*)?|[xхXХ]{1,4}%?|0?\d", t))
