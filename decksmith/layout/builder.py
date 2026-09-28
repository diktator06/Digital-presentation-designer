"""Сборщик колоды: применяет решения по вёрстке к копии шаблона.

clone   : копирует пример-слайд, сокращает повторитель, раскладывает контент по слотам
          по ролям, вписывает текст по шкале кеглей шаблона, удаляет неиспользованные слоты
compose : новый слайд на макете-«холсте» шаблона (плейсхолдер заголовка сохраняется),
          нативные диаграмма/таблица/схема рисуются в зоне контента, найденной CV

От слайда-примера не остаётся ничего, кроме заполненного новым контентом и чистого
оформления (поэтому никаких «Заголовок»/«Lorem ipsum»).
"""
from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass, field
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn

from decksmith.core.models import (
    TABLE_MAX_COLS,
    TABLE_MAX_ROWS,
    Box,
    DeckPlan,
    Item,
    Pattern,
    PatternKind,
    SlideLayout,
    SlideSpec,
    Slot,
    SlotRole,
    TemplateProfile,
)
from decksmith.layout import compose as C
from decksmith.layout.pptx_ops import (
    delete_shape,
    delete_slides,
    duplicate_slide,
    fill_chart,
    fill_table,
    fit_table_text,
    move_shape,
    replace_picture,
    scale_font_sizes,
    scale_paragraph_sizes,
    set_font_size,
    set_paragraphs,
    shape_by_id,
    table_has_merges,
)
from decksmith.layout.selector import Variant
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR

from decksmith.layout.room import (clear_title_box, draws, grow_frame, is_plate, owned_height, paragraph_styles, spacing,
                                   text_height, text_room)
from decksmith.parsing.ooxml import flatten_shapes
from decksmith.layout.textfit import DEFAULT_INSET_TB, fit_composite, fit_font_size, max_chars_for, measure, snap_down
from decksmith.parsing.ooxml import contrast_ratio, rel_luminance
from decksmith.parsing.template_parser import region_color, region_colors, title_clear_box

log = logging.getLogger(__name__)
# заглушка названия презентации в колонтитулах шаблона (RU/EN)
TITLE_FILLER = re.compile(
    r"^\s*((название|заголовок|тема)\s+(вашей\s+|нашей\s+)?(презентации|доклада|выступления)"
    r"|(your\s+)?(presentation|deck)\s+(title|name)|(title|name)\s+of\s+(the\s+|your\s+)?presentation)\s*$",
    re.IGNORECASE)
EMU_PT = 12700


@dataclass
class Overflow:
    slide: int
    shape_id: int
    role: str
    paragraphs: list[str]
    budget: int


@dataclass
class BuildReport:
    slides: list[dict] = field(default_factory=list)
    overflows: list[Overflow] = field(default_factory=list)


def _layouts(prs):
    """Все макеты всех мастеров презентации."""
    return [layout for m in prs.slide_masters for layout in m.slide_layouts]


def _is_decorative_picture(shape) -> bool:
    """Прозрачная PNG-графика (3D-объекты, пятна) — оформление бренда; непрозрачные фото — контент."""
    try:
        blob = shape.image.blob
        im = Image.open(io.BytesIO(blob))
        if im.mode in ("RGBA", "LA") or (im.mode == "P" and "transparency" in im.info):
            a = im.convert("RGBA").getchannel("A").resize((32, 32))
            return a.getextrema()[0] < 200
    except Exception:
        return False
    return False


def number_format(sample: str, i: int) -> str | None:
    """Номер элемента в формате образца шаблона («01» или «1»); None, если образец — не номер."""
    s = sample.strip()
    if re.fullmatch(r"0\d", s):
        return f"{i + 1:02d}"
    if re.fullmatch(r"\d{1,2}", s):
        return f"{i + 1}"
    return None


class DeckBuilder:
    def __init__(self, profile: TemplateProfile, variant: Variant):
        """Открывает шаблон и готовит индексы макетов и слайдов-примеров."""
        self.profile = profile
        self.variant = variant
        self.prs = Presentation(profile.file)
        self.n_src = len(self.prs.slides)
        self.src = list(self.prs.slides)
        self.layouts = _layouts(self.prs)
        self.report = BuildReport()
        self._bold_first: set[str] = set()
        self._inline_values = False
        self._bg_cache: dict[tuple, str | None] = {}
        self._clear_cache: dict[tuple, Box | None] = {}
        self._clear_prompt_texts()

    def _clear_prompt_texts(self) -> None:
        """Образец текста в плейсхолдерах макета/мастера («Click to edit», «Образец текста») — подсказка, а не
        дизайн. Некоторые рендеры (LibreOffice) рисуют его на каждом слайде этого макета, поэтому в
        выходной колоде он очищается; раны и абзацы остаются, так что унаследованное форматирование не
        меняется. Плейсхолдеры колонтитула, даты и номера слайда сохраняют содержимое.
        """
        keep = {"FOOTER", "DATE", "SLIDE_NUMBER"}
        for container in list(self.prs.slide_masters) + self.layouts:
            for ph in container.placeholders:
                try:
                    if str(ph.placeholder_format.type).split(".")[-1].split(" ")[0] in keep:
                        continue
                except Exception:
                    continue
                for t in ph._element.iter(qn("a:t")):
                    t.text = ""

    # ------------------------------------------------------------------ публичные методы
    def build(self, plan: DeckPlan, decisions: list[SlideLayout], images: dict[str, str] | None = None,
              icons: dict[str, list[str | None]] | None = None) -> BuildReport:
        images = images or {}
        icons = icons or {}
        specs = {s.id: s for s in plan.slides}
        # заглушка «Название презентации» в колонтитулах мастеров и макетов получает название колоды
        self._title_footers(list(self.prs.slide_masters) + self.layouts, plan.title)
        for n, d in enumerate(decisions):
            spec = specs[d.spec_id]
            before = len(self.prs.slides)
            try:
                if d.mode == "clone" and d.pattern_id:
                    pat = self.profile.pattern(d.pattern_id)
                    slide = self._clone(pat)
                    self._fill_clone(slide, n, pat, spec, d, plan, images.get(spec.id), icons.get(spec.id))
                else:
                    slide = self._compose(n, spec, d.compose_kind or "bullets", images.get(spec.id), icons.get(spec.id))
            # слайд не теряется никогда: полусобранный удаляется, вместо него собирается список
            except Exception as e:
                log.exception("slide %s failed (%s), falling back", spec.id, e)
                if len(self.prs.slides) > before:
                    delete_slides(self.prs, list(range(before, len(self.prs.slides))))
                slide = self._compose(n, spec, "bullets", None, None)
                d.rationale += f" | fallback after error: {e}"
            if spec.notes:
                slide.notes_slide.notes_text_frame.text = spec.notes
            self._title_footers([slide], plan.title)
            for fld in slide._element.iter(qn("a:fld")):  # кэш номеров страниц у скопированных примеров
                if fld.get("type") == "slidenum" and fld.find(qn("a:t")) is not None:
                    fld.find(qn("a:t")).text = str(n + 1)
            source = None
            kind = d.compose_kind or ""
            if d.mode == "clone" and d.pattern_id:
                pat = self.profile.pattern(d.pattern_id)
                source, kind = pat.slide_index, pat.kind.value
            if spec.intent.value in ("title", "section", "thanks", "quote", "contacts"):
                kind = spec.intent.value
            self.report.slides.append({"spec": spec.id, "mode": d.mode, "pattern": d.pattern_id or d.compose_kind,
                                       "source": source, "kind": kind, "rationale": d.rationale})
        delete_slides(self.prs, list(range(self.n_src)))
        return self.report

    @staticmethod
    def _title_footers(containers: list, title: str) -> None:
        """Колонтитул с заглушкой названия презентации («Название презентации», «Presentation title») —
        частая боковая или нижняя надпись шаблонов — получает название колоды. Настоящие фирменные
        колонтитулы (компания, пометка о конфиденциальности) не трогаются.
        """
        title = " ".join(title.split())
        if not title:
            return
        for c in containers:
            for ph in c.placeholders:
                try:
                    footer = "FOOTER" in str(ph.placeholder_format.type)
                except Exception:
                    continue
                if not (footer and ph.has_text_frame and TITLE_FILLER.match(ph.text_frame.text)):
                    continue
                # колонтитул рассчитан на короткую надпись: длинное название сокращается по границе слова
                limit = max(2 * len(ph.text_frame.text.strip()), 32)
                text = title
                if len(text) > limit:
                    text = text[:limit].rsplit(" ", 1)[0].rstrip(" ,.;:—-") + "…"
                set_paragraphs(ph, [text])

    def save(self, path: str | Path) -> Path:
        """Сохраняет колоду."""
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(path))
        return path

    # ------------------------------------------------------------------ клонирование
    def _clone(self, pat: Pattern):
        """Новый слайд из слайда-примера или макета паттерна."""
        if pat.source == "slide" and pat.slide_index is not None:
            return duplicate_slide(self.prs, self.src[pat.slide_index])
        return self._slide_from_layout(self.layouts[pat.layout_index])

    def _slide_from_layout(self, layout):
        """Новый слайд, у плейсхолдеров которого явно задана геометрия макета. Некоторые программы записывают
        все плейсхолдеры с одним idx (например, idx=0), и наследование по idx дало бы основному тексту
        рамку заголовка; фиксация геометрии пары каждого плейсхолдера в макете снимает неоднозначность.
        Бренд-элементы, повторяющиеся на примерах шаблона, тоже переносятся.
        """
        import copy as _copy

        slide = self.prs.slides.add_slide(layout)
        cloneable = list(layout.iter_cloneable_placeholders())
        sw, sh_ = self.profile.tokens.slide_w, self.profile.tokens.slide_h
        pad = int(0.01 * sw)
        # обе стороны в порядке документа (slide.placeholders отсортированы по idx, копии — нет)
        for sph, lph in zip([s for s in slide.shapes if s.is_placeholder], cloneable):
            try:
                x, y, w, h = lph.left, lph.top, lph.width, lph.height
                if None in (x, y, w, h):
                    continue
                # в эти рамки идёт наш текст: оставляем их на слайде, даже если макет выходит за край
                x, y = max(int(x), 0), max(int(y), 0)
                w = min(int(w), sw - x - pad) if x + int(w) > sw else int(w)
                h = min(int(h), sh_ - y - pad) if y + int(h) > sh_ else int(h)
                sph.left, sph.top, sph.width, sph.height = x, y, max(w, pad), max(h, pad)
            except Exception:
                continue
        masters = list(self.prs.slide_masters)
        m_idx = next((i for i, m in enumerate(masters) if m.part is layout.slide_master.part), 0)
        for f in self.profile.brand_furniture:
            if f.get("master", 0) != m_idx or f["slide"] >= len(self.src):
                continue
            src_shape = shape_by_id(self.src[f["slide"]], f["shape_id"])
            if src_shape is not None:
                el = _copy.deepcopy(src_shape._element)
                cnv = el.find(".//" + qn("p:cNvPr"))
                if cnv is not None:
                    cnv.set("id", str(slide.shapes._next_shape_id))  # уникальный id на новом слайде
                slide.shapes._spTree.append(el)
        return slide

    def _clear_title_box(self, layout_index: int | None, box: Box) -> Box | None:
        """Рамка заголовка, укороченная до фоновой графики, которую макет рисует в полосе заголовка (полоса
        логотипов в фоне, угловая графика). None, если рамка уже свободна или не в этой полосе.
        """
        layouts = self.profile.layouts
        if layout_index is None or not 0 <= layout_index < len(layouts):
            return None
        return clear_title_box(box, layouts[layout_index].title_clear)

    def _clear_text_box(self, layout_index: int | None, box: Box) -> Box | None:
        """Текстовая рамка, укороченная до графики, которую макет рисует на её уровне (орнамент, фото в фоне):
        подзаголовок или текст длиннее, чем в примере, не должен уходить под неё. None — рамка свободна.
        """
        layouts, t = self.profile.layouts, self.profile.tokens
        if layout_index is None or not 0 <= layout_index < len(layouts):
            return None
        li = layouts[layout_index]
        if not li.thumbnail or not Path(li.thumbnail).exists():
            return None
        key = (layout_index, box.x, box.y, box.w, box.h)
        if key not in self._clear_cache:
            # тот же разбор рендера пустого макета, что и для полосы заголовка
            self._clear_cache[key] = title_clear_box(Path(li.thumbnail), box, t.slide_w, t.slide_h,
                                                     li.background_hex or t.background_hex)
        return self._clear_cache[key]

    def _shape(self, slide, pat: Pattern, slot: Slot):
        if pat.source == "layout":
            # плейсхолдеры наследуют геометрию макета: сопоставление по idx, если он уникален, иначе по
            # положению
            phs = list(slide.placeholders)
            by_idx = [ph for ph in phs if ph.placeholder_format.idx == slot.placeholder_idx]
            if len(by_idx) == 1:
                return by_idx[0]

            def dist(ph):
                """Расстояние между плейсхолдером и рамкой слота."""
                try:
                    return abs(ph.left - slot.box.x) + abs(ph.top - slot.box.y) + abs(ph.width - slot.box.w) + abs(ph.height - slot.box.h)
                except TypeError:
                    return 1 << 62
            best = min(phs, key=dist, default=None)
            if best is not None and dist(best) < 0.05 * (self.profile.tokens.slide_w + self.profile.tokens.slide_h):
                return best
            return None
        return shape_by_id(slide, slot.shape_id)

    def _reduce_repeater(self, slide, pat: Pattern, keep: int) -> None:
        rep = pat.repeaters[0]
        n = rep.n_items
        removed_ids = [sid for ids in rep.item_shape_ids[keep:] for sid in ids]
        for sid in removed_ids:
            sh = shape_by_id(slide, sid)
            if sh is not None:
                delete_shape(sh)
        # оставшиеся элементы распределяются по исходной ширине (только фигуры верхнего уровня)
        boxes = rep.item_boxes
        if rep.direction not in ("row", "column") or keep < 1:
            return
        shapes = [[shape_by_id(slide, sid) for sid in ids] for ids in rep.item_shape_ids[:keep]]
        if any(s is None or s._element.getparent().tag.endswith("grpSp") for ss in shapes for s in ss):
            return
        if rep.direction == "row":
            span0, span1 = boxes[0].x, boxes[n - 1].r
            w = boxes[0].w
            # оставшиеся карточки сохраняют промежуток примера и встают по центру ряда (а не расходятся к краям)
            gap = max(boxes[1].x - boxes[0].r, 0) if n > 1 else 0
            start = span0 + (span1 - span0 - (keep * w + (keep - 1) * gap)) / 2
            for i in range(keep):
                nx = start + i * (w + gap)
                for s in shapes[i]:
                    move_shape(s, int(nx - boxes[i].x), 0)
        else:
            span0, span1 = boxes[0].y, boxes[n - 1].b
            h = boxes[0].h
            gap = (span1 - span0 - keep * h) / max(keep - 1, 1) if keep > 1 else 0
            gap = min(gap, h * 0.6)
            for i in range(keep):
                ny = span0 + i * (h + gap)
                for s in shapes[i]:
                    move_shape(s, 0, int(ny - boxes[i].y))

    def _fill_clone(self, slide, n: int, pat: Pattern, spec: SlideSpec, d: SlideLayout, plan: DeckPlan,
                    image: str | None, icons: list[str | None] | None) -> None:
        self._bold_first = set()
        rep = pat.repeaters[0] if pat.repeaters else None
        items = list(spec.items)
        if rep and not items and spec.bullets:
            items = [Item(title=b) for b in spec.bullets]
        keep = d.keep_items if d.keep_items else (rep.n_items if rep else 0)
        if rep and keep < rep.n_items:
            self._reduce_repeater(slide, pat, keep)
        # шаблон обложки, взятый для разделителя, сохраняет текст разделителя
        is_cover = spec.intent == PatternKind.title
        body_queue: list[str] = list(spec.bullets) if not rep else []
        free_numbers = [s for s in pat.slots if s.role == SlotRole.number and s.item_index is None]
        free_texts_for_numbers = items[len([1 for _ in range(keep)]):] if rep else items
        stat_units = list(free_texts_for_numbers) if free_numbers else []
        fills: dict[str, tuple[list[str], list[int] | None]] = {}
        body_slots = sorted([s for s in pat.slots if s.role == SlotRole.body and s.item_index is None], key=lambda s: (s.box.y, s.box.x))

        # какие текстовые слоты есть у элемента? (чтобы упаковать заголовок и текст, если слот один)
        item_text_slots: dict[int, list[Slot]] = {}
        for s in pat.slots:
            if s.kind == "text" and s.item_index is not None and s.role not in (SlotRole.number, SlotRole.label):
                item_text_slots.setdefault(s.item_index, []).append(s)
        is_quote = pat.kind == PatternKind.quote and bool(spec.quote)
        # у элементов со значениями (KPI) на шаблоне без слота числа значение остаётся в заголовке
        self._inline_values = not any(s.role == SlotRole.number and s.item_index is not None for s in pat.slots) and not any(
            s.para_roles and s.para_roles[0] == SlotRole.number for s in pat.slots)
        used_item_roles: set[tuple[int, str]] = set()
        used_item_texts: dict[int, set[str]] = {}
        used_texts: set[str] = set()

        def once(txt: str) -> str:
            """Каждый фрагмент свободного текста попадает на слайд не больше одного раза."""
            if not txt or txt in used_texts:
                return ""
            used_texts.add(txt)
            return txt

        for slot in sorted(pat.slots, key=lambda s: (s.box.y, s.box.x)):
            if slot.kind != "text" or slot.role == SlotRole.decor:
                continue
            if slot.item_index is not None:
                if slot.item_index >= keep:
                    continue
                it = items[slot.item_index] if slot.item_index < len(items) else None
                if it is None:
                    continue
                key = (slot.item_index, slot.role.value if not slot.para_roles else "composite")
                if key in used_item_roles:
                    continue  # например, текст таймлайна, продублированный над и под осью
                used_item_roles.add(key)
                only = len(item_text_slots.get(slot.item_index, [])) == 1 and not slot.para_roles
                if only and slot.role not in (SlotRole.number, SlotRole.label) and it.title and it.text:
                    fills[slot.id] = ([self._head(it), it.text], [0, 0])
                    self._bold_first.add(slot.id)
                else:
                    paras, tidx = self._item_text(slot, it, slot.item_index)
                    # элемент без заголовка не должен показывать свой текст дважды (слот заголовка + слот
                    # текста)
                    item_seen = used_item_texts.setdefault(slot.item_index, set())
                    if paras and all(p in item_seen for p in paras if p):
                        continue
                    item_seen.update(p for p in paras if p)
                    fills[slot.id] = (paras, tidx)
                continue
            role = slot.role
            if role == SlotRole.title:
                if is_quote:
                    fills[slot.id] = ([spec.quote], None)
                else:
                    fills[slot.id] = ([plan.title if is_cover else spec.title], None)
            elif role == SlotRole.subtitle:
                txt = once((plan.subtitle if is_cover else "") or spec.subtitle) or once(spec.message)
                if txt:
                    fills[slot.id] = ([txt], None)
            elif role == SlotRole.number and stat_units:
                it = stat_units.pop(0)
                fills[slot.id] = ([it.value or it.title], None)
                # пара — ближайший подписеподобный текстовый слот под числом
                cap = self._nearest_below(slot, pat, fills)
                if cap is not None and (it.text or it.title):
                    fills[cap.id] = ([it.text or it.title], None)
            elif role == SlotRole.body:
                if slot.para_roles:
                    # составная рамка: оформленный ведущий абзац + абзацы основного текста
                    rest = list(body_queue[: self.variant.max_bullets]) if body_queue else ([m] if (m := once(spec.message)) else [])
                    head = once(spec.subtitle) if rest else ""
                    if body_queue:
                        body_queue = []
                    if rest:
                        paras = ([head] if head else []) + rest
                        fills[slot.id] = (paras, ([0] if head else []) + [1] * len(rest))
                    continue
                if len(body_slots) > 1 and body_queue:
                    # пункты распределяются по нескольким рамкам основного текста в порядке чтения
                    k = body_slots.index(slot)
                    per = -(-len(spec.bullets) // len(body_slots))
                    chunk = spec.bullets[k * per:(k + 1) * per]
                    if chunk:
                        fills[slot.id] = (chunk, None)
                elif body_queue:
                    fills[slot.id] = (body_queue[: self.variant.max_bullets], None)
                    body_queue = []
                elif is_quote and spec.quote_author:
                    if once(spec.quote_author):
                        fills[slot.id] = ([spec.quote_author], None)
                elif spec.quote and not is_quote:
                    if once(spec.quote):
                        fills[slot.id] = ([spec.quote], None)
                else:
                    txt = once(spec.message) or once((plan.subtitle if is_cover else "") or spec.subtitle)
                    if txt:
                        fills[slot.id] = ([txt], None)
            elif role in (SlotRole.caption, SlotRole.label, SlotRole.person) and is_quote and spec.quote_author:
                if once(spec.quote_author):
                    fills[slot.id] = ([spec.quote_author], None)

        # макеты цитаты без слота заголовка: цитата идёт в самую большую текстовую рамку
        if is_quote and not any(spec.quote in paras for paras, _ in fills.values()):
            frames = sorted([s for s in pat.slots if s.kind == "text" and s.role != SlotRole.decor and s.item_index is None],
                            key=lambda s: -s.box.area)
            if frames:
                displaced = fills.get(frames[0].id)
                fills[frames[0].id] = ([spec.quote], None)
                if displaced and spec.quote_author in displaced[0] and len(frames) > 1:
                    fills[frames[1].id] = ([spec.quote_author], None)

        # ведущая строка разделителя или финального слайда (призыв к действию в питче) идёт в свободную
        # подписеподобную рамку, если у примера нет рамки подзаголовка, — вместо того чтобы потеряться
        lead = spec.message or spec.subtitle
        lead_slot = None
        if spec.intent in (PatternKind.section, PatternKind.thanks) and lead and lead not in used_texts:
            spare = sorted([s for s in pat.slots if s.kind == "text" and s.item_index is None and s.id not in fills
                            and s.role in (SlotRole.subtitle, SlotRole.body, SlotRole.person, SlotRole.caption, SlotRole.label)],
                           key=lambda s: -s.box.area)
            if spare:
                fills[spare[0].id] = ([once(lead)], None)
                lead_slot = spare[0]

        # ведущая строка содержательного слайда, не нашедшая рамки подзаголовка, идёт в свободную рамку
        # примечания на плашке (полоса-вывод примера), а не оставляет полосу пустой
        lead = spec.subtitle or spec.message
        if spec.intent not in (PatternKind.title, PatternKind.section, PatternKind.thanks, PatternKind.quote) and lead \
                and lead not in used_texts:
            plates = [r.box for r in flatten_shapes(slide.shapes) if r.shape.shape_type != MSO_SHAPE_TYPE.GROUP
                      and is_plate(r.shape) and r.box.area < 0.85 * self.profile.tokens.slide_w * self.profile.tokens.slide_h]
            notes = sorted([s for s in pat.slots if s.kind == "text" and s.item_index is None and s.id not in fills
                            and s.role in (SlotRole.label, SlotRole.caption) and any(b.contains(s.box, tol=12700) for b in plates)],
                           key=lambda s: -s.box.w)
            if notes:
                fills[notes[0].id] = ([once(lead)], None)

        # применяем текст (подгонка — в конце)
        deleted: set[str] = set()
        placed: list[tuple[Slot, object, list[str], list[int] | None]] = []
        for slot in pat.slots:
            if slot.kind != "text":
                continue
            if slot.item_index is not None and slot.item_index >= keep:
                continue
            sh = self._shape(slide, pat, slot)
            if sh is None:
                continue
            if slot.id not in fills or not any(p.strip() for p in fills[slot.id][0] if p):
                delete_shape(sh)
                deleted.add(slot.id)
                continue
            paras, tidx = fills[slot.id]
            paras = [p for p in paras if p is not None]
            set_paragraphs(sh, paras, tidx)
            if slot.id in self._bold_first:
                runs = sh.text_frame.paragraphs[0].runs
                if runs:
                    runs[0].font.bold = True
            placed.append((slot, sh, paras, tidx))
        # картинки
        pics = [s for s in pat.slots if s.kind == "picture" and (s.item_index is None or s.item_index < keep)]
        content_pics = sorted([s for s in pics if s.role == SlotRole.image], key=lambda s: -s.box.area)
        used_image = False
        for slot in content_pics:
            sh = self._shape(slide, pat, slot)
            if sh is None:
                continue
            if image and not used_image and slot.box.area > 0.04 * self.profile.tokens.slide_w * self.profile.tokens.slide_h:
                try:
                    if sh.is_placeholder and not hasattr(sh, "image"):
                        sh.insert_picture(image)
                    else:
                        replace_picture(slide, sh, image)
                    used_image = True
                    continue
                except Exception as e:
                    log.warning("image replace failed: %s", e)
            if getattr(sh, "is_placeholder", False) and sh.shape_type != 13:
                delete_shape(sh)  # пустой плейсхолдер картинки
                deleted.add(slot.id)
            elif not _is_decorative_picture(sh):
                delete_shape(sh)
                deleted.add(slot.id)
        for slot in pics:
            if slot.role == SlotRole.icon and icons and slot.item_index is not None and slot.item_index < len(icons):
                ic = icons[slot.item_index]
                sh = self._shape(slide, pat, slot)
                if ic and sh is not None and hasattr(sh, "image"):
                    try:
                        replace_picture(slide, sh, ic, mode="contain")
                    except Exception:
                        pass
        # рамки и аватары, служившие только удалённому слоту, тоже удаляются
        keep_ids = {s.shape_id for s in pat.slots if s.id not in deleted}
        by_container: dict[int, list[str]] = {}
        for sid, cids in pat.containers.items():
            for c in cids:
                by_container.setdefault(c, []).append(sid)
        for cid, sids in by_container.items():
            if cid in keep_ids:
                continue
            if all(s in deleted for s in sids):
                sh = shape_by_id(slide, cid)
                if sh is not None:
                    delete_shape(sh)
        self._drop_empty_plates(slide, [s.box for s in pat.slots if s.id in deleted and s.kind == "text"])
        # рамка спикера, получившая ведущую строку, теряет соседний аватар
        if lead_slot is not None and lead_slot.role == SlotRole.person:
            for cid in pat.containers.get(lead_slot.id, []):
                sh = shape_by_id(slide, cid)
                if sh is not None and sh.width is not None and not Box(
                        x=int(sh.left), y=int(sh.top), w=int(sh.width), h=int(sh.height)).contains(lead_slot.box, tol=12700):
                    delete_shape(sh)
        # нативные таблицы и диаграммы примера: заполняются нашими данными в стиле шаблона
        t = self.profile.tokens
        for slot in [s for s in pat.slots if s.kind in ("table", "chart")]:
            sh = self._shape(slide, pat, slot)
            if sh is None:
                continue
            table = spec.table or (_chart_as_table(spec).model_copy() if spec.chart and slot.kind == "table" else None)
            try:
                if slot.kind == "table" and table and not table_has_merges(sh):
                    fill_table(sh, table.columns[:TABLE_MAX_COLS], [r[:TABLE_MAX_COLS] for r in table.rows[:TABLE_MAX_ROWS - 1]],
                               t.slide_h - t.margins.bottom)
                    if fit_table_text(sh, t.slide_h - t.margins.bottom, t.body_font, max(t.type_scale.caption, 10.0)):
                        continue
                    log.info("template table cannot hold the data readably, redrawing natively")
                if slot.kind == "chart" and spec.chart and spec.chart.series:
                    fill_chart(sh, spec.chart.categories[:8], [(s.name, s.values) for s in spec.chart.series], spec.chart.unit)
                    continue
            except Exception as e:  # необычный XML диаграммы/таблицы: рисуем нативно в той же рамке
                log.warning("native refill failed (%s), redrawing", e)
            box = slot.box
            if slot.kind == "table" and not any(o.box.y >= box.b and min(o.box.r, box.r) > max(o.box.x, box.x)
                                                for o in pat.slots if o is not slot and o.kind != "picture"):
                # ниже ничего нет
                box = Box(x=box.x, y=box.y, w=box.w, h=max(box.h, t.slide_h - t.margins.bottom - box.y))
            delete_shape(sh)
            st = self._style_for(slide_bg=self.profile.tokens.background_hex)
            if slot.kind == "chart" and spec.chart:
                C.draw_chart(slide, box, spec.chart, st)
            elif table:
                C.draw_table(slide, box, table, st)
        # оставшиеся пустые плейсхолдеры (у копий python-pptx их не добавляет, макеты могут оставить)
        for ph in list(slide.placeholders):
            if ph.has_text_frame and not ph.text_frame.text.strip():
                delete_shape(ph)
        # текст подгоняется последним: рамка может вырасти в место, освобождённое неиспользованными слотами
        for slot, sh, paras, tidx in placed:
            fit_slot = slot
            if slot.item_index is None and slot.role in (SlotRole.title, SlotRole.subtitle, SlotRole.body):
                if slot.role == SlotRole.title:
                    clear = self._clear_title_box(pat.layout_index, slot.box)
                else:
                    clear = self._clear_text_box(pat.layout_index, slot.box)
                if clear is not None:
                    # фиксируем все четыре: плейсхолдеры могут наследовать
                    x, y, h = sh.left, sh.top, sh.height
                    sh.left, sh.top, sh.width, sh.height = x, y, clear.w, h
                    fit_slot = slot.model_copy(update={"box": clear})
            owned = owned_height(slot, [o for o, *_ in placed])
            if owned < fit_slot.box.h and slot.role in (SlotRole.title, SlotRole.subtitle):
                # однострочная рамка заголовка примера с подзаголовком прямо под первой строкой:
                # более длинный заголовок держится в части рамки над следующей заполненной рамкой
                b = fit_slot.box
                fit_slot = fit_slot.model_copy(update={"box": Box(x=b.x, y=b.y, w=b.w, h=owned)})
                x, y, w = sh.left, sh.top, sh.width
                sh.left, sh.top, sh.width, sh.height = x, y, w, int(owned * sh.height / max(slot.box.h, 1))
            self._fit(slide, n, sh, fit_slot, paras, tidx)
            if slot.id in pat.backdrops:
                self._fit_backdrop(slide, pat.backdrops[slot.id], sh, fit_slot, paras)
            else:
                self._ensure_contrast(sh, slot, pat)

    def _ensure_contrast(self, sh, slot: Slot, pat: Pattern) -> None:
        """Слот, собственный цвет которого нечитаем на том, что шаблон рисует под ним (ниже 3:1, например
        фиолетовый заголовок на фиолетовой карточке), получает самый читаемый из текстовых цветов шаблона.
        Умеренные случаи остаются как задумано; о них сообщает аудит.
        """
        col = slot.style.color_hex
        if not col or not pat.thumbnail or not Path(pat.thumbnail).exists():
            return
        key = (pat.id, slot.box.x, slot.box.y, slot.box.w, slot.box.h)
        if key not in self._bg_cache:
            try:
                t = self.profile.tokens
                self._bg_cache[key] = region_color(Path(pat.thumbnail), slot.box, t.slide_w, t.slide_h)
            except Exception:
                self._bg_cache[key] = None
        bg = self._bg_cache[key]
        if bg is None or contrast_ratio(col, bg) >= 3.0:  # минимум WCAG для крупного текста
            return
        own = [s.style.color_hex for s in pat.slots if s.style.color_hex and s.style.color_hex != col]
        cands = own + [self.profile.tokens.text_hex] + [c.hex for c in self.profile.tokens.palette[:8]]
        best = max(cands, key=lambda c: contrast_ratio(c, bg), default=None)
        if best is None or contrast_ratio(best, bg) < 4.5:
            best = "FFFFFF" if rel_luminance(bg) < 0.4 else "000000"
        for r in sh._element.iter(qn("a:r")):
            rpr = r.find(qn("a:rPr"))
            if rpr is None:
                rpr = etree.Element(qn("a:rPr"))
                r.insert(0, rpr)
            for old_fill in rpr.findall(qn("a:solidFill")):
                rpr.remove(old_fill)
            fill = etree.Element(qn("a:solidFill"))
            etree.SubElement(fill, qn("a:srgbClr")).set("val", best)
            # порядок схемы: ln?, затем заливки до effects/latin/ea/cs
            ln = rpr.find(qn("a:ln"))
            rpr.insert(list(rpr).index(ln) + 1 if ln is not None else 0, fill)

    def _fit_backdrop(self, slide, backdrop_id: int, sh, slot: Slot, paras: list[str]) -> None:
        """Подгоняет плашку под заголовком под новый текст: те же отступы, что в примере, но не шире
        (свободной) рамки заголовка.
        """
        bd = shape_by_id(slide, backdrop_id)
        if bd is None or bd.width is None:
            return
        rpr = sh._element.find(".//" + qn("a:rPr"))
        size = int(rpr.get("sz")) / 100 if rpr is not None and rpr.get("sz") else (slot.style.size or self.profile.tokens.type_scale.title)
        font = slot.style.font or self.profile.tokens.heading_font
        m = measure(paras, font, size, slot.box.w, slot.style.bold, caps=slot.style.caps, tracking=slot.style.tracking)
        x0, pad = int(bd.left), max(int(sh.left) - int(bd.left), 0)
        right = min(int(sh.left) + m.width_emu + pad, slot.box.r + pad)
        bd.width = max(right - x0, 2 * pad + int(0.05 * self.profile.tokens.slide_w))
        if m.lines > 1:  # у заголовка в несколько строк плашка остаётся под каждой строкой
            bd.height = max(int(bd.height), m.height_emu + 2 * max(int(sh.top) - int(bd.top), 0))

    def _nearest_below(self, num: Slot, pat: Pattern, fills) -> Slot | None:
        """Ближайший свободный текстовый слот под номером элемента (для подписи к номеру)."""
        best, bd = None, None
        for s in pat.slots:
            if s.kind != "text" or s.id in fills or s.role in (SlotRole.title, SlotRole.number) or s.item_index is not None:
                continue
            if s.box.y < num.box.y + num.box.h * 0.5:
                continue
            dx = abs(s.box.cx - num.box.cx) + abs(s.box.x - num.box.x)
            dy = s.box.y - num.box.b
            dist = dx + max(dy, 0) * 2
            if bd is None or dist < bd:
                best, bd = s, dist
        return best

    def _head(self, it: Item) -> str:
        """Заголовок элемента; значение идёт первым, если в паттерне больше негде показать числа."""
        if it.value and it.title and self._inline_values:
            return f"{it.value} {it.title}"
        return it.title or it.value

    def _item_text(self, slot: Slot, it: Item, i: int) -> tuple[list[str], list[int] | None]:
        """Текст элемента для слота (с учётом ролей абзацев составной рамки) и индексы стилей."""
        if slot.para_roles:
            head_role = slot.para_roles[0]
            if head_role == SlotRole.number:
                head = it.value or number_format(slot.sample_text.split("\n")[0], i) or f"{i + 1}"
                body = it.text or it.title
            else:
                head = self._head(it)
                body = it.text
            return ([head, body], [0, 1]) if body else ([head], [0])
        role = slot.role
        if role == SlotRole.number:
            fmt = number_format(slot.sample_text, i)
            return ([it.value or fmt or f"{i + 1}"] if (it.value or fmt) else [it.title], None)
        if role == SlotRole.item_title:
            return ([self._head(it) or it.text], None)
        if role == SlotRole.item_text:
            return ([it.text or it.title or it.value], None)
        if role == SlotRole.label:
            return ([it.value], None) if it.value else ([], None)
        return ([it.text or it.title], None)

    def _drop_empty_plates(self, slide, removed: list[Box]) -> None:
        """Плашка примера, обрамлявшая только удалённые теперь тексты (полоса примечания, выноска), осталась
        бы пустой цветной областью: удаляется и она. Плашки, на которых что-то осталось, сохраняются.
        """
        if not removed:
            return
        t = self.profile.tokens
        area, tol = t.slide_w * t.slide_h, int(0.004 * t.slide_w)
        recs = [r for r in flatten_shapes(slide.shapes) if r.shape.shape_type != MSO_SHAPE_TYPE.GROUP]
        for r in recs:
            sh, b = r.shape, r.box
            if etree.QName(sh._element).localname != "sp" or not is_plate(sh) or not 0.03 * area < b.area < 0.85 * area:
                continue
            if sh.has_text_frame and sh.text_frame.text.strip():
                continue
            if not any(b.contains(x, tol=tol) for x in removed):
                continue
            if any(o is not r and b.contains(o.box, tol=tol) and o.box.area < b.area and draws(o.shape) for o in recs):
                continue
            delete_shape(sh)

    def _brand_zones(self, slide) -> list[Box]:
        """Зоны логотипа, колонтитула и номера страницы, которые рисуют макет и мастер слайда."""
        li = next((l.index for l, lay in zip(self.profile.layouts, self.layouts) if lay is slide.slide_layout), None)
        if li is None:
            return []
        own = {f"layout:{li}", f"master:{self.profile.layouts[li].master_index}"}
        return [b.box for b in self.profile.brand_elements
                if b.kind in ("logo", "footer", "page_number") and own & set(b.source.split(","))]

    def _fit(self, slide, n: int, sh, slot: Slot, paras: list[str], tidx: list[int] | None = None) -> None:
        if not paras or not any(p.strip() for p in paras):
            return
        size = slot.style.size or self.profile.tokens.type_scale.body
        font = slot.style.font or self.profile.tokens.body_font
        # заголовки должны оставаться в своей рамке (они стоят прямо над контентом);
        # другие растущие рамки могут занимать свободное место под собой
        if slot.role in (SlotRole.title, SlotRole.subtitle) or not slot.autofit:
            avail_h = slot.box.h
        else:
            avail_h = max(slot.box.h, int(slot.max_lines * size * 1.2 * EMU_PT + 91440))
        # остальные рамки растут внутри своего блока до уменьшения кегля (разъяснение организаторов)
        room = None
        if slot.role not in (SlotRole.title, SlotRole.subtitle) and not slot.painted:
            room = text_room(slide, sh, self.profile.tokens, owned_h=slot.box.h, keep_clear=self._brand_zones(slide))
        if room is not None:
            avail_h = max(slot.box.h, min(avail_h, room[2] - room[1]) if slot.autofit else room[2] - room[1])
        box_w = slot.box.w
        # интервалы между абзацами и межстрочный интервал занимают свою долю высоты
        styles = paragraph_styles(slide, sh)[1]
        gaps, line = spacing(styles)
        # капс и разрядка шаблона расширяют слова: замер идёт так, как текст нарисует рендерер
        caps = slot.style.caps or any(s.caps for s in styles)
        trk = max([slot.style.tracking] + [s.tracking for s in styles])
        if styles and styles[0].size and not slot.para_roles and abs(styles[0].size - size) > 0.5:
            # рамка рисуется унаследованным кеглем: масштаб автоподбора примера не переносится
            size = styles[0].size
        fit_h = max(int((avail_h - gaps * EMU_PT - 2 * DEFAULT_INSET_TB) / line) + 2 * DEFAULT_INSET_TB, int(avail_h * 0.3))
        if slot.para_roles:
            # рамка с несколькими стилями (крупный лид над основным текстом): абзацы сохраняют роли примера;
            # сначала уступает лид (до 1.2× основного текста), затем всё уменьшается одним множителем
            sizes = [ps.size or size for ps in slot.para_styles] or [size]
            tidx = tidx or [0] + [1] * (len(paras) - 1)
            psizes = [sizes[min(i, len(sizes) - 1)] for i in tidx]
            if len(styles) == len(paras) and all(st.size for st in styles):
                psizes = [st.size for st in styles]  # как рисуется: масштаб автоподбора примера не переносится
            heads = [i for i, t in enumerate(tidx) if t == 0]
            body = min((s for s, t in zip(psizes, tidx) if t != 0), default=0)
            fh = 1.0
            if heads and body and fit_composite(paras, psizes, font, box_w, fit_h, caps=caps, tracking=trk) < 0.99:
                for k in (0.9, 0.8, 0.7, 0.6, 0.5):
                    if psizes[heads[0]] * k < 1.2 * body:
                        break
                    fh = k
                    if fit_composite(paras, [s * (k if t == 0 else 1) for s, t in zip(psizes, tidx)], font, box_w, fit_h,
                                     caps=caps, tracking=trk) >= 0.99:
                        break
                scale_paragraph_sizes(sh, heads, fh, default_size=psizes[heads[0]])
                psizes = [s * (fh if t == 0 else 1) for s, t in zip(psizes, tidx)]
            f = fit_composite(paras, psizes, font, box_w, fit_h, caps=caps, tracking=trk)
            if f < 0.99:
                scale_font_sizes(sh, f, default_size=size)
            if room is not None:
                grow_frame(sh, room, text_height(slide, sh, box_w, font))
            return
        ts = self.profile.tokens.type_scale
        allowed = ts.sizes
        # раздутые кегли плейсхолдеров по умолчанию (например, 32 pt в основном тексте) могут уменьшаться до
        # кегля текста шаблона
        floor_ratio = 0.7 if size <= 1.3 * ts.body else max(ts.body * 0.85 / size, 0.35)
        if slot.role == SlotRole.title:
            floor_ratio = min(floor_ratio, 0.6)
        fitted = fit_font_size(paras, font, size, box_w, fit_h, slot.style.bold, floor_ratio, allowed, caps, trk)
        if fitted is None:
            min_size = snap_down(max(size * floor_ratio, ts.caption), allowed)
            set_font_size(sh, min_size)
            self.report.overflows.append(Overflow(slide=n, shape_id=sh.shape_id, role=slot.role.value, paragraphs=paras,
                                                  budget=max_chars_for(box_w, fit_h, font, min_size, slot.style.bold,
                                                                       caps, trk)))
        elif fitted < size - 0.05:
            set_font_size(sh, fitted)
        if room is not None:
            grow_frame(sh, room, text_height(slide, sh, box_w, font))

    # ---------------------------------------------------------------- компоновка
    def _canvas(self) -> tuple[int, bool]:
        """Макет-холст для композиции варианта: (индекс макета, тёмный ли он)."""
        prof = self.profile
        tone = self.variant.tone
        if tone == "auto":
            tone = "dark" if prof.tokens.dark_background else "light"
        idx = prof.canvas_layouts.get(tone)
        if idx is None:
            idx = next(iter(prof.canvas_layouts.values()), 0)
        return idx, prof.layouts[idx].dark

    def _region_colors(self, layout_info, region: Box | None) -> list[str]:
        """Цвета, которые макет рисует под `region` (несколько — для градиентов и свечений)."""
        if region is None or not layout_info.thumbnail or not Path(layout_info.thumbnail).exists():
            return []
        try:
            return region_colors(Path(layout_info.thumbnail), region, self.profile.tokens.slide_w, self.profile.tokens.slide_h)
        except Exception:
            return []

    def _style_for(self, slide_bg: str, samples: list[str] | None = None) -> C.Style:
        """Стиль композиции для фона слайда."""
        return C.Style(tokens=self.profile.tokens, bg=slide_bg, dark=rel_luminance(slide_bg) < 0.4,
                       bg_samples=tuple(samples or ()))

    def _title_style(self):
        """Преобладающий стиль заголовка в содержательных примерах шаблона (для макетов без плейсхолдера
        заголовка).
        """
        from collections import Counter

        styles = Counter()
        for p in self.profile.usable_patterns():
            if p.kind.value in ("title", "section", "thanks", "quote"):
                continue
            for s in p.slots:
                if s.role == SlotRole.title and s.style.size:
                    styles[(s.style.font, s.style.size, s.style.bold, s.style.color_hex)] += 1
        return styles.most_common(1)[0][0] if styles else (None, None, False, None)

    def _compose_section(self, spec: SlideSpec, plan_title: str | None = None):
        """Разделитель или финальный слайд, если в шаблоне нет такого примера: крупный заголовок на
        макете-«холсте», акцентная линия, подзаголовок — всё в токенах шаблона.
        """
        t = self.profile.tokens
        idx, _ = self._canvas()
        layout_info = self.profile.layouts[idx]
        slide = self._slide_from_layout(self.layouts[idx])
        for ph in list(slide.placeholders):
            delete_shape(ph)
        bg = layout_info.content_bg_hex or layout_info.background_hex or t.background_hex
        m = t.margins
        # разделитель ставится в свободную зону макета, найденную CV (вне фоновой графики), иначе посередине
        # слайда
        region = Box(x=m.left, y=int(0.18 * t.slide_h), w=t.slide_w - m.left - m.right, h=int(0.7 * t.slide_h))
        cb = layout_info.content_box
        if cb is not None and cb.area >= 0.15 * t.slide_w * t.slide_h:
            x0, y0 = max(cb.x, m.left), max(cb.y, m.top)
            x1, y1 = min(cb.r, t.slide_w - m.right), min(cb.b, t.slide_h - m.bottom)
            if x1 - x0 > 0.3 * t.slide_w and y1 - y0 > 0.25 * t.slide_h:
                region = Box(x=x0, y=y0, w=x1 - x0, h=y1 - y0)
        st = self._style_for(bg, self._region_colors(layout_info, region))
        title = plan_title or spec.title
        size = snap_down(min(t.type_scale.title * 1.3, t.type_scale.number), t.type_scale.sizes)
        box = Box(x=region.x, y=region.y + int(0.1 * region.h), w=int(region.w * 0.9), h=int(0.42 * region.h))
        C.add_text(slide, box, [title], st, role="title", size=size, anchor="bottom", bold=True, name="Title")
        rule_y = box.b + int(0.04 * region.h)
        C.add_line(slide, region.x + 45720, rule_y, region.x + int(region.w * 0.18), rule_y, st.accent, 3.0, name="Accent rule")
        sub = spec.subtitle or spec.message
        if sub:
            C.add_text(slide, Box(x=region.x, y=box.b + int(0.09 * region.h), w=int(region.w * 0.85), h=int(0.3 * region.h)),
                       [sub], st, role="subtitle", color=st.muted, name="Subtitle")
        return slide

    def _compose(self, n: int, spec: SlideSpec, kind: str, image: str | None, icons: list[str | None] | None):
        if kind == "section":
            return self._compose_section(spec)
        prof = self.profile
        t = prof.tokens
        idx, dark = self._canvas()
        layout_info = prof.layouts[idx]
        slide = self._slide_from_layout(self.layouts[idx])
        title_ph = None
        for ph in list(slide.placeholders):
            pt = str(ph.placeholder_format.type)
            if "TITLE" in pt and title_ph is None:
                title_ph = ph
            else:
                delete_shape(ph)
        bg = layout_info.content_bg_hex or layout_info.background_hex or t.background_hex
        st = self._style_for(bg)
        region = layout_info.content_box or t.content_box or Box(x=t.margins.left, y=int(0.22 * t.slide_h),
                                                                  w=t.slide_w - t.margins.left - t.margins.right,
                                                                  h=int(0.68 * t.slide_h))
        tb = layout_info.title_box or t.title_box or Box(x=t.margins.left, y=t.margins.top,
                                                         w=t.slide_w - t.margins.left - t.margins.right, h=int(0.13 * t.slide_h))
        allowed = t.type_scale.sizes
        if title_ph is not None:
            title_ph.text_frame.text = spec.title
            tb = Box(x=int(title_ph.left), y=int(title_ph.top), w=int(title_ph.width), h=int(title_ph.height))
            clear = self._clear_title_box(idx, tb)
            if clear is not None:
                title_ph.left, title_ph.top, title_ph.width, title_ph.height = tb.x, tb.y, clear.w, tb.h
                tb = clear
            size, font, bold = t.type_scale.title, t.heading_font, False
            # плейсхолдер заголовка макета может наследовать капс и разрядку
            ps = paragraph_styles(slide, title_ph)[1]
            caps, trk = any(s.caps for s in ps), max([s.tracking for s in ps] or [0.0])
        else:
            caps, trk = False, 0.0
            tb = self._clear_title_box(idx, tb) or tb
            font, size, bold, color = self._title_style()
            font, size = font or t.heading_font, size or t.type_scale.title
        # заголовок должен помещаться в свою рамку: переполненный заголовок наезжает на контент или уходит за
        # слайд
        fitted = fit_font_size([spec.title], font, size, tb.w, tb.h, bold, 0.6, allowed, caps, trk)
        overflow = fitted is None
        if overflow:
            fitted = snap_down(size * 0.6, allowed)
        if title_ph is not None:
            if fitted < size - 0.05:
                set_font_size(title_ph, fitted)
            anchor_bottom = title_ph.text_frame.vertical_anchor == MSO_ANCHOR.BOTTOM
            title_id = title_ph.shape_id
        else:
            tcolor = color if color and C.contrast_ratio(color, bg) >= 3 else st.text
            shp = C.add_text(slide, tb, [spec.title], st, role="title", color=tcolor, bold=bold, size=fitted, font=font,
                             fit=False, name="Title")
            anchor_bottom = False
            title_id = shp.shape_id
        if overflow:  # сокращается шагом LLM (скилл shortener), если модель подключена
            self.report.overflows.append(Overflow(slide=n, shape_id=title_id, role="title", paragraphs=[spec.title],
                                                  budget=max_chars_for(tb.w, tb.h, font, fitted, bold, caps, trk)))
        need = measure([spec.title], font, fitted, tb.w, bold, caps=caps, tracking=trk).height_emu
        title_bottom = tb.b if anchor_bottom else max(tb.b, min(tb.y + need, tb.b + need // 2))
        # контент ниже начинается под последней строкой заголовка: рамка может вырасти туда
        if title_bottom > tb.b:
            title_sh = shape_by_id(slide, title_id)
            if title_sh is not None:
                title_sh.left, title_sh.top, title_sh.width, title_sh.height = tb.x, tb.y, tb.w, title_bottom - tb.y
        m = t.margins
        # область остаётся в полях и под заголовком (возможно, двухстрочным);
        # левый край выровнен по рамке заголовка, чтобы текст начинался от одной направляющей
        on_title_guide = bool(tb and abs(tb.x - region.x) < 0.06 * t.slide_w)
        if on_title_guide:
            region = Box(x=tb.x, y=region.y, w=region.r - tb.x, h=region.h)
        x0 = region.x if on_title_guide else max(region.x, m.left)
        y0 = max(region.y, title_bottom + int(0.025 * t.slide_h))
        x1 = min(region.r, t.slide_w - m.right)
        # низ: над зоной колонтитула / даты / номера макета, с небольшим запасом
        foot = [Box(**ph["box"]) for ph in layout_info.placeholders
                if ph.get("type") in ("FOOTER", "DATE", "SLIDE_NUMBER") and ph.get("box")]
        foot_top = min((b.y for b in foot if b.y > 0.75 * t.slide_h and b.h > 0), default=t.slide_h)
        y1 = min(region.b, t.slide_h - m.bottom, foot_top - int(0.015 * t.slide_h), int(0.94 * t.slide_h))
        if x1 - x0 < 0.3 * t.slide_w:  # свободна лишь узкая полоса: используем ширину между полями
            x0, x1 = m.left, t.slide_w - m.right
        if y1 - y0 < 0.3 * t.slide_h:  # контент никогда не уходит ниже слайда: берём область под заголовком
            y1 = t.slide_h - m.bottom
            y0 = max(title_bottom + int(0.02 * t.slide_h), min(y0, y1 - int(0.3 * t.slide_h)))
        region = Box(x=x0, y=y0, w=max(x1 - x0, 1), h=max(y1 - y0, 1))
        st = self._style_for(bg, self._region_colors(layout_info, region))
        if layout_info.textured:
            # пёстрый фон: спокойная подложка в цветах шаблона сохраняет читаемость текста
            pad = int(0.02 * t.slide_w)
            plate_color = C.mix(t.background_hex if rel_luminance(t.background_hex) > 0.5 else "FFFFFF", bg, 0.08) \
                if rel_luminance(bg) > 0.35 else C.mix("000000", bg, 0.15)
            C.add_rect(slide, Box(x=region.x - pad, y=region.y - pad, w=region.w + 2 * pad, h=region.h + 2 * pad), plate_color,
                       rounded=True, name="Plate")
            st = self._style_for(plate_color)
        lead = spec.subtitle or (spec.message if kind != "bullets" else "")
        if lead:  # по размеру текста (кеглем основного текста), от 12% до 30% области
            need = measure([lead], t.body_font, t.type_scale.body, region.w).height_emu
            lh = min(max(need, int(region.h * 0.12)), int(region.h * 0.3))
            C.add_text(slide, Box(x=region.x, y=region.y, w=region.w, h=lh), [lead], st, role="body", color=st.muted,
                       name="Lead")
            region = Box(x=region.x, y=region.y + lh + int(region.h * 0.03), w=region.w, h=region.h - lh - int(region.h * 0.03))
        items = spec.items or [Item(title=b) for b in spec.bullets]
        if kind == "chart" and spec.chart:
            C.draw_chart(slide, region, spec.chart, st)
        elif kind == "table" and (spec.table or spec.chart):
            table = spec.table or _chart_as_table(spec)
            C.draw_table(slide, region, table, st)
        elif kind == "kpi":
            kitems = items
            if spec.chart and not spec.items:
                s0 = spec.chart.series[0]
                kitems = [Item(value=_fmt_num(v, spec.chart.unit), text=c) for c, v in zip(spec.chart.categories, s0.values)]
            C.draw_kpis(slide, region, kitems, st)
        elif kind == "process" and items:
            C.draw_process(slide, region, items, st)
        elif kind == "cards" and items:
            C.draw_cards(slide, region, items, st, icons)
        elif kind == "image":
            C.draw_image_text(slide, region, image, spec.bullets or [spec.message], st)
        elif kind == "quote":
            C.draw_quote(slide, region, spec.quote or spec.message, spec.quote_author, st)
        else:
            bullets = spec.bullets or [_item_line(i) for i in spec.items] or [spec.message]
            C.draw_bullets(slide, region, bullets, st, self.variant.max_bullets)
        return slide


def _item_line(it: Item) -> str:
    """Один пункт на элемент: сначала значение (KPI не должен терять число), затем «заголовок: текст»."""
    head = " ".join(x for x in (it.value, it.title) if x)
    return f"{head}: {it.text}" if head and it.text else (head or it.text)


def _fmt_num(v: float, unit: str) -> str:
    """Число с разделителями разрядов и единицей измерения."""
    s = f"{v:,.0f}".replace(",", " ") if abs(v) >= 100 or float(v).is_integer() else f"{v:.1f}".replace(".", ",")
    return f"{s}{unit if unit in ('%', '₽', '$', '€') else (' ' + unit if unit else '')}"


def _chart_as_table(spec: SlideSpec):
    """Данные диаграммы в виде таблицы."""
    from decksmith.core.models import TableSpec

    ch = spec.chart
    cols = [ch.x_title or ""] + [s.name + (f", {ch.unit}" if ch.unit else "") for s in ch.series]
    rows = [[c] + [_fmt_num(s.values[i], "") if i < len(s.values) else "" for s in ch.series] for i, c in enumerate(ch.categories)]
    return TableSpec(columns=cols, rows=rows)
