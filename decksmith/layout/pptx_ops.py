"""Низкоуровневые операции над объектами python-pptx с сохранением форматирования."""
from __future__ import annotations

import copy
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx.oxml.ns import qn

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_SKIP_RELS = ("notesSlide", "slideLayout", "comments")


# ----------------------------------------------------------------------------
# Слайды
# ----------------------------------------------------------------------------
def _remap_rids(root, rid_map: dict[str, str]) -> None:
    """Переназначает ссылки r:id в XML скопированного слайда."""
    for el in root.iter():
        for attr in list(el.attrib):
            if attr.startswith(f"{{{R_NS}}}") and el.get(attr) in rid_map:
                el.set(attr, rid_map[el.get(attr)])


def copy_chart_part(src_part, package):
    """Независимая копия части диаграммы (+ встроенная книга Excel), чтобы у каждого клона были свои данные.
    """
    from pptx.parts.chart import ChartPart
    from pptx.parts.embeddedpackage import EmbeddedXlsxPart

    new = ChartPart.load(package.next_partname(ChartPart.partname_template), src_part.content_type, package, src_part.blob)
    rid_map = {}
    for rid, rel in src_part.rels.items():
        if rel.is_external:
            rid_map[rid] = new.relate_to(rel.target_ref, rel.reltype, is_external=True)
        elif rel.reltype.endswith("/package"):
            xlsx = EmbeddedXlsxPart.new(rel.target_part.blob, package)
            rid_map[rid] = new.relate_to(xlsx, rel.reltype)
        else:
            rid_map[rid] = new.relate_to(rel.target_part, rel.reltype)
    _remap_rids(new._element, rid_map)
    return new


def duplicate_slide(prs, src):
    """Добавляет глубокую копию `src` (фигуры, фон, изображения, медиа, ссылки, собственные части диаграмм).
    """
    new = prs.slides.add_slide(src.slide_layout)
    for shp in list(new.shapes):
        shp._element.getparent().remove(shp._element)
    rid_map: dict[str, str] = {}
    for rid, rel in src.part.rels.items():
        if any(s in rel.reltype for s in _SKIP_RELS):
            continue
        if rel.is_external:
            rid_map[rid] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        elif rel.reltype.endswith("/chart"):
            rid_map[rid] = new.part.relate_to(copy_chart_part(rel.target_part, prs.part.package), rel.reltype)
        else:
            rid_map[rid] = new.part.relate_to(rel.target_part, rel.reltype)
    src_csld = src._element.find(qn("p:cSld"))
    new_csld = new._element.find(qn("p:cSld"))
    bg = src_csld.find(qn("p:bg"))
    if bg is not None:
        new_csld.insert(0, copy.deepcopy(bg))
    tree = new.shapes._spTree
    for el in src.shapes._spTree.iterchildren():
        tag = etree.QName(el).localname
        if tag in ("nvGrpSpPr", "grpSpPr"):
            continue
        tree.append(copy.deepcopy(el))
    _remap_rids(new._element, rid_map)
    return new


def delete_slides(prs, indices: list[int]) -> None:
    """Удаляет слайды по индексам."""
    lst = prs.slides._sldIdLst
    ids = list(lst)
    for i in sorted(set(indices), reverse=True):
        sld = ids[i]
        prs.part.drop_rel(sld.rId)
        lst.remove(sld)


# ----------------------------------------------------------------------------
# Фигуры
# ----------------------------------------------------------------------------
def iter_all_shapes(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:  # ГРУППА
            yield from iter_all_shapes(sh.shapes)


def shape_by_id(slide, shape_id: int):
    """Фигура по id (включая фигуры внутри групп)."""
    for sh in iter_all_shapes(slide.shapes):
        if sh.shape_id == shape_id:
            return sh
    return None


def delete_shape(shape) -> None:
    el = shape._element
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)
        # удаляем опустевшие группы
        if etree.QName(parent).localname == "grpSp":
            children = [c for c in parent if etree.QName(c).localname not in ("nvGrpSpPr", "grpSpPr")]
            if not children and parent.getparent() is not None:
                parent.getparent().remove(parent)


def move_shape(shape, dx: int, dy: int) -> None:
    """Сдвигает фигуру."""
    shape.left = int(shape.left + dx)
    shape.top = int(shape.top + dy)


# ----------------------------------------------------------------------------
# Текст (с сохранением форматирования)
# ----------------------------------------------------------------------------
def _clean_paragraph_copy(p_el, text: str):
    p = copy.deepcopy(p_el)
    runs = p.findall(qn("a:r"))
    for child in list(p):
        tag = etree.QName(child).localname
        if tag in ("br", "fld") or (tag == "r" and runs and child is not runs[0]):
            p.remove(child)
    runs = p.findall(qn("a:r"))
    if runs:
        r = runs[0]
    else:
        r = etree.SubElement(p, qn("a:r"))
        end = p.find(qn("a:endParaRPr"))
        rpr = copy.deepcopy(end) if end is not None else etree.Element(qn("a:rPr"))
        rpr.tag = qn("a:rPr")
        r.append(rpr)
        etree.SubElement(r, qn("a:t"))
        if end is not None:
            p.remove(end)
            p.append(end)
    t = r.find(qn("a:t"))
    if t is None:
        t = etree.SubElement(r, qn("a:t"))
    t.text = text
    # endParaRPr должен оставаться последним дочерним элементом
    end = p.find(qn("a:endParaRPr"))
    if end is not None:
        p.remove(end)
        p.append(end)
    return p


def set_paragraphs(shape, paragraphs: list[str], template_index: list[int] | None = None) -> None:
    """Заменяет текст, сохраняя pPr/rPr абзацев шаблона.

        template_index[i] = индекс исходного абзаца, форматирование которого копирует i-й
        новый абзац (составные рамки: стиль заголовка для 0, стиль текста для 1..).
    """
    txb = shape._element.find(qn("p:txBody"))
    if txb is None:
        shape.text_frame.text = "\n".join(paragraphs)
        return
    set_txbody_paragraphs(txb, paragraphs, template_index)


def set_txbody_paragraphs(txb, paragraphs: list[str], template_index: list[int] | None = None) -> None:
    """Заменяет абзацы txBody, сохраняя форматирование абзацев-образцов."""
    src_ps = txb.findall(qn("a:p"))
    non_empty = [p for p in src_ps if "".join(t.text or "" for t in p.iter(qn("a:t"))).strip()] or src_ps
    if not non_empty:
        non_empty = [etree.SubElement(txb, qn("a:p"))]
    new_ps = []
    for i, text in enumerate(paragraphs):
        ti = template_index[i] if template_index and i < len(template_index) else min(i, len(non_empty) - 1)
        ti = max(0, min(ti, len(non_empty) - 1))
        new_ps.append(_clean_paragraph_copy(non_empty[ti], text))
    for p in src_ps:
        txb.remove(p)
    for p in new_ps:
        txb.append(p)


def set_font_size(shape, size_pt: float) -> None:
    """Задаёт один кегль всем фрагментам фигуры."""
    sz = str(int(round(size_pt * 100)))
    txb = shape._element.find(qn("p:txBody"))
    if txb is None:
        return
    for r in txb.iter(qn("a:r")):
        rpr = r.find(qn("a:rPr"))
        if rpr is None:
            rpr = etree.Element(qn("a:rPr"))
            r.insert(0, rpr)
        rpr.set("sz", sz)
    for e in txb.iter(qn("a:endParaRPr")):
        e.set("sz", sz)


def scale_font_sizes(shape, factor: float, default_size: float) -> None:
    """Масштабирует кегли всех фрагментов фигуры."""
    txb = shape._element.find(qn("p:txBody"))
    if txb is None:
        return
    for r in txb.iter(qn("a:r")):
        rpr = r.find(qn("a:rPr"))
        if rpr is None:
            rpr = etree.Element(qn("a:rPr"))
            r.insert(0, rpr)
        cur = int(rpr.get("sz")) / 100 if rpr.get("sz") else default_size
        rpr.set("sz", str(int(round(cur * factor * 100))))


def scale_paragraph_sizes(shape, indices: list[int], factor: float, default_size: float) -> None:
    """Масштабирует только часть абзацев (вводная фраза уступает место основному тексту)."""
    txb = shape._element.find(qn("p:txBody"))
    if txb is None or factor >= 0.999:
        return
    paras = txb.findall(qn("a:p"))
    for i in indices:
        if i >= len(paras):
            continue
        for el in list(paras[i].iter(qn("a:rPr"))) + list(paras[i].iter(qn("a:endParaRPr"))):
            cur = int(el.get("sz")) / 100 if el.get("sz") else default_size
            el.set("sz", str(int(round(cur * factor * 100))))


# ----------------------------------------------------------------------------
# Нативные таблицы / диаграммы из примеров шаблона
# ----------------------------------------------------------------------------
def table_has_merges(gf) -> bool:
    """В таблице есть объединённые ячейки."""
    tbl = gf._element.graphic.graphicData.tbl
    return any(tc.get(a) for tc in tbl.iter(qn("a:tc")) for a in ("gridSpan", "rowSpan", "hMerge", "vMerge"))


def fill_table(gf, columns: list[str], rows: list[list[str]], max_bottom: int | None = None) -> None:
    """Перезаполняет таблицу шаблона с сохранением стиля: строки/колонки добавляются клонированием последней
    строки/колонки (полосы и форматирование ячеек сохраняются), удаляются с конца.
    """
    tbl = gf._element.graphic.graphicData.tbl
    grid = tbl.find(qn("a:tblGrid"))
    total_w = sum(int(gc.get("w")) for gc in grid.findall(qn("a:gridCol")))
    n_cols, n_rows = len(columns), len(rows) + 1
    # колонки
    while len(grid.findall(qn("a:gridCol"))) > n_cols:
        grid.remove(grid.findall(qn("a:gridCol"))[-1])
        for tr in tbl.findall(qn("a:tr")):
            tr.remove(tr.findall(qn("a:tc"))[-1])
    while len(grid.findall(qn("a:gridCol"))) < n_cols:
        grid.append(copy.deepcopy(grid.findall(qn("a:gridCol"))[-1]))
        for tr in tbl.findall(qn("a:tr")):
            tr.append(copy.deepcopy(tr.findall(qn("a:tc"))[-1]))
    lens = [max([len(str(columns[c]))] + [len(str(r[c])) if c < len(r) else 0 for r in rows]) for c in range(n_cols)]
    weights = [max(l, 4) ** 0.8 for l in lens]
    for gc, w in zip(grid.findall(qn("a:gridCol")), weights):
        gc.set("w", str(int(total_w * w / sum(weights))))
    # строки (сохраняем шапку и строки тела; клонируем образец строки тела)
    trs = tbl.findall(qn("a:tr"))
    while len(trs) > n_rows:
        tbl.remove(trs[-1])
        trs = tbl.findall(qn("a:tr"))
    while len(trs) < n_rows:
        proto = trs[-2] if len(trs) >= 3 else trs[-1]  # чередование полос, когда возможно
        tbl.append(copy.deepcopy(proto))
        trs = tbl.findall(qn("a:tr"))
    # таблица остаётся внутри слайда: при необходимости сжимаем высоты строк
    heights = [int(tr.get("h", "0")) for tr in trs]
    if max_bottom is not None and sum(heights) and gf.top + sum(heights) > max_bottom:
        k = max((max_bottom - gf.top) / sum(heights), 0.5)
        for tr, h in zip(trs, heights):
            tr.set("h", str(int(h * k)))
    for r, tr in enumerate(trs):
        values = columns if r == 0 else [str(v) for v in rows[r - 1]]
        for c, tc in enumerate(tr.findall(qn("a:tc"))):
            txb = tc.find(qn("a:txBody"))
            if txb is None:
                continue
            set_txbody_paragraphs(txb, [str(values[c]) if c < len(values) else ""])
    gf.height = int(sum(int(tr.get("h", "0")) for tr in trs)) or gf.height


def fit_table_text(gf, max_bottom: int, font: str, min_size: float = 10.0, default_size: float = 18.0) -> bool:
    """После перезаполнения: один коэффициент масштаба для текста всех ячеек, при котором каждое слово
    помещается в свою колонку, а строки, выросшие под текст, заканчиваются выше `max_bottom`; строки
    получают высоты, которые нарисуют рендереры. False, если не помещается даже наименьший читаемый кегль
    (тогда вызывающий код перерисовывает таблицу).
    """
    from decksmith.layout.textfit import measure

    tbl = gf._element.graphic.graphicData.tbl
    widths = [int(gc.get("w", "0")) for gc in tbl.find(qn("a:tblGrid")).findall(qn("a:gridCol"))]
    rows = []
    for tr in tbl.findall(qn("a:tr")):
        cells = []
        for c, tc in enumerate(tr.findall(qn("a:tc"))):
            txb = tc.find(qn("a:txBody"))
            if txb is None or c >= len(widths):
                continue
            rpr = txb.find(".//" + qn("a:rPr"))
            latin = rpr.find(qn("a:latin")) if rpr is not None else None
            pr = tc.find(qn("a:tcPr"))
            mar = [int(pr.get(k, d)) if pr is not None else d for k, d in (("marL", 91440), ("marR", 91440), ("marT", 45720), ("marB", 45720))]
            cells.append(("".join(t.text or "" for t in txb.iter(qn("a:t"))),
                          int(rpr.get("sz")) / 100 if rpr is not None and rpr.get("sz") else default_size,
                          rpr is not None and rpr.get("b") in ("1", "true"),
                          (latin.get("typeface") if latin is not None and not latin.get("typeface", "+").startswith("+") else font),
                          widths[c], mar))
        rows.append((tr, int(tr.get("h", "0")), cells))
    if not rows:
        return True
    avail = max_bottom - int(gf.top)
    k = min(1.0, avail / (sum(h for _, h, _ in rows) or 1))  # высоты строк шаблона — лишь минимальные
    big = max((c[1] for _, _, cells in rows for c in cells), default=default_size)
    f = 1.0
    while True:
        heights, ok = [], True
        for _, h0, cells in rows:
            need = int(h0 * k)
            for text, size, bold, face, w, (ml, mr, mt, mb) in cells:
                if not text.strip():
                    continue
                m = measure([text], face, size * f, w - ml - mr + 2 * 91440, bold, inset_tb=(mt + mb) // 2)
                ok = ok and m.longest_word_emu <= w - ml - mr
                need = max(need, m.height_emu)
            heights.append(need)
        if ok and sum(heights) <= avail:
            break
        f -= 0.05
        if big * f < min_size - 0.01:
            return False
    if f < 0.999:
        for tr, _, _ in rows:
            for el in list(tr.iter(qn("a:rPr"))) + list(tr.iter(qn("a:endParaRPr"))):
                size = int(el.get("sz")) / 100 if el.get("sz") else default_size
                el.set("sz", str(int(round(size * f * 2)) * 50))
    for (tr, _, _), h in zip(rows, heights):
        tr.set("h", str(h))
    gf.height = sum(heights)
    return True


def fill_chart(gf, categories: list[str], series: list[tuple[str, list[float]]], unit: str = "") -> None:
    """Заменяет данные диаграммы на месте: оформление диаграммы шаблона (цвета, шрифты, оси) сохраняется."""
    from pptx.chart.data import CategoryChartData

    cd = CategoryChartData()
    cd.categories = categories
    for name, values in series[:5]:
        cd.add_series(name, [float(v) for v in values])
    chart = gf.chart
    chart.replace_data(cd)
    plot = chart.plots[0]
    xml = chart._chartSpace.xml
    has_axis_title = "<c:valAx>" in xml and "<c:title>" in xml.split("<c:valAx>")[-1]
    if not has_axis_title:
        plot.has_data_labels = True  # значения читаются и без названия оси
        try:
            plot.data_labels.show_value = True
        except Exception:
            pass
    if len(series) > 1 or "pieChart" in xml or "doughnutChart" in xml:
        chart.has_legend = True


# ----------------------------------------------------------------------------
# Изображения
# ----------------------------------------------------------------------------
def replace_picture(slide, shape, image_path: str | Path, mode: str = "cover") -> None:
    """Меняет картинку в фигуре-изображении, сохраняя рамку; обрезает, чтобы избежать искажений."""
    image_part, rid = slide.part.get_or_add_image_part(str(image_path))
    blip = shape._element.find(".//" + qn("a:blip"))
    if blip is None:
        return
    blip.set(qn("r:embed"), rid)
    with Image.open(image_path) as im:
        iw, ih = im.size
    fw, fh = shape.width, shape.height
    blipfill = shape._element.find(".//" + qn("p:blipFill"))
    if blipfill is None:
        return
    for sr in blipfill.findall(qn("a:srcRect")):
        blipfill.remove(sr)
    if mode == "cover" and fw and fh:
        img_ratio, box_ratio = iw / ih, fw / fh
        src = etree.Element(qn("a:srcRect"))
        if img_ratio > box_ratio:  # обрезка слева/справа
            crop = (1 - box_ratio / img_ratio) / 2
            src.set("l", str(int(crop * 100000)))
            src.set("r", str(int(crop * 100000)))
        elif img_ratio < box_ratio:
            crop = (1 - img_ratio / box_ratio) / 2
            src.set("t", str(int(crop * 100000)))
            src.set("b", str(int(crop * 100000)))
        idx = list(blipfill).index(blip) + 1
        blipfill.insert(idx, src)
