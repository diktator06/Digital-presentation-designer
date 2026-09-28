"""Сквозная оркестрация.

  до Enter (с кэшем, без лимита времени):   analyze_template, ingest_content
  после Enter (≤ 5 мин на 3 колоды):        план -> изображения/иконки -> 3 варианта параллельно:
        выбор -> сборка -> сокращение -> рендер -> аудит -> автоисправление -> повторный аудит -> экспорт

Каждый запуск пишет workspace/runs/<run_id>/manifest.json с замерами времени, версиями
скиллов (+sha), телеметрией модели и обоснованием вёрстки каждого слайда.
"""
from __future__ import annotations

import asyncio
import json
import logging
import shutil
import time
import uuid
from dataclasses import dataclass, field
from pathlib import Path
from typing import Awaitable, Callable

from decksmith.audit.context import AuditContext
from decksmith.audit.engine import run_audit
from decksmith.audit.fixers import AUTO_SAFE, apply_fixes
from decksmith.content.ingest import ingest
from decksmith.core.config import settings
from decksmith.core.models import AuditReport, ContentCorpus, DeckPlan, Item, PatternKind, SlideSpec, TemplateProfile
from decksmith.export.exporters import to_html
from decksmith.generation.images import generate_images
from decksmith.generation.llm import LLMClient, Telemetry
from decksmith.generation.planner import Brief, make_plan
from decksmith.generation.skills import Agent, load_agent, load_skill, skills_manifest
from decksmith.layout.builder import DeckBuilder
from decksmith.layout.icons import icon_for
from decksmith.layout.pptx_ops import set_paragraphs, shape_by_id
from decksmith.layout.selector import Variant, load_variants, select_layouts
from decksmith.parsing.template_parser import analyze_template
from decksmith.render.soffice import RenderError, pdf_to_pngs, pptx_to_pdf

log = logging.getLogger(__name__)
Emit = Callable[[dict], Awaitable[None] | None]


@dataclass
class VariantResult:
    name: str
    title: str
    pptx: str = ""
    pdf: str = ""
    html: str = ""
    pngs: list[str] = field(default_factory=list)
    audit: AuditReport | None = None
    audit_before_fix: dict | None = None
    plan: DeckPlan | None = None  # план варианта, по которому идёт аудит (и повторный аудит после правок)
    slides: list[dict] = field(default_factory=list)
    timings: dict[str, float] = field(default_factory=dict)
    error: str = ""


@dataclass
class RunResult:
    run_id: str
    dir: str
    plan: DeckPlan
    plan_mode: str
    variants: list[VariantResult]
    timings: dict[str, float]
    manifest: str


async def _emit(cb: Emit | None, **ev) -> None:
    """Отправляет событие прогресса в колбэк (синхронный или асинхронный)."""
    if cb is None:
        return
    r = cb(ev)
    if asyncio.iscoroutine(r):
        await r


# ----------------------------------------------------------------------------
# Преобразования плана на уровне варианта (детерминированные)
# ----------------------------------------------------------------------------
def _bullet_item(b: str) -> Item:
    """«Заголовок: подробности» становится карточкой с заголовком; слишком длинный для карточки заголовок не
    обрезается посреди слова — весь пункт остаётся текстом карточки.
    """
    head, sep, rest = b.partition(":")
    if sep and rest.strip() and len(head.strip()) <= 40:
        return Item(title=head.strip(), text=rest.strip(), icon=b)
    return Item(title="", text=b.strip(), icon=b)


_CONTENT = {PatternKind.text, PatternKind.cards, PatternKind.steps, PatternKind.stats, PatternKind.chart, PatternKind.table,
            PatternKind.image_text, PatternKind.two_column}


def _section_digest(sec: SlideSpec, rest: list[SlideSpec]) -> SlideSpec:
    """Разделитель аналитического варианта становится итогом своего раздела: его вводная фраза и выводы
    (заголовки) его слайдов. Разделитель, которому нечего подытожить, остаётся разделителем.
    """
    body = []
    for s in rest:
        if s.intent in (PatternKind.section, PatternKind.thanks, PatternKind.contacts):
            break
        if s.intent in _CONTENT:
            body.append(s.title)
    lead = sec.message or sec.subtitle
    bullets = ([lead] if lead else []) + body[:5]
    if len(bullets) < 2:
        return sec
    return sec.model_copy(deep=True, update={"intent": PatternKind.text, "bullets": bullets, "subtitle": "", "message": ""})


def apply_variant(plan: DeckPlan, v: Variant) -> DeckPlan:
    p = plan.model_copy(deep=True)
    slides = p.slides
    conclusions = [s.title for s in slides if s.intent in _CONTENT]
    summary_first = False
    if v.exec_summary_first:
        summ = next((s for s in slides if "[summary]" in (s.notes or "") or s.notes == "summary"), None)
        agenda = next((s for s in slides if s.intent == PatternKind.agenda), None)
        if summ is None and agenda is not None:
            # в плане нет итогового слайда: содержание становится выводами в начале
            concl = conclusions[: v.max_bullets]
            if len(concl) >= 2:
                agenda.intent, agenda.items, agenda.bullets = PatternKind.text, [], concl
                agenda.title = "Главное" if p.language == "ru" else "Key takeaways"
                agenda.notes = (agenda.notes or "") + "\n[summary]"
                summ = agenda
        if summ is not None:
            slides.remove(summ)
            pos = 2 if len(slides) > 2 and slides[1].intent == PatternKind.agenda else 1
            slides.insert(pos, summ)
            summary_first = True
    if v.drop_sections and not summary_first:
        # в этом варианте нет разделителей, но заданное пользователем число слайдов сохраняется (ТЗ):
        # разделитель подводит итог своего раздела (если выводы уже стоят в начале, он остаётся обычным
        # разделителем)
        slides = [_section_digest(s, slides[i + 1:]) if s.intent == PatternKind.section else s for i, s in enumerate(slides)]
    if v.prefer_icons:
        for s in slides:  # короткие списки становятся карточками с иконками
            if s.intent == PatternKind.text and 3 <= len(s.bullets) <= 5 and all(len(b.split()) <= 14 for b in s.bullets) and not s.items:
                s.items = [_bullet_item(b) for b in s.bullets]
                s.intent = PatternKind.cards
    for s in slides:
        s.bullets = s.bullets[: v.max_bullets]
    p.slides = slides
    return p


def apply_rewrites(plan: DeckPlan, rewrites: list[tuple[list[str], list[str]]]) -> DeckPlan:
    """План с текстами после сокращения: аудит сверяет содержание слайда с тем, что на нём стоит."""
    if not rewrites:
        return plan
    table: dict[str, str] = {}
    for old, new in rewrites:
        # абзацы сопоставляются по порядку; если модель изменила их число, старый абзац получает весь новый текст
        for i, o in enumerate(old):
            table[o.strip()] = new[i] if len(old) == len(new) else " ".join(new)

    def sub(t: str) -> str:
        return table.get(t.strip(), t) if t else t

    out = plan.model_copy(deep=True)
    for spec in out.slides:
        spec.title, spec.subtitle, spec.message, spec.quote = sub(spec.title), sub(spec.subtitle), sub(spec.message), sub(spec.quote)
        spec.bullets = [sub(b) for b in spec.bullets]
        for it in spec.items:
            it.title, it.text = sub(it.title), sub(it.text)
    return out


def plan_variants(plan: DeckPlan, profile: TemplateProfile, names: list[str],
                  images: dict[str, str] | None = None) -> dict[str, tuple[DeckPlan, list]]:
    """Планы вариантов + решения по вёрстке выбираются по очереди, чтобы каждый вариант избегал паттернов,
    которые предыдущие варианты использовали для того же слайда (видимое различие).
    """
    all_v = load_variants()
    avoid: dict[str, set[str]] = {}
    out = {}
    for n in names:
        v = all_v[n]
        vplan = apply_variant(plan, v)
        vimages = set(list((images or {}).keys())[: v.images]) if v.images else set()
        decisions = select_layouts(vplan, profile, v, avoid, vimages)
        for d in decisions:
            avoid.setdefault(d.spec_id, set()).add(d.pattern_id if d.mode == "clone" else f"compose:{d.compose_kind}")
        out[n] = (vplan, decisions)
    return out


def icons_for_plan(plan: DeckPlan, color: str) -> dict[str, list[str | None]]:
    """Иконки для элементов слайдов плана в цвете шаблона."""
    out = {}
    for s in plan.slides:
        if s.items:
            out[s.id] = [icon_for(it.icon or it.title or it.text, color) for it in s.items]
    return out


# ----------------------------------------------------------------------------
class Pipeline:
    def __init__(self, agent: Agent | None = None, transport=None):
        self.cfg = settings()
        self.agent = agent or load_agent()
        self.telemetry = Telemetry()
        # собственный HTTP-транспорт для endpoint модели (тесты: эмулятор в том же процессе)
        self.transport = transport

    # ---- до Enter ------------------------------------------------------------
    def prepare_template(self, path: str | Path, name: str | None = None) -> TemplateProfile:
        """Разбор шаблона (с кэшем)."""
        return analyze_template(path, name=name)

    def prepare_content(self, files: list[str | Path], extra_text: str = "") -> ContentCorpus:
        """Приём контент-пакета."""
        return ingest(files, extra_text)

    # ---- после Enter ---------------------------------------------------------
    async def run(self, profile: TemplateProfile, corpus: ContentCorpus, brief: Brief, variants: list[str] | None = None,
                  emit: Emit | None = None, run_id: str | None = None, plan: DeckPlan | None = None) -> RunResult:
        """Бриф + контент -> варианты. Готовый `plan` (например, утверждённая структура) пропускает шаг
        планирования.
        """
        t0 = time.time()
        run_id = run_id or time.strftime("%Y%m%d-%H%M%S-") + uuid.uuid4().hex[:6]
        rdir = self.cfg.workspace / "runs" / run_id
        rdir.mkdir(parents=True, exist_ok=True)
        timings: dict[str, float] = {}
        all_variants = load_variants()
        names = variants or self.cfg.pipeline.variants
        await _emit(emit, stage="plan", status="start", t=0)
        async with LLMClient(self.cfg.llm, telemetry=self.telemetry, transport=self.transport) as llm:
            if plan is None:
                plan, mode = await make_plan(brief, corpus, profile, llm, self.agent)
            else:
                mode = "given"
            timings["plan"] = round(time.time() - t0, 2)
            (rdir / "plan.json").write_text(plan.model_dump_json(indent=1), encoding="utf-8")
            await _emit(emit, stage="plan", status="done", t=timings["plan"], mode=mode, slides=len(plan.slides),
                        titles=[s.title for s in plan.slides], materials=plan.materials_fit)

            t1 = time.time()
            palette_desc = f"#{profile.tokens.accent_hex}"
            img_task = asyncio.create_task(generate_images(plan, rdir / "images", palette_desc,
                                                           max(all_variants[n].images for n in names) if names else 0))
            icons = await asyncio.to_thread(icons_for_plan, plan, profile.tokens.accent_hex)
            images = await img_task
            timings["assets"] = round(time.time() - t1, 2)
            await _emit(emit, stage="assets", status="done", images=len(images), t=round(time.time() - t0, 2))

            chosen = plan_variants(plan, profile, names, images)  # детерминированно, миллисекунды
            results = await asyncio.gather(*[
                self._variant(all_variants[n], chosen[n], profile, corpus, brief, images, icons, rdir, llm, emit, t0) for n in names
            ])
        timings["total"] = round(time.time() - t0, 2)
        try:
            manifest = str(self._manifest(run_id, rdir, profile, corpus, brief, plan, mode, results, timings))
        except Exception as e:
            # колоды уже готовы: сбой записи манифеста не должен превращать запуск в ошибку
            log.exception("manifest failed: %s", e)
            manifest = ""
        await _emit(emit, stage="done", status="done", t=timings["total"], run_id=run_id)
        return RunResult(run_id=run_id, dir=str(rdir), plan=plan, plan_mode=mode, variants=list(results), timings=timings,
                         manifest=manifest)

    async def _variant(self, v: Variant, chosen: tuple[DeckPlan, list], profile: TemplateProfile, corpus: ContentCorpus,
                       brief: Brief, images: dict[str, str], icons: dict, rdir: Path, llm: LLMClient, emit: Emit | None,
                       t0: float) -> VariantResult:
        vr = VariantResult(name=v.name, title=v.title)
        vdir = rdir / v.name
        vdir.mkdir(parents=True, exist_ok=True)
        try:
            ts = time.time()
            vplan, decisions = chosen
            vimages = dict(list(images.items())[: v.images]) if v.images else {}
            builder = DeckBuilder(profile, v)
            report = builder.build(vplan, decisions, vimages, icons)
            pptx = builder.save(vdir / "deck_v1.pptx")
            vr.slides = report.slides
            vr.timings["build"] = round(time.time() - ts, 2)
            await _emit(emit, stage="build", variant=v.name, status="done", t=round(time.time() - t0, 2), slides=len(report.slides))

            if report.overflows and llm.enabled and self.agent.enabled("shorten"):
                ts = time.time()
                # аудит сверяет содержание с тем, что стоит на слайде после сокращения
                vplan = apply_rewrites(vplan, await self._shorten(pptx, report.overflows, llm))
                vr.timings["shorten"] = round(time.time() - ts, 2)

            ts = time.time()
            pngs, pdf, render_ok = await self._render(pptx, vdir / "render_v1")
            vr.timings["render"] = round(time.time() - ts, 2)
            await _emit(emit, stage="render", variant=v.name, status="done", t=round(time.time() - t0, 2))

            ts = time.time()
            ctx = AuditContext(pptx=pptx, profile=profile, plan=vplan, corpus=corpus, pngs=pngs, brief_text=brief.text,
                               slide_kinds=[s["kind"] for s in report.slides], slide_sources=[s["source"] for s in report.slides],
                               render_ok=render_ok)
            # после аудита остаются автоисправления, повторный рендер и экспорт: им оставляется запас
            audit = await run_audit(ctx, llm, self.agent, visual=self.cfg.pipeline.visual_audit and self.agent.enabled("audit_visual"),
                                    deadline=t0 + self.cfg.pipeline.deadline_s - self.cfg.pipeline.finish_reserve_s)
            vr.timings["audit"] = round(time.time() - ts, 2)
            await _emit(emit, stage="audit", variant=v.name, status="done", score=audit.score, issues=len(audit.issues),
                        t=round(time.time() - t0, 2))

            # детерминированный безопасный раунд автоисправлений (остальное пользователь может просмотреть и
            # выбрать)
            safe = [i for i in audit.issues if i.deterministic and i.fix in AUTO_SAFE and i.severity.value != "info"]
            if safe and self.cfg.pipeline.auto_fix and self.agent.enabled("autofix"):
                ts = time.time()
                vr.audit_before_fix = {"score": audit.score, "issues": len(audit.issues)}
                fixed = vdir / "deck_v2.pptx"
                await apply_fixes(pptx, safe, profile, fixed)
                pptx = fixed
                pngs, pdf, render_ok = await self._render(pptx, vdir / "render_v2")
                ctx = AuditContext(pptx=pptx, profile=profile, plan=vplan, corpus=corpus, pngs=pngs, brief_text=brief.text,
                                   slide_kinds=[s["kind"] for s in report.slides], slide_sources=[s["source"] for s in report.slides],
                                   render_ok=render_ok)
                visual_issues = [i for i in audit.issues if not i.deterministic]
                audit = await run_audit(ctx, None, None, visual=False)
                # контекстные замечания хранятся, пока пользователь по ним не решит
                audit.issues += visual_issues
                vr.timings["autofix"] = round(time.time() - ts, 2)
                await _emit(emit, stage="autofix", variant=v.name, status="done", score=audit.score, t=round(time.time() - t0, 2))
            vr.audit = audit
            vr.plan = vplan
            (vdir / "plan.json").write_text(vplan.model_dump_json(indent=1), encoding="utf-8")
            (vdir / "audit.json").write_text(audit.model_dump_json(indent=1), encoding="utf-8")

            ts = time.time()
            final = vdir / f"{v.name}.pptx"
            shutil.copy(pptx, final)
            vr.pptx = str(final)
            if pdf:
                final_pdf = vdir / f"{v.name}.pdf"
                shutil.copy(pdf, final_pdf)
                vr.pdf = str(final_pdf)
                vr.html = str(await asyncio.to_thread(to_html, final_pdf, final, vdir / f"{v.name}.html", vplan.title,
                                                      profile.tokens.accent_hex, vplan.language))
            vr.pngs = [str(p) for p in pngs]
            vr.timings["export"] = round(time.time() - ts, 2)
            await _emit(emit, stage="export", variant=v.name, status="done", t=round(time.time() - t0, 2))
        except Exception as e:
            log.exception("variant %s failed", v.name)
            vr.error = str(e)
            await _emit(emit, stage="error", variant=v.name, message=str(e))
        return vr

    async def _render(self, pptx: Path, out: Path):
        """Рендер колоды в PDF и PNG: (картинки, pdf, успех); при сбое рендера — пустой результат."""
        try:
            pdf = await asyncio.to_thread(pptx_to_pdf, pptx, out, self.cfg.render.timeout_s)
            pngs = await asyncio.to_thread(pdf_to_pngs, pdf, out / "png", self.cfg.render.dpi)
            return pngs, pdf, True
        except RenderError as e:
            log.error("render failed: %s", e)
            return [], None, False

    async def _shorten(self, pptx: Path, overflows, llm: LLMClient) -> list[tuple[list[str], list[str]]]:
        """Сокращает текст, который не поместился в рамку, через скилл `shortener`; возвращает пары
        (абзацы до, абзацы после) для сверки содержания в аудите.
        """
        from pptx import Presentation

        skill = load_skill("shortener")
        items = [{"id": f"{o.slide}:{o.shape_id}", "paragraphs": o.paragraphs, "budget": o.budget} for o in overflows[:40]]
        try:
            res = await llm.run_skill(skill, None, items=items)
        except Exception as e:
            log.warning("shortener failed: %s", e)
            return []
        before = {f"{o.slide}:{o.shape_id}": o.paragraphs for o in overflows[:40]}
        rewrites = []
        prs = Presentation(str(pptx))
        slides = list(prs.slides)
        for it in (res or {}).get("items", []):
            try:
                si, sid = map(int, str(it["id"]).split(":"))
                sh = shape_by_id(slides[si], sid)
                if sh is not None and it.get("paragraphs"):
                    new = [str(p) for p in it["paragraphs"]]
                    set_paragraphs(sh, new)
                    rewrites.append((before.get(str(it["id"]), []), new))
            except Exception:
                continue
        prs.save(str(pptx))
        return rewrites

    def _manifest(self, run_id, rdir: Path, profile, corpus, brief, plan, mode, results, timings) -> Path:
        """Пишет манифест запуска: время этапов, версии скиллов, телеметрия модели, обоснования вёрстки."""
        skills_used = skills_manifest(self.agent)
        m = {
            "run_id": run_id,
            "agent": {"name": self.agent.name, "version": self.agent.version, "sha256": self.agent.sha256},
            "skills": skills_used,
            "models": {"llm": self.cfg.llm.model if self.cfg.llm.enabled else "offline", "vlm": self.cfg.llm.vlm_model or self.cfg.llm.model,
                       "t2i": self.cfg.image.model if self.cfg.image.provider != "none" else None},
            "template": {"id": profile.id, "name": profile.name, "sha256": profile.sha256, "parser": profile.parser_version},
            "content": {"id": corpus.id, "files": corpus.files, "chunks": len(corpus.chunks), "fits_brief": plan.materials_fit},
            "brief": brief.model_dump(),
            "plan_mode": mode,
            "timings_s": timings,
            "llm": self.telemetry.summary(),
            "variants": [
                {"name": r.name, "title": r.title, "pptx": r.pptx, "pdf": r.pdf, "html": r.html, "timings_s": r.timings,
                 "audit_score": r.audit.score if r.audit else None, "issues": len(r.audit.issues) if r.audit else None,
                 "audit_before_fix": r.audit_before_fix, "slides": r.slides, "error": r.error}
                for r in results
            ],
        }
        p = rdir / "manifest.json"
        p.write_text(json.dumps(m, ensure_ascii=False, indent=1), encoding="utf-8")
        return p
