"""Слепые шаблоны: на финальной защите используется шаблон, которого никто не видел.

Эти тесты собирают полные колоды на синтетических шаблонах с намеренно трудными
свойствами (нет слайдов-примеров, 4:3, колонтитулы на каждом слайде, заголовки свободными
текстовыми блоками, вообще нет плейсхолдеров, тёмная тема, вход .potx) и проверяют, что
результат чистый, соответствует шаблону и три варианта различаются.
"""
import asyncio
import re
import shutil
import zipfile
from pathlib import Path

import pytest

from decksmith.audit.context import AuditContext
from decksmith.audit.engine import run_audit
from decksmith.layout.builder import DeckBuilder
from decksmith.layout.selector import load_variants
from decksmith.parsing.template_parser import analyze_template
from decksmith.pipeline import plan_variants
from decksmith.render.soffice import render_pptx
from decksmith.testing.synthetic import GENERATORS

pytestmark = pytest.mark.slow
ROOT = Path(__file__).resolve().parents[1]
HARD_ERRORS = {"integrity.title_missing", "integrity.content_lost", "integrity.empty", "integrity.placeholder_text",
               "integrity.raster", "layout.out_of_bounds", "layout.text_clipped"}


def _build_all(template: Path, plan, out: Path):
    """Разбор шаблона и сборка трёх вариантов с рендером и аудитом."""
    prof = analyze_template(template, force=True)
    chosen = plan_variants(plan, prof, list(load_variants()))
    results = {}
    for name, v in load_variants().items():
        vplan, dec = chosen[name]
        b = DeckBuilder(prof, v)
        rep = b.build(vplan, dec)
        pptx = b.save(out / f"{template.stem}_{name}.pptx")
        _, pngs = render_pptx(pptx, out / f"{template.stem}_{name}_r", dpi=40)
        audit = asyncio.run(run_audit(AuditContext(pptx=pptx, profile=prof, plan=vplan, pngs=pngs,
                                                   slide_kinds=[s["kind"] for s in rep.slides],
                                                   slide_sources=[s["source"] for s in rep.slides])))
        results[name] = (rep, audit, len(pngs), len(vplan.slides))
    return prof, chosen, results


@pytest.mark.parametrize("gen", sorted(GENERATORS))
def test_synthetic_template_end_to_end(gen, plan, tmp_path):
    """Синтетический шаблон: полная колода без ошибок во всех вариантах."""
    t = GENERATORS[gen](tmp_path / f"{gen}.pptx")
    prof, chosen, results = _build_all(t, plan, tmp_path)
    for name, (rep, audit, n_pngs, n_expected) in results.items():
        assert n_pngs == n_expected, f"{gen}/{name}: slide count"
        hard = [i for i in audit.issues if i.check in HARD_ERRORS and i.severity.value == "error"]
        assert not hard, f"{gen}/{name}: {[(i.check, i.slide, i.message) for i in hard]}"
        assert audit.score >= 80, f"{gen}/{name}: score {audit.score}"
        assert not any("fallback after error" in s["rationale"] for s in rep.slides)
    keys = {n: [(d.mode, d.pattern_id or d.compose_kind) for d in dec] for n, (_, dec) in chosen.items()}
    diff = sum(a != b for a, b in zip(keys["balanced"], keys["visual"])) / len(keys["balanced"])
    assert diff >= 0.3, f"{gen}: variants too similar ({diff:.2f})"


def test_brand_texts_survive(plan, tmp_path):
    """Текст колонтитула, повторённый на каждом слайде-примере, — брендовый элемент: никогда не удаляется."""
    t = GENERATORS["brand_footer"](tmp_path / "brand.pptx")
    prof, chosen, results = _build_all(t, plan, tmp_path)
    from pptx import Presentation

    prs = Presentation(str(tmp_path / "brand_balanced.pptx"))
    cloned = [s for s, meta in zip(prs.slides, results["balanced"][0].slides) if meta["mode"] == "clone"]
    assert cloned
    for s in cloned:
        assert any("Конфиденциально" in sh.text_frame.text for sh in s.shapes if sh.has_text_frame)


def test_potx_input_is_accepted(tmp_path):
    """Шаблон .potx принимается наравне с .pptx."""
    src = GENERATORS["dark_minimal"](tmp_path / "t.pptx")
    potx = tmp_path / "t.potx"
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(potx, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"presentationml.presentation.main+xml", b"presentationml.template.main+xml")
            zout.writestr(item, data)
    prof = analyze_template(potx, force=True)
    assert prof.usable_patterns()


def test_profile_does_not_depend_on_file_name(tmp_path):
    """Профиль шаблона не зависит от имени файла."""
    a = GENERATORS["blank_layouts"](tmp_path / "a.pptx")
    b = tmp_path / "совсем_другое_имя.pptx"
    shutil.copy(a, b)
    pa, pb = analyze_template(a, force=True), analyze_template(b, force=True)
    assert [p.kind for p in pa.patterns] == [p.kind for p in pb.patterns]
    assert pa.tokens.model_dump() == pb.tokens.model_dump()


def test_no_dataset_specific_code():
    """Проверка против подгонки: пакет не должен упоминать шаблоны датасета."""
    banned = re.compile(r"vk[\s_-]?(tech|education|workspace)|google shape|свободный дизайн|\bVK\b", re.I)
    hits = []
    for p in (ROOT / "decksmith").rglob("*.py"):
        if "testing" in p.parts:
            continue
        for n, line in enumerate(p.read_text(encoding="utf-8").splitlines(), 1):
            if banned.search(line):
                hits.append(f"{p.relative_to(ROOT)}:{n}: {line.strip()[:80]}")
    assert not hits, "\n".join(hits)


def test_layout_frames_follow_document_order(unknown_template, tmp_path):
    """python-pptx перечисляет плейсхолдеры слайда по idx: рамки всё равно должны сопоставляться со своими
    двойниками в макете, когда порядок idx отличается от порядка в документе.
    """
    from pptx import Presentation

    prs = Presentation(unknown_template)
    layout = prs.slide_layouts[3]  # заголовок + две колонки контента
    left, right = sorted([ph for ph in layout.placeholders if ph.placeholder_format.idx in (1, 2)], key=lambda p: p.left)
    left._element.ph.set("idx", "20")
    right._element.ph.set("idx", "13")
    path = tmp_path / "idx_order.pptx"
    prs.save(path)
    prof = analyze_template(path, force=True)
    b = DeckBuilder(prof, load_variants()["balanced"])
    slide = b._slide_from_layout(b.layouts[3])
    frames = {ph.placeholder_format.idx: (ph.left, ph.top, ph.width, ph.height) for ph in b.layouts[3].placeholders}
    for ph in slide.placeholders:
        assert (ph.left, ph.top, ph.width, ph.height) == frames[ph.placeholder_format.idx], ph.name
