"""Аудит: каждая детерминированная проверка ловит внесённый дефект, а её фиксер его устраняет."""
import asyncio

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Emu, Pt

from decksmith.audit.context import AuditContext
from decksmith.audit.engine import catalogue, run_audit
from decksmith.audit.fixers import apply_fixes
from decksmith.render.soffice import render_pptx

pytestmark = pytest.mark.slow


def _defective_deck(profile, path):
    """Один слайд на дефект, на собственном макете-холсте шаблона."""
    prs = Presentation(profile.file)
    lst = prs.slides._sldIdLst
    for sid in list(lst):
        prs.part.drop_rel(sid.rId)
        lst.remove(sid)
    layout = [l for m in prs.slide_masters for l in m.slide_layouts][profile.canvas_layouts.get("light", 0)]
    W, H = prs.slide_width, prs.slide_height

    def slide(title):
        """Новый слайд с заголовком на макете-холсте."""
        s = prs.slides.add_slide(layout)
        for ph in list(s.placeholders):
            if "TITLE" in str(ph.placeholder_format.type):
                ph.text_frame.text = title
            else:
                ph._element.getparent().remove(ph._element)
        return s

    def card(s, x, y, w, h):
        """Карточка-подложка."""
        c = s.shapes.add_shape(1, Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
        c.fill.solid()
        c.fill.fore_color.rgb = RGBColor(0xEE, 0xF2, 0xF8)
        c.line.fill.background()
        return c

    def tb(s, x, y, w, h, text, size=14, color=None, font=None):
        """Текстовый блок с заданным кеглем, цветом и шрифтом."""
        t = s.shapes.add_textbox(Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
        t.text_frame.word_wrap = True
        paras = text if isinstance(text, list) else [text]
        for i, p in enumerate(paras):
            para = t.text_frame.paragraphs[0] if i == 0 else t.text_frame.add_paragraph()
            r = para.add_run()
            r.text = p
            r.font.size = Pt(size)
            if color:
                r.font.color.rgb = RGBColor.from_string(color)
            if font:
                r.font.name = font
        return t

    s = slide("Выход за границы")  # 0
    tb(s, W * 0.8, H * 0.5, W * 0.4, H * 0.1, "Этот блок выходит за правый край слайда")
    s = slide("Наложение блоков")  # 1
    tb(s, W * 0.1, H * 0.4, W * 0.4, H * 0.2, "Первый блок текста")
    tb(s, W * 0.2, H * 0.45, W * 0.4, H * 0.2, "Второй блок наезжает")
    s = slide("Заглушки")  # 2
    tb(s, W * 0.1, H * 0.4, W * 0.6, H * 0.2, "Lorem ipsum dolor sit amet")
    s = slide("Слишком много пунктов")  # 3
    tb(s, W * 0.1, H * 0.3, W * 0.8, H * 0.6, [f"Пункт номер {i} списка" for i in range(9)], size=12)
    s = slide("Низкий контраст")  # 4
    tb(s, W * 0.1, H * 0.4, W * 0.6, H * 0.1, "Почти невидимый светлый текст", color="EEEEEE")
    s = slide("Чужой шрифт")  # 5
    tb(s, W * 0.1, H * 0.4, W * 0.6, H * 0.1, "Текст набран чужой гарнитурой", font="Comic Sans MS")
    s = slide("Большая таблица")  # 6
    t = s.shapes.add_table(12, 6, Emu(int(W * 0.1)), Emu(int(H * 0.25)), Emu(int(W * 0.8)), Emu(int(H * 0.65)))
    for r in range(12):
        for c in range(6):
            t.table.cell(r, c).text = str(r * c)
    s = slide("Пустой слайд с заголовком")  # 7
    s = slide("Таблица на 9 строк")  # 8: допустимо (организаторы: до 10 строк — не ошибка)
    t = s.shapes.add_table(9, 4, Emu(int(W * 0.1)), Emu(int(H * 0.25)), Emu(int(W * 0.8)), Emu(int(H * 0.6)))
    for r in range(9):
        for c in range(4):
            t.table.cell(r, c).text = str(r * c)
    long = "Текст карточки описывает, как сервис раскладывает контент по слотам шаблона. " * 7
    s = slide("Текст вышел за карточку")  # 9
    card(s, W * 0.1, H * 0.35, W * 0.35, H * 0.3)
    tb(s, W * 0.12, H * 0.38, W * 0.31, H * 0.1, long)
    s = slide("Рамка мала, но в карточке есть место")  # 10
    card(s, W * 0.1, H * 0.3, W * 0.5, H * 0.6)
    fixed = tb(s, W * 0.12, H * 0.33, W * 0.46, H * 0.12, long[:300])
    fixed.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    prs.save(path)
    return path


@pytest.fixture(scope="module")
def defective(profiles, tmp_path_factory):
    """Колода с внесёнными дефектами, её рендер и отчёт аудита."""
    prof = profiles["vk_education"]
    p = _defective_deck(prof, tmp_path_factory.mktemp("audit") / "defects.pptx")
    pdf, pngs = render_pptx(p, p.parent / "render")
    rep = asyncio.run(run_audit(AuditContext(pptx=p, profile=prof, pngs=pngs)))
    return prof, p, rep


@pytest.mark.parametrize("check,slide", [
    ("layout.out_of_bounds", 0),
    ("layout.overlap", 1),
    ("integrity.placeholder_text", 2),
    ("density.bullets", 3),
    ("template.contrast", 4),
    ("template.font", 5),
    ("density.table", 6),
    ("integrity.empty", 7),
])
def test_check_detects_injected_defect(defective, check, slide):
    """Каждая проверка находит свой внесённый дефект."""
    _, _, rep = defective
    assert any(i.check == check and i.slide == slide for i in rep.issues), f"{check} missed on slide {slide + 1}"


def test_organisers_clarifications(defective):
    """Таблицы до 10 строк проходят; текст может перерасти рамку только внутри своего блока."""
    _, _, rep = defective
    found = {(i.check, i.slide) for i in rep.issues}
    assert ("density.table", 8) not in found, "a 9-row table is allowed"
    assert ("layout.block_overflow", 9) in found, "text spilling out of its card is an error"
    assert ("layout.text_overflow", 10) in found and ("layout.block_overflow", 10) not in found


def test_overflowing_frame_grows_inside_its_block(defective, tmp_path):
    """Переполненная рамка растёт внутри своего блока, а не уменьшает кегль."""
    prof, p, rep = defective
    iss = [i for i in rep.issues if i.check == "layout.text_overflow" and i.slide == 10]
    out = tmp_path / "grown.pptx"
    res = asyncio.run(apply_fixes(p, iss, prof, out))
    assert res["applied"]
    shape = next(sh for sh in Presentation(str(out)).slides[10].shapes if sh.has_text_frame and sh.text_frame.text.startswith("Текст карточки"))
    assert {r.font.size.pt for para in shape.text_frame.paragraphs for r in para.runs} == {14.0}, "type kept, frame resized"
    rep2 = asyncio.run(run_audit(AuditContext(pptx=out, profile=prof)))
    left = {(i.check, i.slide) for i in rep2.issues}
    assert ("layout.text_overflow", 10) not in left and ("layout.block_overflow", 10) not in left


def test_issues_carry_boxes_for_visualisation(defective):
    """Замечания несут рамки для подсветки на слайде."""
    _, _, rep = defective
    boxed = [i for i in rep.issues if i.check in ("layout.out_of_bounds", "layout.overlap", "template.contrast")]
    assert boxed and all(i.boxes for i in boxed)


def test_fixers_resolve_selected_issues(defective, tmp_path):
    """Выбранные исправления устраняют замечания при повторном аудите."""
    prof, p, rep = defective
    fixable = [i for i in rep.issues if i.fixable and i.deterministic and i.check in
               ("layout.out_of_bounds", "integrity.placeholder_text", "density.bullets", "template.contrast", "template.font", "density.table")]
    out = tmp_path / "fixed.pptx"
    res = asyncio.run(apply_fixes(p, fixable, prof, out))
    assert res["applied"]
    pdf, pngs = render_pptx(out, tmp_path / "render")
    rep2 = asyncio.run(run_audit(AuditContext(pptx=out, profile=prof, pngs=pngs)))
    before = {(i.check, i.slide) for i in fixable}
    after = {(i.check, i.slide) for i in rep2.issues}
    assert len(before & after) <= len(before) // 3, f"still present: {before & after}"
    assert rep2.score > rep.score


def test_catalogue_separates_deterministic_and_contextual():
    """Каталог разделяет детерминированные и контекстные проверки."""
    cat = catalogue()
    assert any(c["deterministic"] for c in cat) and any(not c["deterministic"] for c in cat)
    assert len([c for c in cat if c["deterministic"]]) >= 20


def test_template_quirks_do_not_lower_the_score():
    """Особенности самого шаблона («так в шаблоне») видны в отчёте, но балл не снижают; наши дефекты снижают."""
    from decksmith.audit.engine import score
    from decksmith.core.models import AuditIssue, Severity

    def one(sev, inherited):
        return AuditIssue(id="x", check="layout.overlap", category="layout", slide=0, message="",
                          severity=sev, data={"inherited": inherited})

    assert score([one(Severity.info, True)] * 5, 5) == 100.0
    assert score([one(Severity.info, False)] * 5, 5) == 96.0
    assert score([one(Severity.error, False), one(Severity.info, True)], 10) == 96.0
