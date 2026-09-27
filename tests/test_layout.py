"""Слой вёрстки: детерминированный выбор, чистый результат, нативные объекты, три различных варианта."""
import re

import pytest
from pptx import Presentation
from pptx.util import Inches

from decksmith.layout.builder import DeckBuilder
from decksmith.layout.selector import load_variants, select_layouts

pytestmark = pytest.mark.slow
LEFTOVER = re.compile(r"lorem ipsum|^\s*(заголовок|текст|пункт|описание|имя фамилия|должность)\s*$|вставить фото", re.I | re.M)


def _build(profile, plan, variant, tmp_path):
    """Выбор паттернов, сборка и сохранение колоды."""
    decisions = select_layouts(plan, profile, variant)
    b = DeckBuilder(profile, variant)
    rep = b.build(plan, decisions)
    return b.save(tmp_path / f"{profile.name}_{variant.name}.pptx"), decisions, rep


def test_selection_is_deterministic(profiles, plan):
    """Одинаковые входные данные дают одинаковый выбор паттернов."""
    v = load_variants()["balanced"]
    a = [d.model_dump() for d in select_layouts(plan, profiles["vk_tech"], v)]
    b = [d.model_dump() for d in select_layouts(plan, profiles["vk_tech"], v)]
    assert a == b


@pytest.mark.parametrize("name", ["vk_tech", "vk_workspace", "vk_education"])
def test_build_is_clean_and_native(profiles, plan, tmp_path, name):
    """Колода чистая: нативные объекты, без текста шаблона и пустых слайдов."""
    prof = profiles[name]
    path, decisions, rep = _build(prof, plan, load_variants()["balanced"], tmp_path)
    prs = Presentation(str(path))
    assert len(prs.slides) == len(plan.slides)
    layout_names = {l.name for l in prof.layouts}
    charts = tables = 0
    for s in prs.slides:
        assert s.slide_layout.name in layout_names, "every slide sits on a template layout"
        for sh in s.shapes:
            if sh.has_text_frame:
                assert not LEFTOVER.search(sh.text_frame.text), f"template filler left: {sh.text_frame.text!r}"
            charts += bool(getattr(sh, "has_chart", False) and sh.has_chart)
            tables += bool(getattr(sh, "has_table", False) and sh.has_table)
        big_pics = [sh for sh in s.shapes if sh.shape_type == 13 and sh.width * sh.height > 0.85 * prs.slide_width * prs.slide_height]
        texts = [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        assert not (big_pics and not texts), "slide exported as a single raster image"
    assert charts >= 1 and tables >= 1, "data slides are native chart/table objects"
    assert all(d.rationale for d in decisions)


def test_variants_differ(profiles, plan, tmp_path):
    """Три варианта заметно различаются выбором паттернов."""
    prof = profiles["vk_education"]
    choices = {}
    for name, v in load_variants().items():
        from decksmith.pipeline import apply_variant

        vplan = apply_variant(plan, v)
        choices[name] = [(d.mode, d.pattern_id or d.compose_kind) for d in select_layouts(vplan, prof, v)]
    assert choices["balanced"] != choices["visual"] != choices["dense"]


def test_unknown_template_builds(unknown_template, plan, tmp_path):
    """Незнакомый шаблон разбирается и собирается."""
    from decksmith.parsing.template_parser import analyze_template

    prof = analyze_template(unknown_template)
    path, decisions, _ = _build(prof, plan, load_variants()["balanced"], tmp_path)
    assert len(Presentation(str(path)).slides) == len(plan.slides)


def test_text_frame_grows_inside_its_card(unknown_template, tmp_path):
    """Уточнение организаторов: текстовая рамка может расти внутри своего блока, поэтому кегль сохраняется."""
    from decksmith.core.models import DeckPlan, Item, PatternKind, SlideLayout, SlideSpec
    from decksmith.parsing.template_parser import analyze_template

    prof = analyze_template(unknown_template)
    pat = next(p for p in prof.patterns if p.repeaters and p.n_items == 3 and p.source == "slide")
    text = ("Сервис читает шаблон как набор правил: палитру, шрифты, сетку и паттерны примеров, а затем раскладывает "
            "контент по слотам без ручной вёрстки и проверяет результат аудитом; три варианта колоды отличаются "
            "плотностью, визуализацией и порядком подачи материала")
    plan = DeckPlan(title="Проба", slides=[SlideSpec(id="s1", intent=PatternKind.cards, title="Карточки с длинным текстом",
                                                  items=[Item(title=f"Карточка {i + 1}", text=text) for i in range(3)])])
    b = DeckBuilder(prof, load_variants()["balanced"])
    b.build(plan, [SlideLayout(spec_id="s1", mode="clone", pattern_id=pat.id, keep_items=3)])
    slide = Presentation(str(b.save(tmp_path / "grown.pptx"))).slides[0]
    cards = [sh for sh in slide.shapes if sh.has_text_frame and not sh.text_frame.text.strip()]
    bodies = [sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.startswith("Сервис читает")]
    assert len(bodies) == 3 and len(cards) == 3
    for body in bodies:
        card = next(c for c in cards if c.left <= body.left and body.left + body.width <= c.left + c.width)
        assert body.height > Inches(1.6), "the frame grew instead of shrinking the type"
        assert body.top + body.height <= card.top + card.height, "and stayed inside its card"
        assert {r.font.size.pt for p in body.text_frame.paragraphs for r in p.runs} == {14.0}
