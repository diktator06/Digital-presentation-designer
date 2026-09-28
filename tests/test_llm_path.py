"""Путь через модель от начала до конца против эмулятора OpenAI-совместимого API в том же процессе:
рендер скиллов, формат запросов (JSON-режим, картинки для VLM), проверка схемы,
раунд исправления JSON, сопоставление визуального аудита и контекстные исправления.
"""
import asyncio
from pathlib import Path


from decksmith.core.config import LLMConfig, settings
from decksmith.testing.llm_emulator import Emulator

ROOT = Path(__file__).resolve().parents[1]
CFG = LLMConfig(provider="openai_compatible", base_url="http://emulator/v1", model="Qwen/Qwen3-32B",
                vlm_model="Qwen/Qwen2.5-VL-32B-Instruct", max_retries=0)


def test_plan_via_llm_with_repair(profiles):
    from decksmith.content.ingest import ingest
    from decksmith.generation.llm import LLMClient
    from decksmith.generation.planner import Brief, make_plan
    from decksmith.generation.skills import load_agent

    emu = Emulator()
    corpus = ingest([], "Сервис генерации презентаций: 10–15 слайдов, 5 минут, 3 варианта.")

    async def go():
        """План через эмулятор модели."""
        async with LLMClient(CFG, transport=emu.transport()) as llm:
            return await make_plan(Brief(text="Сервис", n_slides=9), corpus, profiles["vk_tech"], llm, load_agent()), llm

    (plan, mode), llm = asyncio.run(go())
    assert mode == "llm"
    skills = [r["skill"] for r in emu.requests]
    assert "outline@v3" in skills and "outline@v3:repair" in skills  # неверная схема -> один раунд исправления
    # авторы текстов параллельно, по одному на содержательный слайд
    assert sum(s == "slide_writer@v3" for s in skills) >= 6
    assert all(r["body"].get("response_format") == {"type": "json_object"} for r in emu.requests)
    intents = [s.intent.value for s in plan.slides]
    assert intents[0] == "title" and intents[-1] == "thanks"
    chart = next(s for s in plan.slides if s.intent.value == "chart")
    assert chart.chart and chart.chart.series[0].values == [40, 10, 20, 30]
    assert llm.telemetry.summary()["calls"] == len(emu.requests)


def test_materials_off_the_brief_topic_are_not_used(profiles):
    """Бриф о лишнем весе, а материалы — о генераторе презентаций: модель помечает материалы как не по теме,
    авторы текстов их не получают, но видят бриф; план помнит оценку (её показывает интерфейс).
    """
    from decksmith.content.ingest import ingest
    from decksmith.generation.llm import LLMClient
    from decksmith.generation.planner import Brief, make_plan
    from decksmith.generation.skills import load_agent

    emu = Emulator(materials_fit="none")
    corpus = ingest([], "Сервис генерации презентаций: колода за 5 минут, экспорт в PPTX и аудит.")
    brief = Brief(text="о проблеме лишнего веса", n_slides=9)

    async def go():
        """План через эмулятор модели."""
        async with LLMClient(CFG, transport=emu.transport()) as llm:
            return await make_plan(brief, corpus, profiles["vk_tech"], llm, load_agent())

    plan, mode = asyncio.run(go())
    assert mode == "llm" and plan.materials_fit == "none"
    writers = [r["body"]["messages"][-1]["content"] for r in emu.requests if r["skill"].startswith("slide_writer")]
    assert writers
    assert all("о проблеме лишнего веса" in w for w in writers)
    assert not any("колода за 5 минут" in w for w in writers)
    # план видит материалы (чтобы их оценить), аудит спрашивает о теме брифа
    outline = next(r["body"]["messages"][-1]["content"] for r in emu.requests if r["skill"] == "outline@v3")
    assert "колода за 5 минут" in outline and "о проблеме лишнего веса" in outline

    # страховка от ошибки модели: слова брифа есть в материалах — материалы по теме хотя бы частично
    from decksmith.generation.planner import Outline, materials_fit
    none = Outline(title="t", materials_fit="none", slides=[])
    assert materials_fit(none, corpus, "Сервис генерации презентаций") == "partial"
    assert materials_fit(none, corpus, "о проблеме лишнего веса") == "none"
    assert materials_fit(none, ingest([], ""), "о проблеме лишнего веса") == "full"  # без материалов сравнивать не с чем


def test_pipeline_with_model_path(profiles, monkeypatch, tmp_path):
    """Весь пайплайн через путь модели (эмулятор): план, тексты, визуальный аудит."""
    from decksmith.content.ingest import ingest
    from decksmith.generation.planner import Brief
    from decksmith.pipeline import Pipeline

    monkeypatch.setattr(settings(), "llm", CFG)
    emu = Emulator(fail_visual_on={2})
    pipe = Pipeline(transport=emu.transport())
    corpus = ingest([], "Сервис генерации презентаций: 10–15 слайдов, 5 минут, 3 варианта.")
    res = asyncio.run(pipe.run(profiles["vk_education"], corpus, Brief(text="Сервис", n_slides=9), ["balanced"]))
    v = res.variants[0]
    assert not v.error and Path(v.pptx).exists()
    vlm = [r for r in emu.requests if r["skill"].startswith("visual_audit")]
    assert vlm and vlm[0]["model"] == CFG.vlm_model
    content = vlm[0]["body"]["messages"][-1]["content"]
    assert any(p.get("type") == "image_url" and p["image_url"]["url"].startswith("data:image/png;base64,") for p in content)
    assert any("q12" in p.get("text", "") and "Тема презентации (бриф пользователя): «Сервис»" in p.get("text", "")
               for p in content)  # вопрос о теме брифа
    ctx_issues = [i for i in v.audit.issues if not i.deterministic]
    assert any(i.check == "content.title_is_conclusion" and i.slide == 2 for i in ctx_issues)
    assert res.plan_mode == "llm"


def test_contextual_fix_via_fixer_skill(profiles, tmp_path):
    """Контекстное исправление идёт через скилл `fixer`."""
    from pptx import Presentation

    from decksmith.audit.fixers import apply_fixes
    from decksmith.core.models import AuditIssue
    from decksmith.generation.llm import LLMClient

    prof = profiles["vk_tech"]
    src = Path(prof.file)
    emu = Emulator()
    issue = AuditIssue(id="x", check="content.title_is_conclusion", category="content", deterministic=False, slide=0,
                       message="Заголовок не содержит вывода", fixable=True, fix="fix_llm")

    async def go():
        """Исправление через эмулятор модели."""
        async with LLMClient(CFG, transport=emu.transport()) as llm:
            return await apply_fixes(src, [issue], prof, tmp_path / "fixed.pptx", llm)

    res = asyncio.run(go())
    assert res["contextual_changed"] >= 1
    texts = [sh.text_frame.text for sh in Presentation(str(tmp_path / "fixed.pptx")).slides[0].shapes if sh.has_text_frame]
    assert "Исправленный заголовок-вывод" in texts
    # правка возвращается парой «было -> стало»: по ней план для повторного аудита получает новый текст
    from decksmith.core.models import DeckPlan, SlideSpec
    from decksmith.pipeline import apply_rewrites

    old, new = res["rewrites"][0]
    plan = DeckPlan(title="t", slides=[SlideSpec(id="s1", intent="title", title=old[0])])
    assert new == ["Исправленный заголовок-вывод"]
    assert apply_rewrites(plan, res["rewrites"]).slides[0].title == "Исправленный заголовок-вывод"
