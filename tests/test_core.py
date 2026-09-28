"""Модульные тесты: геометрия, цвета, помощники наследования OOXML, извлечение JSON, реестр скиллов."""
import pytest

from decksmith.core.models import Box
from decksmith.generation.llm import extract_json
from decksmith.parsing.ooxml import color_distance, contrast_ratio, flatten_shapes


def test_box_ops():
    """Операции с рамками: пересечение, объединение, вложенность."""
    a, b = Box(x=0, y=0, w=10, h=10), Box(x=5, y=5, w=10, h=10)
    assert a.intersection(b) == 25
    assert a.union(b) == Box(x=0, y=0, w=15, h=15)
    assert a.contains(Box(x=1, y=1, w=2, h=2))


def test_contrast_wcag():
    assert round(contrast_ratio("000000", "FFFFFF"), 1) == 21.0
    assert contrast_ratio("8F8F8F", "FFFFFF") < 4.5  # серый шаблона не проходит WCAG AA
    assert color_distance("0077FF", "0077FF") == 0


def test_group_transform_flattening():
    """Дочерние фигуры масштабированной группы получают абсолютные координаты слайда."""
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    g = s.shapes.add_group_shape()
    g.shapes.add_shape(1, Emu(1000), Emu(1000), Emu(1000), Emu(1000))
    recs = [r for r in flatten_shapes(s.shapes) if r.depth == 1]
    assert recs and recs[0].box.w > 0


@pytest.mark.parametrize("raw", [
    '```json\n{"a": 1}\n```',
    'Вот ответ: {"a": 1} надеюсь помог',
    '<think>...</think>{"a": {"b": [1, 2]}}',
    '{"a": [1, 2',  # обрезанный вывод закрывается
])
def test_extract_json(raw):
    """JSON извлекается из ответа модели в любом обрамлении."""
    assert "a" in extract_json(raw)


def test_skills_registry_and_rendering():
    """Реестр скиллов: версии, загрузка и рендер промптов."""
    from decksmith.generation.skills import list_versions, load_agent, load_skill

    v = list_versions()
    for name in ("outline", "slide_writer", "shortener", "visual_audit", "fixer"):
        assert v["skills"][name]["active"] in v["skills"][name]["versions"]
        sk = load_skill(name)
        assert sk.sha256 and sk.user
    agent = load_agent()
    assert agent.step("outline")["skill"].startswith("outline@")
    ctx = dict(brief="b", purpose="product", purpose_hint="", audience="", n_slides=12, capabilities="- cards",
               context="ctx", language="ru", title_chars=60)
    _, user = load_skill("outline").render(**ctx, with_images=False)
    assert "12" in user and "ctx" in user and "intent=image_text в" not in user
    # при включённой генерации изображений структура просит слайд-иллюстрацию
    _, user = load_skill("outline").render(**ctx, with_images=True)
    assert "intent=image_text в" in user


def test_font_download_is_time_bounded(monkeypatch):
    """Отсутствие сети (или медленная сеть) никогда не должно задерживать разбор шаблона: один неудачный
    запрос на время отключает загрузки.
    """
    import time

    import httpx

    from decksmith.core import fonts

    calls = []

    class Offline(httpx.BaseTransport):
        def handle_request(self, request):
            """Транспорт без сети: каждый запрос падает по таймауту."""
            calls.append(str(request.url))
            raise httpx.ConnectTimeout("offline")

    real_client = httpx.Client
    monkeypatch.setattr(fonts.httpx, "Client", lambda **kw: real_client(transport=Offline(), **kw))
    monkeypatch.setattr(fonts, "_NET_DOWN_UNTIL", 0.0)
    t0 = time.monotonic()
    assert not fonts.try_download_google_font("Some Open Family")
    assert not fonts.try_download_google_font("Another Open Family")
    assert len(calls) == 1 and time.monotonic() - t0 < 2
    assert fonts.split_weight("Montserrat Medium") == ("Montserrat", "medium")
    assert fonts.split_weight("Open Sans Extra Bold") == ("Open Sans", "extrabold")
    assert fonts.split_weight("Open Sans") == ("Open Sans", None)


def test_render_timeout_kills_process_group():
    """Зависший рендерер, чьи дочерние процессы продолжают работать, не должен блокировать дольше таймаута."""
    import subprocess
    import sys
    import time

    import pytest

    from decksmith.render.soffice import _run

    if sys.platform.startswith("win"):
        pytest.skip("POSIX process groups")
    t0 = time.monotonic()
    with pytest.raises(subprocess.TimeoutExpired):
        _run(["/bin/sh", "-c", "sleep 30 & sleep 30"], 1)
    assert time.monotonic() - t0 < 5


def test_offline_plan_meets_slide_range():
    """Без модели в колоде заданное число слайдов (ТЗ: 10–15 или сколько задал пользователь) и чистая обложка.
    """
    from pathlib import Path

    from decksmith.content.ingest import ingest
    from decksmith.generation.planner import Brief, plan_offline

    root = Path(__file__).resolve().parents[1]
    corpus = ingest(sorted((root / "data" / "content").glob("*.md")))
    brief = Brief(text="Цифровой дизайнер презентаций — сервис, который по брифу и контент-пакету собирает "
                       "презентацию в фирменном шаблоне компании и сам проверяет её качество.",
                  purpose="product", n_slides=12, language="ru")
    plan = plan_offline(brief, corpus)
    assert 10 <= len(plan.slides) <= 15
    assert plan.title == "Цифровой дизайнер презентаций"
    assert plan.subtitle.startswith("Сервис") and plan.subtitle != "product"
    assert not any(s.title.startswith("#") or "|" in s.title for s in plan.slides)
    for n in (6, 8, 12):
        brief.n_slides = n
        assert len(plan_offline(brief, corpus).slides) == n, n


def test_model_plan_is_completed_to_the_requested_count():
    """Модель, ответившая на один слайд меньше (замечено у Qwen на хостинге), всё равно даёт заданное число.
    """
    from decksmith.core.models import DeckPlan, Item, PatternKind, SlideSpec
    from decksmith.generation.planner import Brief, Outline, OutlineSlide, _normalize_outline, complete_plan

    o = Outline(title="Колода", slides=[OutlineSlide(intent="title", title="Колода")]
                + [OutlineSlide(intent="cards", title=f"Вывод {i}") for i in range(8)]
                + [OutlineSlide(intent="thanks", title="Спасибо")])
    o = _normalize_outline(o, Brief(text="Колода", n_slides=12, language="ru"))
    assert len(o.slides) == 12
    assert o.slides[-2].message == "summary" and o.slides[1].intent == "agenda"

    plan = DeckPlan(title="Колода", slides=[
        SlideSpec(id="s1", intent=PatternKind.title, title="Колода"),
        SlideSpec(id="s2", intent=PatternKind.cards, title="Шесть причин", items=[Item(title=f"П{i}") for i in range(6)]),
        SlideSpec(id="s3", intent=PatternKind.text, title="Главное", notes="\n[summary]"),
        SlideSpec(id="s4", intent=PatternKind.thanks, title="Спасибо")])
    complete_plan(plan, 5)
    assert len(plan.slides) == 5 and [s.id for s in plan.slides] == [f"s{i}" for i in range(1, 6)]
    assert [len(s.items) for s in plan.slides[1:3]] == [3, 3], "the longest list is split in two"
    assert plan.slides[3].bullets, "an empty summary lists the conclusions"


def test_every_variant_keeps_the_slide_count():
    """ТЗ: заданное пользователем число слайдов выдерживается во всех трёх вариантах."""
    from pathlib import Path

    from decksmith.core.models import DeckPlan
    from decksmith.layout.selector import load_variants
    from decksmith.pipeline import apply_variant

    plan = DeckPlan.model_validate_json((Path(__file__).parent / "fixtures" / "plan_sample.json").read_text(encoding="utf-8"))
    for name, v in load_variants().items():
        vp = apply_variant(plan, v)
        assert len(vp.slides) == len(plan.slides), name
    dense = apply_variant(plan, load_variants()["dense"])
    assert "[summary]" in (dense.slides[1].notes or ""), "the analytic variant opens with its conclusions"


def test_empty_env_value_takes_default(monkeypatch):
    """Как ${VAR:-default} в shell: пустая переменная из .env берёт значение по умолчанию."""
    from decksmith.core.config import _interpolate

    monkeypatch.setenv("DECKSMITH_TEST_VAR", "")
    assert _interpolate("${DECKSMITH_TEST_VAR:-по умолчанию}") == "по умолчанию"
    monkeypatch.setenv("DECKSMITH_TEST_VAR", "задано")
    assert _interpolate({"a": ["${DECKSMITH_TEST_VAR:-по умолчанию}"]}) == {"a": ["задано"]}
    monkeypatch.delenv("DECKSMITH_TEST_VAR")
    assert _interpolate("${DECKSMITH_TEST_VAR:-}") == ""


def test_skill_files_are_reread_after_change(tmp_path, monkeypatch):
    """Правка промпта на диске видна работающему сервису без перезапуска; недоступный скилл отмечается в
    манифесте, а не роняет запуск.
    """
    import os

    from decksmith.generation import skills

    monkeypatch.setattr(skills, "_root", lambda: tmp_path)
    (tmp_path / "registry.yaml").write_text("skills:\n  demo: {active: v1}\nagents: {}\n", encoding="utf-8")
    (tmp_path / "demo").mkdir()
    f = tmp_path / "demo" / "v1.yaml"
    f.write_text("user: первая версия\n", encoding="utf-8")
    assert skills.load_skill("demo").user == "первая версия"
    f.write_text("user: вторая версия\n", encoding="utf-8")
    st = f.stat()
    os.utime(f, ns=(st.st_atime_ns, st.st_mtime_ns + 1_000_000))  # гарантированно новое время изменения
    assert skills.load_skill("demo").user == "вторая версия"

    agent = skills.Agent(name="a", version="v1", path=f, sha256="", steps=[
        {"name": "ok", "skill": "demo@v1"}, {"name": "gone", "skill": "removed@v1", "enabled": False}])
    m = skills.skills_manifest(agent)
    assert m["ok"]["skill"] == "demo@v1" and m["ok"]["sha256"]
    assert "FileNotFoundError" in m["gone"]["error"]


def test_caps_and_letter_spacing_are_measured():
    """Капс (cap="all") и разрядка (spc) шаблона читаются из OOXML и расширяют замер: иначе рендерер рвёт
    слова заголовка посередине.
    """
    from pptx import Presentation
    from pptx.oxml.ns import qn
    from pptx.util import Emu

    from decksmith.layout.textfit import measure
    from decksmith.parsing.ooxml import StyleResolver, parse_theme

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(3000000), Emu(600000))
    tb.text_frame.text = "Визуализация"
    rpr = tb.text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr()
    rpr.set("cap", "all")
    rpr.set("spc", "300")
    p = tb.text_frame.paragraphs[0]
    st = StyleResolver(slide, parse_theme(slide.slide_layout.slide_master)).resolve(tb, p, p.runs[0])
    assert st.caps is True and st.tracking == 3.0
    assert rpr.tag == qn("a:rPr")

    plain = measure(["Визуализация"], "Arial", 24, 3000000, True)
    drawn = measure(["Визуализация"], "Arial", 24, 3000000, True, caps=True, tracking=3.0)
    assert drawn.longest_word_emu > plain.longest_word_emu * 1.2


def test_footer_title_filler_gets_the_deck_title():
    """Колонтитул-заглушка «Название презентации» получает название колоды (длинное — сокращается по слову),
    фирменный колонтитул остаётся как в шаблоне.
    """
    from pptx import Presentation

    from decksmith.layout.builder import DeckBuilder

    prs = Presentation()
    footers = [next(ph for ph in lay.placeholders if "FOOTER" in str(ph.placeholder_format.type))
               for lay in list(prs.slide_layouts)[:3]]
    for ph, text in zip(footers, ["НАЗВАНИЕ ПРЕЗЕНТАЦИИ", "ООО «Альфа», конфиденциально", "Presentation title"]):
        ph.text_frame.text = text

    DeckBuilder._title_footers(list(prs.slide_layouts)[:2], "Решение проблемы веса")
    assert [f.text_frame.text for f in footers[:2]] == ["Решение проблемы веса", "ООО «Альфа», конфиденциально"]

    long = "Цифровой дизайнер презентаций для корпоративных шаблонов любой сложности"
    DeckBuilder._title_footers([prs.slide_layouts[2]], long)
    short = footers[2].text_frame.text
    assert short.endswith("…") and len(short) <= 2 * len("Presentation title") + 1 and long.startswith(short[:-1])
