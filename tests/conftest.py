import os
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
os.environ.setdefault("DECKSMITH_LLM_PROVIDER", "offline")  # тесты никогда не вызывают модель
# свой кэш: тестовые шаблоны не должны появиться в списке шаблонов сервиса
os.environ.setdefault("DECKSMITH_WORKSPACE", str(ROOT / "workspace" / "_test"))
TEMPLATES = sorted((ROOT / "data" / "templates").glob("*.pptx"))


@pytest.fixture(scope="session")
def plan():
    """Пример плана колоды из фикстуры."""
    from decksmith.core.models import DeckPlan

    return DeckPlan.model_validate_json((ROOT / "tests" / "fixtures" / "plan_sample.json").read_text(encoding="utf-8"))


@pytest.fixture(scope="session")
def profiles():
    """Шаблоны датасета организатора (не распространяются в репозитории: положите их в data/templates/)."""
    from decksmith.parsing.template_parser import analyze_template

    if not TEMPLATES:
        pytest.skip("dataset templates not present in data/templates/")
    return {t.stem: analyze_template(t) for t in TEMPLATES}


@pytest.fixture(scope="session")
def unknown_template(tmp_path_factory) -> Path:
    """Шаблон, который система никогда не видела: стандартная тема Office с несколькими слайдами-примерами."""
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Название презентации"
    s.placeholders[1].text = "Подзаголовок"
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Заголовок"
    s.placeholders[1].text_frame.text = "Пункт"
    for _ in range(3):
        s.placeholders[1].text_frame.add_paragraph().text = "Пункт"
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Заголовок"
    for i in range(3):
        box = s.shapes.add_shape(5, Inches(0.8 + i * 4.1), Inches(2.2), Inches(3.8), Inches(3.5))
        box.fill.solid()
        box.fill.fore_color.rgb = __import__("pptx.dml.color", fromlist=["RGBColor"]).RGBColor(0xEE, 0xF2, 0xF8)
        t = s.shapes.add_textbox(Inches(1.0 + i * 4.1), Inches(2.4), Inches(3.4), Inches(0.6))
        t.text_frame.text = "Заголовок"
        t.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
        t.text_frame.paragraphs[0].runs[0].font.bold = True
        d = s.shapes.add_textbox(Inches(1.0 + i * 4.1), Inches(3.1), Inches(3.4), Inches(1.6))
        d.text_frame.text = "Описание"
        d.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    p = tmp_path_factory.mktemp("unknown") / "unknown_office.pptx"
    prs.save(p)
    return p
